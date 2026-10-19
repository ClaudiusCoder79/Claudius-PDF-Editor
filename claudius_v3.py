"""
Claudius PDF Editor - version 3 features, part 1:
settings window, remembered window / panel / reading position, autosave settings, undo history, crop & resize pages,
presets for header / footer / watermark / page numbers, edit images, signature blocks, highlight colour sets and export,
look up / translate, redaction check, metadata cleaner.
(Part 2 - forms, batch, compress, merge picker, folder search - is in claudius_v3b.py.)
Mixed into PdfEditor by claudius.py.
"""
import copy
import datetime
import io
import json
import os
import re
import shutil
import time
import tkinter as tk
import urllib.parse
import webbrowser
from tkinter import ttk, filedialog, messagebox

try:
    import pymupdf as fitz
except ImportError:
    import fitz
from PIL import Image, ImageDraw, ImageTk

import claudius_theme as TH
from claudius_i18n import tr as _tr
import claudius_extra as EX
from claudius_extra import APP, UI_FONT, BG, ACCENT, S, CONFIG_DIR, find_font, load_pil_font
from claudius_view import get_setting, set_setting

# ====================================================================================== settings store
V3_DEFAULTS = {
    "remember_window": True, "remember_view": True, "remember_position": True,
    "default_view": "page",                    # page | width | 100
    "autosave": True, "autosave_minutes": 1, "recovery_days": 14,
    "printer": "", "copies": 1,
    "panel": True, "sidebar_mode": "pages",
    "hl_names": ["Important", "Agree", "Question", "Reference"],
    "hl_colors": None,
    "presets": {}, "positions": {},
}
MAX_POSITIONS = 300


def load_cfg():
    cfg = copy.deepcopy(V3_DEFAULTS)
    try:
        saved = get_setting("v3", {}) or {}
        if isinstance(saved, dict):
            for k, v in saved.items():
                if k in cfg and type(v) is type(cfg[k]) or (k in cfg and cfg[k] is None and isinstance(v, list)):
                    cfg[k] = v
    except Exception:
        pass
    cfg["autosave_minutes"] = max(1, min(30, int(cfg["autosave_minutes"] or 1)))
    cfg["recovery_days"] = max(1, min(90, int(cfg["recovery_days"] or 14)))
    cfg["copies"] = max(1, min(99, int(cfg["copies"] or 1)))
    if cfg["default_view"] not in ("page", "width", "100"):
        cfg["default_view"] = "page"
    if len(cfg["hl_names"]) != 4:
        cfg["hl_names"] = list(V3_DEFAULTS["hl_names"])
    return cfg


def save_cfg(cfg):
    try:
        set_setting("v3", cfg)
    except Exception:
        pass


def norm_path(p):
    try:
        return os.path.normcase(os.path.abspath(p))
    except Exception:
        return p


# ====================================================================================== helpers
def pix_to_photo(pix, master=None):
    mode = "RGBA" if pix.alpha else "RGB"
    img = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
    return ImageTk.PhotoImage(img, master=master)


def image_png_bytes(doc, xref):
    """PNG (or the original JPEG/PNG) bytes of an image object, with transparency when it has a mask"""
    info = doc.extract_image(xref)
    if info and not info.get("smask") and info.get("ext") in ("png", "jpg", "jpeg"):
        return info["image"], info["ext"]
    pix = fitz.Pixmap(doc, xref)
    if info and info.get("smask"):
        try:
            pix = fitz.Pixmap(fitz.Pixmap(fitz.csRGB, pix) if pix.colorspace and pix.colorspace.n != 3 else pix,
                              fitz.Pixmap(doc, info["smask"]))
        except Exception:
            pass
    if pix.colorspace is not None and pix.colorspace.n not in (1, 3):
        pix = fitz.Pixmap(fitz.csRGB, pix)
    return pix.tobytes("png"), "png"


PAPER_MM = {"A3": (297, 420), "A4": (210, 297), "A5": (148, 210), "A6": (105, 148), "Letter": (215.9, 279.4),
            "Legal": (215.9, 355.6), "Tabloid": (279.4, 431.8)}


def resize_page_content(doc, pno, vw, vh, mode="fit"):
    """Change the VIEW size of page `pno` to vw x vh points; the content is scaled (fit = keep proportions, centred)."""
    pg = doc[pno]
    rot = pg.rotation
    cb = pg.cropbox
    mb = pg.mediabox
    ow, oh = cb.width, cb.height
    nw, nh = (vw, vh) if rot in (0, 180) else (vh, vw)
    sx, sy = nw / ow, nh / oh
    if mode == "fit":
        sx = sy = min(sx, sy)
    tx, ty = (nw - ow * sx) / 2, (nh - oh * sy) / 2
    left = mb.x0 + cb.x0                       # crop box in PDF user space (y up)
    bottom = mb.y0 + (mb.height - cb.y1)
    links = pg.get_links()
    widgets = [(w.xref, fitz.Rect(w.rect)) for w in (pg.widgets() or [])] if rot == 0 else []
    pg.wrap_contents()
    x_open = doc.get_new_xref()
    doc.update_object(x_open, "<< >>")
    doc.update_stream(x_open, ("q %.6f 0 0 %.6f %.4f %.4f cm\n" % (sx, sy, tx - left * sx, ty - bottom * sy)).encode())
    x_close = doc.get_new_xref()
    doc.update_object(x_close, "<< >>")
    doc.update_stream(x_close, b"\nQ\n")
    kind, val = doc.xref_get_key(pg.xref, "Contents")
    if kind == "array":
        body = val.strip()[1:-1]
    elif kind == "xref":
        body = val
    else:
        body = ""
    doc.xref_set_key(pg.xref, "Contents", "[%d 0 R %s %d 0 R]" % (x_open, body, x_close))
    box = "[0 0 %.4f %.4f]" % (nw, nh)
    doc.xref_set_key(pg.xref, "MediaBox", box)
    doc.xref_set_key(pg.xref, "CropBox", box)
    for key in ("TrimBox", "BleedBox", "ArtBox"):
        doc.xref_set_key(pg.xref, key, "null")
    pg = doc[pno]
    # links: move them with the content
    for lk in links:
        try:
            pg.delete_link(lk)
        except Exception:
            pass
    for lk in links:
        r = fitz.Rect(lk["from"])
        nr = fitz.Rect(r.x0 * sx + tx, nh - ty - oh * sy + r.y0 * sy, r.x1 * sx + tx, nh - ty - oh * sy + r.y1 * sy)
        new = {k: v for k, v in lk.items() if k not in ("xref", "id")}
        new["from"] = nr
        try:
            pg.insert_link(new)
        except Exception:
            pass
    for xref, r in widgets:
        for w in pg.widgets() or []:
            if w.xref == xref:
                w.rect = fitz.Rect(r.x0 * sx + tx, nh - ty - oh * sy + r.y0 * sy, r.x1 * sx + tx, nh - ty - oh * sy + r.y1 * sy)
                try:
                    w.update()
                except Exception:
                    pass



# ====================================================================================== font matching for "Edit text"
SYS_FAMILIES = [                      # keyword in the PDF font name -> (regular, bold, italic, bold italic) file names
    ("arial", ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")),
    ("helvetica", ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")),
    ("times", ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf")),
    ("calibri", ("calibri.ttf", "calibrib.ttf", "calibrii.ttf", "calibriz.ttf")),
    ("verdana", ("verdana.ttf", "verdanab.ttf", "verdanai.ttf", "verdanaz.ttf")),
    ("tahoma", ("tahoma.ttf", "tahomabd.ttf", "tahoma.ttf", "tahomabd.ttf")),
    ("segoe", ("segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf", "segoeuiz.ttf")),
    ("georgia", ("georgia.ttf", "georgiab.ttf", "georgiai.ttf", "georgiaz.ttf")),
    ("consolas", ("consola.ttf", "consolab.ttf", "consolai.ttf", "consolaz.ttf")),
    ("courier", ("cour.ttf", "courbd.ttf", "couri.ttf", "courbi.ttf")),
    ("trebuchet", ("trebuc.ttf", "trebucbd.ttf", "trebucit.ttf", "trebucbi.ttf")),
    ("dejavusans", ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf")),
    ("liberationsans", ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf", "LiberationSans-Italic.ttf", "LiberationSans-BoldItalic.ttf")),
    ("carlito", ("Carlito-Regular.ttf", "Carlito-Bold.ttf", "Carlito-Italic.ttf", "Carlito-BoldItalic.ttf")),
]


def _norm(n):
    return re.sub(r"[^a-z0-9]", "", re.sub(r"^[A-Z]{6}\+", "", n or "").lower())


def _covers(font, text):
    try:
        return all(font.has_glyph(ord(c)) for c in text if c not in "\n\r\t ")
    except Exception:
        return False


def match_font(doc, pg, span, text):
    """Font settings for insert_text that look like the original text: the font embedded in the PDF when it has every
    letter that is needed, else an installed font of the same family, else the nearest standard PDF font."""
    name = span.get("font", "") or ""
    flags = span.get("flags", 0)
    plain = re.sub(r"^[A-Z]{6}\+", "", name)
    try:
        for f in pg.get_fonts(full=True):
            if _norm(f[3]) == _norm(plain) or _norm(f[3]).startswith(_norm(plain)) or _norm(plain).startswith(_norm(f[3])):
                if f[1] in ("ttf", "otf", "cff", "pfa", "pfb"):
                    buf = doc.extract_font(f[0])[3]
                    if buf and _covers(fitz.Font(fontbuffer=buf), text):
                        return {"fontname": "CF%d" % f[0], "fontbuffer": buf}
                break
    except Exception:
        pass
    key = re.sub(r"[^a-z]", "", plain.lower())
    bold = bool(flags & 16) or "bold" in key or "black" in key
    ital = bool(flags & 2) or "italic" in key or "oblique" in key
    for kw, files in SYS_FAMILIES:
        if kw in key:
            path = EX._find_in_dirs([files[int(bold) + 2 * int(ital)]])
            if path and _covers(fitz.Font(fontfile=path), text):
                return {"fontname": "SF" + re.sub(r"[^A-Za-z0-9]", "", os.path.splitext(os.path.basename(path))[0]), "fontfile": path}
            break
    if EX.needs_unicode(text):
        return EX.font_kwargs(text, bold)
    return {"fontname": EX.base14_name(name, flags)}


# ====================================================================================== signature blocks
def blocks_dir():
    return os.path.join(CONFIG_DIR, "blocks")


def list_blocks():
    out = []
    try:
        for f in sorted(os.listdir(blocks_dir())):
            if f.endswith(".json"):
                with open(os.path.join(blocks_dir(), f), "r", encoding="utf-8") as fh:
                    d = json.load(fh)
                d["_id"] = f[:-5]
                out.append(d)
    except Exception:
        pass
    return out


DATE_FORMATS = ["%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%d %B %Y", "%B %d, %Y"]


def render_block(defn, sig_path=None):
    """The picture of a signature block: signature, line, name, position, today's date"""
    sig = None
    p = sig_path
    if p and os.path.isfile(p):
        sig = Image.open(p).convert("RGBA")
        if sig.height > 150:
            sig = sig.resize((max(1, int(sig.width * 150 / sig.height)), 150), Image.LANCZOS)
    f_b = load_pil_font(find_font("bold"), 30)
    f_r = load_pil_font(find_font("regular"), 26)
    lines = []
    if defn.get("who"):
        lines.append((defn["who"], f_b))
    if defn.get("role"):
        lines.append((defn["role"], f_r))
    if defn.get("date"):
        fmt = defn.get("fmt") if defn.get("fmt") in DATE_FORMATS else DATE_FORMATS[0]
        lines.append((datetime.date.today().strftime(fmt), f_r))
    probe = ImageDraw.Draw(Image.new("RGBA", (10, 10)))
    widths = [int(probe.textlength(t, font=f)) for t, f in lines]
    w = max([sig.width if sig else 0] + widths + [260]) + 30
    h = 12 + (sig.height + 6 if sig else 40) + (6 if defn.get("line", True) else 0) + 36 * len(lines) + 12
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    y = 12
    if sig:
        img.alpha_composite(sig, (15, y))
        y += sig.height + 6
    else:
        y += 40
    d = ImageDraw.Draw(img)
    if defn.get("line", True):
        d.line((15, y, w - 15, y), fill=(60, 60, 60, 255), width=2)
        y += 6
    for t, f in lines:
        d.text((15, y), t, font=f, fill=(30, 30, 30, 255))
        y += 36
    return img


# ====================================================================================== the mixin
class V3Tools:
    # ------------------------------------------------------------------ start-up / state
    def _v3_init(self):
        self.cfg = load_cfg()
        try:                                              # highlight colours chosen by the user
            cols = self.cfg.get("hl_colors")
            if cols and len(cols) == 4:
                EX.HL_COLORS[:] = [tuple(max(0.0, min(1.0, float(x))) for x in c[:3]) for c in cols]
        except Exception:
            pass
        try:
            import claudius_pro
            claudius_pro.RECOVERY_KEEP_DAYS = self.cfg["recovery_days"]
        except Exception:
            pass
        self.hl_buttons = []
        self.hl_color = EX.HL_COLORS[0]

    def _v3_after_ui(self):
        c = self.cfg
        if c.get("remember_view", True):
            if c.get("sidebar_mode") in ("pages", "bookmarks", "comments"):
                self.set_sidebar(c["sidebar_mode"])
            if not c.get("panel", True):
                self.toggle_sidebar(False)
        self.bind("<Control-comma>", lambda e: self.dlg_settings())

    def _v3_save_state(self):
        try:
            c = self.cfg
            c["panel"] = self.side_frame.winfo_manager() == "pack"
            c["sidebar_mode"] = self.sidebar_mode
            if self.tabs:
                self._tab_stash()
                for st in self.tabs:
                    self._v3_remember(st)
            elif self.doc is not None:
                self._v3_remember({"path": self.path, "page_no": self.page_no, "zoom": self.zoom, "fit_mode": self.fit_mode,
                                   "doc": self.doc})
            save_cfg(c)
        except Exception:
            pass

    def _v3_remember(self, st):
        """note the reading position of one document"""
        try:
            p = st.get("path")
            doc = st.get("doc")
            if not p or doc is None or not self.cfg.get("remember_position", True):
                return
            pos = self.cfg["positions"]
            pos[norm_path(p)] = {"page": int(st.get("page_no") or 0), "zoom": round(float(st.get("zoom") or 1.0), 3),
                                 "fit": st.get("fit_mode") or "", "t": int(time.time())}
            if len(pos) > MAX_POSITIONS:
                for k in sorted(pos, key=lambda k: pos[k].get("t", 0))[:len(pos) - MAX_POSITIONS]:
                    del pos[k]
        except Exception:
            pass

    def _v3_restore_position(self, path):
        """called after a file was opened from disk"""
        try:
            if not self.cfg.get("remember_position", True):
                return
            pos = self.cfg["positions"].get(norm_path(path))
            if not pos or self.doc is None:
                return
            if pos.get("fit") == "width":
                self.fit_width(render=False)
            elif pos.get("fit") == "page":
                self.fit_page(render=False)
            else:
                self.set_zoom(float(pos.get("zoom", 1.0)), render=False)
            n = max(0, min(int(pos.get("page", 0)), len(self.doc) - 1))
            self.page_no = n
            self.sel_pages = {n}
            self.render()
            if n:
                self.goto(n)
            self.update_ui()
        except Exception:
            pass

    def close_tab(self, i=None):
        try:
            k = self.cur_tab if i is None else i
            if self.tabs and 0 <= k < len(self.tabs):
                self._tab_stash()
                self._v3_remember(self.tabs[k])
        except Exception:
            pass
        return super().close_tab(i)

    def _adopt_doc(self, doc, path=None, modified=False):
        super()._adopt_doc(doc, path, modified)
        try:
            dv = self.cfg.get("default_view", "page")
            if dv == "width":
                self.fit_width()
            elif dv == "100":
                self.set_zoom(1.0)
        except Exception:
            pass

    def _autosave_ms(self):
        return int(max(1, self.cfg.get("autosave_minutes", 1)) * 60000)

    # ------------------------------------------------------------------ settings window
    def dlg_settings(self):
        import claudius_i18n as i18n
        c = self.cfg
        d = self._dlg("Settings")
        v = {k: tk.BooleanVar(value=bool(c[k])) for k in ("remember_window", "remember_view", "remember_position", "autosave")}
        r = 0
        self._label(d, "Language", bold=True).grid(row=r, column=0, sticky="w", pady=(0, 3))
        lang = tk.StringVar(value=dict(i18n.LANGS).get(i18n.current, i18n.current))
        ttk.Combobox(d, textvariable=lang, values=[n for _c, n in i18n.LANGS], state="readonly", width=22).grid(row=r, column=1, sticky="w")
        r += 1
        self._label(d, "Appearance", bold=True).grid(row=r, column=0, sticky="w", pady=3)
        themes = (("auto", "System (follow Windows)"), ("light", "Light"), ("dark", "Dark"))
        cur_theme = TH.load_setting()
        theme = tk.StringVar(value=_tr(dict(themes).get(cur_theme, "System (follow Windows)")))
        ttk.Combobox(d, textvariable=theme, values=[_tr(n) for _k, n in themes], state="readonly", width=22).grid(row=r, column=1, sticky="w")
        r += 1
        self._label(d, "Start-up", bold=True).grid(row=r, column=0, sticky="w", pady=(10, 3))
        for key, text in (("remember_window", "Remember the window size and position"),
                          ("remember_view", "Remember the side panel and its tab"),
                          ("remember_position", "Reopen documents at the page where I stopped")):
            tk.Checkbutton(d, text=text, variable=v[key], bg=BG, activebackground=BG, fg=TH.LABEL, selectcolor=TH.CARD,
                           font=(UI_FONT, 9)).grid(row=r, column=1, columnspan=2, sticky="w")
            r += 1
        self._label(d, "View of a new document").grid(row=r, column=0, sticky="w", pady=3)
        views = (("page", "Fit page"), ("width", "Fit width"), ("100", "100 %"))
        dv = tk.StringVar(value=_tr(dict(views).get(c["default_view"], "Fit page")))
        ttk.Combobox(d, textvariable=dv, values=[_tr(n) for _k, n in views], state="readonly", width=22).grid(row=r, column=1, sticky="w")
        r += 1
        self._label(d, "Safety copy", bold=True).grid(row=r, column=0, sticky="w", pady=(10, 3))
        tk.Checkbutton(d, text="Save unsaved work automatically", variable=v["autosave"], bg=BG, activebackground=BG, fg=TH.LABEL,
                       selectcolor=TH.CARD, font=(UI_FONT, 9)).grid(row=r, column=1, columnspan=2, sticky="w")
        r += 1
        mins, days = tk.IntVar(value=c["autosave_minutes"]), tk.IntVar(value=c["recovery_days"])
        self._label(d, "Every (minutes)").grid(row=r, column=0, sticky="w", pady=2)
        ttk.Spinbox(d, from_=1, to=30, width=6, textvariable=mins).grid(row=r, column=1, sticky="w")
        r += 1
        self._label(d, "Keep safety copies (days)").grid(row=r, column=0, sticky="w", pady=2)
        ttk.Spinbox(d, from_=1, to=90, width=6, textvariable=days).grid(row=r, column=1, sticky="w")
        r += 1
        self._label(d, "Printing", bold=True).grid(row=r, column=0, sticky="w", pady=(10, 3))
        printer = tk.StringVar(value=c["printer"])
        copies = tk.IntVar(value=c["copies"])
        try:
            import claudius_print as cp
            names, _default = cp.list_printers()
        except Exception:
            names, default = [], ""
        self._label(d, "Default printer").grid(row=r, column=0, sticky="w", pady=2)
        ttk.Combobox(d, textvariable=printer, values=[""] + list(names), width=34).grid(row=r, column=1, columnspan=2, sticky="w")
        r += 1
        self._label(d, "Copies").grid(row=r, column=0, sticky="w", pady=2)
        ttk.Spinbox(d, from_=1, to=99, width=6, textvariable=copies).grid(row=r, column=1, sticky="w")
        r += 1
        self._label(d, "Highlight colours", bold=True).grid(row=r, column=0, sticky="w", pady=(10, 3))
        ttk.Button(d, text="Change names and colours...", command=self.dlg_highlight_colors).grid(row=r, column=1, sticky="w", pady=(10, 3))
        r += 1

        def go():
            for k in v:
                c[k] = bool(v[k].get())
            c["default_view"] = {_tr(n): k for k, n in views}.get(dv.get(), "page")
            try:
                c["autosave_minutes"] = max(1, min(30, int(mins.get())))
                c["recovery_days"] = max(1, min(90, int(days.get())))
                c["copies"] = max(1, min(99, int(copies.get())))
            except (tk.TclError, ValueError):
                pass
            c["printer"] = printer.get().strip()
            try:
                import claudius_pro
                claudius_pro.RECOVERY_KEEP_DAYS = c["recovery_days"]
            except Exception:
                pass
            save_cfg(c)
            new_lang = {n: k for k, n in i18n.LANGS}.get(lang.get(), i18n.current)
            new_theme = {_tr(n): k for k, n in themes}.get(theme.get(), cur_theme)
            d.destroy()
            self.status.config(text=_tr("Settings saved."))
            if new_lang != i18n.current:
                self.lang_var.set(new_lang)
                self.set_ui_language(new_lang)
            elif new_theme != cur_theme:
                self.theme_var.set(new_theme)
                self.set_theme(new_theme)
        self._buttons(d, "Save", go, r, colspan=3)

    # ------------------------------------------------------------------ undo history
    def dlg_undo_history(self):
        if not self.need_doc():
            return
        self.commit_floats()
        stack = list(self.undo_stack)
        if not stack:
            messagebox.showinfo(APP, "There is nothing to undo yet.")
            return
        d = self._dlg("Undo history")
        self._label(d, "Newest change first. Choose a step to go back to the state before it.", fg=TH.TERTIARY).grid(row=0, column=0, sticky="w")
        tree = ttk.Treeview(d, columns=("what", "page"), show="headings", height=12, selectmode="browse")
        tree.heading("what", text=_tr("Change"))
        tree.heading("page", text=_tr("Page"))
        tree.column("what", width=S(340))
        tree.column("page", width=S(60), anchor="center")
        tree.grid(row=1, column=0, pady=6)
        for k in range(len(stack) - 1, -1, -1):
            snap = stack[k]
            label = getattr(snap, "label", None) or _tr("Change")
            tree.insert("", "end", iid=str(k), values=(label, int(snap[1]) + 1))
        first = tree.get_children()
        if first:
            tree.selection_set(first[0])

        def go():
            sel = tree.selection()
            if not sel:
                return
            steps = len(self.undo_stack) - int(sel[0])
            d.destroy()
            for _ in range(max(0, steps)):
                if not self.undo_stack:
                    break
                self.undo()
            self.status.config(text=_tr("Went back {} step(s).").replace("{}", str(steps)))
        self._buttons(d, "Go back to here", go, 2, colspan=1)

    # ------------------------------------------------------------------ presets (header / footer / watermark / numbers)
    def _preset_bar(self, d, row, kind, getter, setter, colspan=4):
        f = tk.Frame(d, bg=BG)
        f.grid(row=row, column=0, columnspan=colspan, sticky="w", pady=(10, 0))
        self._label(f, "Preset").pack(side="left")
        var = tk.StringVar()
        cb = ttk.Combobox(f, textvariable=var, values=sorted(self.cfg["presets"].get(kind, {})), width=20)
        cb.pack(side="left", padx=4)

        def refresh():
            cb["values"] = sorted(self.cfg["presets"].get(kind, {}))

        def load():
            p = self.cfg["presets"].get(kind, {}).get(var.get())
            if p is None:
                messagebox.showinfo(APP, "Choose a saved preset first.", parent=d)
                return
            try:
                setter(p)
            except Exception as ex:
                messagebox.showerror(APP, f"{ex}", parent=d)

        def save():
            name = var.get().strip()[:40]
            if not name:
                messagebox.showinfo(APP, "Type a name for the preset first.", parent=d)
                return
            self.cfg["presets"].setdefault(kind, {})[name] = getter()
            save_cfg(self.cfg)
            refresh()
            self.status.config(text=_tr("Preset saved."))

        def delete():
            if self.cfg["presets"].get(kind, {}).pop(var.get(), None) is not None:
                save_cfg(self.cfg)
                var.set("")
                refresh()
        for t, c in (("Load", load), ("Save", save), ("Delete", delete)):
            ttk.Button(f, text=t, command=c, width=7).pack(side="left", padx=2)

    # ------------------------------------------------------------------ crop and resize pages
    def dlg_page_crop(self):
        if not self.need_doc():
            return
        self.commit_floats()
        pg = self.page()
        r = pg.rect
        side = S(420)
        sc = min(side / r.width, side / r.height)
        pix = pg.get_pixmap(matrix=fitz.Matrix(sc, sc), alpha=False)
        pw, ph = pix.width, pix.height
        d = self._dlg("Crop and resize pages")
        photo = pix_to_photo(pix, d)
        cv = tk.Canvas(d, width=pw, height=ph, highlightthickness=1, highlightbackground=TH.LINE, cursor="crosshair")
        cv.grid(row=0, column=0, rowspan=12, padx=(0, 16))
        cv.create_image(0, 0, anchor="nw", image=photo)
        cv._photo = photo
        box = [0.0, 0.0, float(pw), float(ph)]
        mm = 72 / 25.4 * sc                                     # canvas pixels per mm
        margins = {k: tk.StringVar(value="0.0") for k in ("top", "bottom", "left", "right")}
        shade = [cv.create_rectangle(0, 0, 0, 0, fill="black", stipple="gray50", outline="") for _ in range(4)]
        frame = cv.create_rectangle(0, 0, 0, 0, outline=ACCENT, width=2)
        hs = S(5)
        handles = [cv.create_rectangle(0, 0, 0, 0, fill="white", outline=ACCENT, width=2) for _ in range(8)]
        lock = {"on": False}

        def redraw():
            x0, y0, x1, y1 = box
            cv.coords(shade[0], 0, 0, pw, y0)
            cv.coords(shade[1], 0, y1, pw, ph)
            cv.coords(shade[2], 0, y0, x0, y1)
            cv.coords(shade[3], x1, y0, pw, y1)
            cv.coords(frame, x0, y0, x1, y1)
            xm, ym = (x0 + x1) / 2, (y0 + y1) / 2
            pts = [(x0, y0), (xm, y0), (x1, y0), (x1, ym), (x1, y1), (xm, y1), (x0, y1), (x0, ym)]
            for h, (hx, hy) in zip(handles, pts):
                cv.coords(h, hx - hs, hy - hs, hx + hs, hy + hs)
            lock["on"] = True
            margins["top"].set("%.1f" % (y0 / mm))
            margins["bottom"].set("%.1f" % ((ph - y1) / mm))
            margins["left"].set("%.1f" % (x0 / mm))
            margins["right"].set("%.1f" % ((pw - x1) / mm))
            lock["on"] = False

        def from_fields(_e=None):
            try:
                t, b, l, rr = (float(margins[k].get().replace(",", ".")) * mm for k in ("top", "bottom", "left", "right"))
            except ValueError:
                return
            x0, y0, x1, y1 = l, t, pw - rr, ph - b
            if x1 - x0 >= 10 and y1 - y0 >= 10 and x0 >= 0 and y0 >= 0:
                box[:] = [x0, y0, x1, y1]
                redraw()
        drag = {}

        def press(e):
            x0, y0, x1, y1 = box
            xm, ym = (x0 + x1) / 2, (y0 + y1) / 2
            pts = {"nw": (x0, y0), "n": (xm, y0), "ne": (x1, y0), "e": (x1, ym), "se": (x1, y1), "s": (xm, y1), "sw": (x0, y1), "w": (x0, ym)}
            for name, (hx, hy) in pts.items():
                if abs(e.x - hx) <= hs + 4 and abs(e.y - hy) <= hs + 4:
                    drag.update(mode=name, last=(e.x, e.y))
                    return
            if x0 < e.x < x1 and y0 < e.y < y1:
                drag.update(mode="move", last=(e.x, e.y))
            else:
                drag.update(mode="new", start=(e.x, e.y), last=(e.x, e.y))

        def move(e):
            m = drag.get("mode")
            if not m:
                return
            x, y = max(0, min(pw, e.x)), max(0, min(ph, e.y))
            x0, y0, x1, y1 = box
            if m == "move":
                dx, dy = e.x - drag["last"][0], e.y - drag["last"][1]
                dx = max(-x0, min(pw - x1, dx))
                dy = max(-y0, min(ph - y1, dy))
                box[:] = [x0 + dx, y0 + dy, x1 + dx, y1 + dy]
                drag["last"] = (e.x, e.y)
            elif m == "new":
                sx, sy = drag["start"]
                box[:] = [min(sx, x), min(sy, y), max(sx, x), max(sy, y)]
            else:
                if "w" in m:
                    x0 = min(x, x1 - 10)
                if "e" in m:
                    x1 = max(x, x0 + 10)
                if "n" in m:
                    y0 = min(y, y1 - 10)
                if "s" in m:
                    y1 = max(y, y0 + 10)
                box[:] = [x0, y0, x1, y1]
            redraw()
        cv.bind("<ButtonPress-1>", press)
        cv.bind("<B1-Motion>", move)
        cv.bind("<ButtonRelease-1>", lambda e: drag.clear())

        col = 1
        self._label(d, "Crop - drag the frame or type the margins (mm):", bold=True).grid(row=0, column=col, columnspan=4, sticky="w")
        for i, (k, t) in enumerate((("top", "Top"), ("bottom", "Bottom"), ("left", "Left"), ("right", "Right"))):
            self._label(d, t).grid(row=1 + i // 2, column=col + (i % 2) * 2, sticky="w", pady=2, padx=(0, 4))
            e = ttk.Entry(d, textvariable=margins[k], width=7)
            e.grid(row=1 + i // 2, column=col + (i % 2) * 2 + 1, sticky="w", padx=(0, 10))
            e.bind("<Return>", from_fields)
            e.bind("<FocusOut>", from_fields)
        self._label(d, "Apply to", bold=True).grid(row=3, column=col, sticky="w", pady=(10, 0))
        scope, rng = tk.StringVar(value="cur"), tk.StringVar(value="")
        for i, (t, val) in enumerate((("Current page only", "cur"), ("All pages", "all"), ("Selected pages", "sel"), ("Page range:", "rng"))):
            tk.Radiobutton(d, text=t, variable=scope, value=val, bg=BG, activebackground=BG, fg=TH.LABEL, selectcolor=TH.CARD,
                           font=(UI_FONT, 9)).grid(row=4 + i, column=col, columnspan=2, sticky="w")
        ttk.Entry(d, textvariable=rng, width=14).grid(row=7, column=col + 2, columnspan=2, sticky="w")
        self._label(d, "Page size", bold=True).grid(row=8, column=col, sticky="w", pady=(10, 0))
        size = tk.StringVar(value="A4")
        ttk.Combobox(d, textvariable=size, values=list(PAPER_MM) + ["Custom (mm)"], state="readonly", width=14).grid(
            row=9, column=col, columnspan=2, sticky="w")
        cw, ch_ = tk.StringVar(value="210"), tk.StringVar(value="297")
        ttk.Entry(d, textvariable=cw, width=6).grid(row=9, column=col + 2, sticky="w")
        ttk.Entry(d, textvariable=ch_, width=6).grid(row=9, column=col + 3, sticky="w")
        how = tk.StringVar(value="Fit (keep proportions)")
        ttk.Combobox(d, textvariable=how, values=["Fit (keep proportions)", "Stretch to the new size"], state="readonly",
                     width=24).grid(row=10, column=col, columnspan=4, sticky="w", pady=3)
        self._label(d, "Comments on resized pages are merged into the page.", fg=TH.TERTIARY).grid(row=11, column=col, columnspan=4, sticky="w")
        redraw()

        def pages_now():
            sc_ = scope.get()
            if sc_ == "all":
                return list(range(len(self.doc)))
            if sc_ == "cur":
                return [self.page_no]
            if sc_ == "sel":
                return self.selected_pages()
            return self._pages_or_warn(rng.get(), d)

        def do_crop():
            pages = pages_now()
            if pages is None:
                return
            x0, y0, x1, y1 = box
            top, bottom, left, right = y0 / sc, (ph - y1) / sc, x0 / sc, (pw - x1) / sc
            d.destroy()
            self.do_crop(pages, top, bottom, left, right)

        def do_resize():
            pages = pages_now()
            if pages is None:
                return
            try:
                if size.get() in PAPER_MM:
                    wmm, hmm = PAPER_MM[size.get()]
                else:
                    wmm, hmm = float(cw.get().replace(",", ".")), float(ch_.get().replace(",", "."))
            except ValueError:
                messagebox.showinfo(APP, "Enter the page size in millimetres.", parent=d)
                return
            if not (20 <= wmm <= 5000 and 20 <= hmm <= 5000):
                messagebox.showinfo(APP, "The page size must be between 20 and 5000 mm.", parent=d)
                return
            d.destroy()
            self.do_resize_pages(pages, wmm * 72 / 25.4, hmm * 72 / 25.4, "fit" if how.get().startswith("Fit") else "stretch")
        bf = tk.Frame(d, bg=BG)
        bf.grid(row=12, column=0, columnspan=5, sticky="e", pady=(12, 0))
        ttk.Button(bf, text="Crop", command=do_crop).pack(side="left", padx=4)
        ttk.Button(bf, text="Resize pages", command=do_resize).pack(side="left", padx=4)
        ttk.Button(bf, text="Cancel", command=d.destroy).pack(side="left")

    def do_resize_pages(self, pages, vw, vh, mode="fit"):
        """Give the pages a new size; orientation follows the current page (landscape stays landscape)."""
        def work():
            if any(len(list(self.doc[i].annots() or [])) for i in pages):
                try:
                    self.doc.bake(annots=True, widgets=False)
                except Exception:
                    pass
            for i in pages:
                v = self.doc[i].rect
                w, h = (max(vw, vh), min(vw, vh)) if v.width > v.height else (min(vw, vh), max(vw, vh))
                resize_page_content(self.doc, i, w, h, mode)
        return self._run_edit(work, rebuild=True, msg=_tr("Resized {} page(s).").replace("{}", str(len(pages))))

    # ------------------------------------------------------------------ edit images already in the PDF
    def _v3_image_at(self, pt):
        pg = self.page()
        hit = None
        for info in pg.get_image_info(xrefs=True):
            if fitz.Rect(info["bbox"]).contains(pt) and info.get("xref"):
                hit = info
        return hit

    def _v3_release_image(self, x0, y0, x1, y1):
        p0 = self._canvas_pt(x0, y0)
        info = self._v3_image_at(p0)
        if not info:
            self.status.config(text=_tr("There is no picture at that spot."))
            return
        bbox = fitz.Rect(info["bbox"])
        xref = info["xref"]
        if abs(x1 - x0) + abs(y1 - y0) > 8:                       # dragged: move the picture
            p1 = self._canvas_pt(x1, y1)
            self.v3_move_image(bbox, xref, p1.x - p0.x, p1.y - p0.y)
            return
        m = tk.Menu(self, tearoff=0)
        self.v3_picture_items(m, bbox, xref)
        m.add_separator()
        m.add_command(label="Tip: drag a picture to move it", state="disabled")
        m.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

    def v3_picture_items(self, m, bbox, xref):
        """Menu entries for a picture that is already part of the page (also used by the right-click menu)."""
        m.add_command(label="Delete this picture", command=lambda: self.v3_delete_image(bbox))
        m.add_command(label="Replace with another picture...", command=lambda: self.v3_replace_image(bbox, xref))
        m.add_command(label="Save this picture as...", command=lambda: self.v3_save_image(xref))
        m.add_separator()
        self.align_menu(m, lambda mode: self.v3_align_image(bbox, xref, mode))

    # ------------------------------------------------------------------ align to the page (like PowerPoint)
    ALIGN_MODES = (("center", "Center on page"), ("h", "Center horizontally"), ("v", "Center vertically"),
                   ("left", "Align left"), ("right", "Align right"), ("top", "Align top"), ("bottom", "Align bottom"))

    @staticmethod
    def align_delta(box, page, mode):
        """How far a box must move (dx, dy) to line up with the page. box/page = (x0, y0, x1, y1)."""
        bx0, by0, bx1, by1 = box
        px0, py0, px1, py1 = page
        dx = dy = 0.0
        if mode in ("center", "h"):
            dx = (px0 + px1) / 2 - (bx0 + bx1) / 2
        if mode in ("center", "v"):
            dy = (py0 + py1) / 2 - (by0 + by1) / 2
        if mode == "left":
            dx = px0 - bx0
        elif mode == "right":
            dx = px1 - bx1
        elif mode == "top":
            dy = py0 - by0
        elif mode == "bottom":
            dy = py1 - by1
        return dx, dy

    def align_menu(self, m, call):
        """Add an 'Align to page' sub-menu to menu m; call(mode) does the work."""
        am = tk.Menu(m, tearoff=0)
        for i, (mode, label) in enumerate(self.ALIGN_MODES):
            if i in (3,):
                am.add_separator()
            am.add_command(label=label, command=lambda mode=mode: call(mode))
        m.add_cascade(label="Align to page", menu=am)

    def img_align(self, mode):
        """Align the selected, not yet applied picture to the page."""
        if not self.need_sel():
            return
        o = self.sel
        r = self.page().rect
        dx, dy = self.align_delta(self.f_bbox(o), (0, 0, r.width, r.height), mode)
        o.cx += dx
        o.cy += dy
        self._touch()
        self.status.config(text=_tr("Picture aligned to the page."))

    def img_align_menu(self):
        if not self.need_sel():
            return
        m = tk.Menu(self, tearoff=0)
        for i, (mode, label) in enumerate(self.ALIGN_MODES):
            if i == 3:
                m.add_separator()
            m.add_command(label=label, command=lambda mode=mode: self.img_align(mode))
        m.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

    def v3_align_image(self, bbox, xref, mode):
        """Align a picture that is already part of the PDF."""
        pg = self.page()
        pr = fitz.Rect(pg.rect * pg.derotation_matrix).normalize()      # image boxes are in unrotated page space
        dx, dy = self.align_delta(tuple(bbox), tuple(pr), mode)
        if abs(dx) < 0.01 and abs(dy) < 0.01:
            self.status.config(text=_tr("The picture is already there."))
            return
        self.v3_move_image(bbox, xref, dx, dy)

    def _v3_remove_image(self, pg, bbox):
        inner = fitz.Rect(bbox)
        inner.x0 += min(2, bbox.width / 4)
        inner.y0 += min(2, bbox.height / 4)
        inner.x1 -= min(2, bbox.width / 4)
        inner.y1 -= min(2, bbox.height / 4)
        pg.add_redact_annot(inner, fill=False)
        pg.apply_redactions(images=fitz.PDF_REDACT_IMAGE_REMOVE, graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                            text=fitz.PDF_REDACT_TEXT_NONE)

    def v3_delete_image(self, bbox):
        pg = self.page()
        self._run_edit(lambda: self._v3_remove_image(pg, bbox), msg=_tr("Picture deleted."))

    def v3_move_image(self, bbox, xref, dx, dy):
        pg = self.page()
        data, _ext = image_png_bytes(self.doc, xref)
        new = fitz.Rect(bbox.x0 + dx, bbox.y0 + dy, bbox.x1 + dx, bbox.y1 + dy)

        def work():
            self._v3_remove_image(pg, bbox)
            pg.insert_image(new, stream=data, keep_proportion=False, overlay=True)
        self._run_edit(work, msg=_tr("Picture moved."))

    def v3_replace_image(self, bbox, xref):
        p = filedialog.askopenfilename(parent=self, filetypes=[("Images", "*.png *.jpg *.jpeg *.bmp *.gif *.webp *.tif *.tiff"), ("All", "*.*")])
        if not p:
            return
        try:
            im = Image.open(p)
            im.load()
            buf = io.BytesIO()
            im.convert("RGBA" if "A" in im.getbands() else "RGB").save(buf, "PNG")
            data = buf.getvalue()
            ratio = im.width / im.height
        except Exception as ex:
            messagebox.showerror(APP, f"Could not open the picture:\n{ex}")
            return
        pg = self.page()
        r = fitz.Rect(bbox)
        if r.width / r.height > ratio:                              # keep the new picture's proportions inside the old frame
            w = r.height * ratio
            r = fitz.Rect(r.x0 + (r.width - w) / 2, r.y0, r.x0 + (r.width + w) / 2, r.y1)
        else:
            h = r.width / ratio
            r = fitz.Rect(r.x0, r.y0 + (r.height - h) / 2, r.x1, r.y0 + (r.height + h) / 2)

        def work():
            self._v3_remove_image(pg, bbox)
            pg.insert_image(r, stream=data, keep_proportion=False, overlay=True)
        self._run_edit(work, msg=_tr("Picture replaced."))

    def v3_save_image(self, xref):
        try:
            data, ext = image_png_bytes(self.doc, xref)
        except Exception as ex:
            messagebox.showerror(APP, f"{ex}")
            return
        p = filedialog.asksaveasfilename(defaultextension="." + ("jpg" if ext == "jpeg" else ext),
                                         filetypes=[("Image", "*.png *.jpg")], initialfile="picture." + ("jpg" if ext == "jpeg" else ext))
        if p:
            with open(p, "wb") as fh:
                fh.write(data)
            self.status.config(text=_tr("Picture saved."))

    # ------------------------------------------------------------------ signature blocks
    def _v3_block_menu(self, m):
        blocks = list_blocks()
        if blocks:
            m.add_separator()
            bm = tk.Menu(m, tearoff=0)
            for b in blocks:
                bm.add_command(label=b.get("name") or b["_id"], command=lambda b=b: self.v3_place_block(b))
            m.add_cascade(label="Signature blocks", menu=bm)
        m.add_command(label="Create a signature block...", command=self.dlg_signature_block)

    def v3_place_block(self, defn):
        try:
            sig = os.path.join(blocks_dir(), defn["_id"] + ".png")
            self.place_image(render_block(defn, sig if os.path.exists(sig) else None), "signature block")
            self.status.config(text=_tr("Drag the signature block into place, then press Enter (or Apply) to fix it on the page."))
        except Exception as ex:
            messagebox.showerror(APP, f"{ex}")

    def dlg_signature_block(self):
        d = self._dlg("Signature block")
        sigs = self.saved_signatures("sign")
        self._label(d, "A signature block is a signature with your name, position and today's date underneath.", fg=TH.TERTIARY).grid(
            row=0, column=0, columnspan=3, sticky="w")
        self._label(d, "Signature").grid(row=1, column=0, sticky="nw", pady=3)
        lb = tk.Listbox(d, height=min(6, max(2, len(sigs))), width=34, exportselection=False, activestyle="none")
        lb.grid(row=1, column=1, sticky="w", pady=3)
        for i, p in enumerate(sigs):
            lb.insert("end", f"{_tr('Signature')} {i + 1}")
        if sigs:
            lb.selection_set(0)
        else:
            lb.insert("end", _tr("(no saved signature - create one first)"))
        ttk.Button(d, text="Create new...", command=lambda: (d.destroy(), self.dlg_new_signature("sign"))).grid(row=1, column=2, padx=6, sticky="n", pady=3)
        name, who, role = tk.StringVar(value=_tr("My signature block")), tk.StringVar(), tk.StringVar()
        for i, (t, var) in enumerate((("Name of this block", name), ("Printed name", who), ("Position / title (optional)", role))):
            self._label(d, t).grid(row=2 + i, column=0, sticky="w", pady=3)
            ttk.Entry(d, textvariable=var, width=36).grid(row=2 + i, column=1, columnspan=2, sticky="w", pady=3)
        usedate, line = tk.BooleanVar(value=True), tk.BooleanVar(value=True)
        fmt = tk.StringVar(value=DATE_FORMATS[0])
        tk.Checkbutton(d, text="Add today's date", variable=usedate, bg=BG, activebackground=BG, fg=TH.LABEL, selectcolor=TH.CARD,
                       font=(UI_FONT, 9)).grid(row=5, column=0, sticky="w")
        examples = [datetime.date.today().strftime(f) for f in DATE_FORMATS]
        ttk.Combobox(d, textvariable=fmt, values=examples, width=34, state="readonly").grid(row=5, column=1, columnspan=2, sticky="w")
        fmt.set(examples[0])
        tk.Checkbutton(d, text="Line above the text", variable=line, bg=BG, activebackground=BG, fg=TH.LABEL, selectcolor=TH.CARD,
                       font=(UI_FONT, 9)).grid(row=6, column=0, columnspan=2, sticky="w")
        prev = tk.Label(d, bg="white", relief="solid", bd=1)
        prev.grid(row=7, column=0, columnspan=3, pady=8)

        def defn():
            return {"name": name.get().strip()[:60] or _tr("My signature block"), "who": who.get().strip()[:80],
                    "role": role.get().strip()[:80], "date": bool(usedate.get()), "line": bool(line.get()),
                    "fmt": DATE_FORMATS[examples.index(fmt.get())] if fmt.get() in examples else DATE_FORMATS[0]}

        def sig_path():
            s = lb.curselection()
            return sigs[s[0]] if sigs and s else None

        def preview(_e=None):
            try:
                im = render_block(defn(), sig_path())
                im.thumbnail((S(420), S(170)))
                bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
                bg.alpha_composite(im)
                prev._ph = ImageTk.PhotoImage(bg, master=d)
                prev.config(image=prev._ph)
            except Exception:
                pass
        for w in (lb,):
            w.bind("<<ListboxSelect>>", preview)
        for var in (who, role, usedate, line, fmt):
            var.trace_add("write", lambda *a: preview())
        preview()

        def save(place=False):
            dd = defn()
            if not dd["who"] and not sig_path():
                messagebox.showinfo(APP, "Choose a signature or type a name.", parent=d)
                return
            bid = "b%d" % int(time.time() * 1000)
            os.makedirs(blocks_dir(), exist_ok=True)
            sp = sig_path()
            if sp:
                shutil.copyfile(sp, os.path.join(blocks_dir(), bid + ".png"))
            with open(os.path.join(blocks_dir(), bid + ".json"), "w", encoding="utf-8") as fh:
                json.dump(dd, fh)
            d.destroy()
            self.status.config(text=_tr("Signature block saved - find it under Sign > Signature."))
            if place and self.doc is not None:
                dd["_id"] = bid
                self.v3_place_block(dd)
        f = tk.Frame(d, bg=BG)
        f.grid(row=8, column=0, columnspan=3, sticky="e", pady=(6, 0))
        ttk.Button(f, text="Save", command=save).pack(side="left", padx=4)
        ttk.Button(f, text="Save and place", command=lambda: save(True)).pack(side="left", padx=4)
        ttk.Button(f, text="Delete a block...", command=lambda: self._v3_delete_block(d)).pack(side="left", padx=4)
        ttk.Button(f, text="Cancel", command=d.destroy).pack(side="left")

    def _v3_delete_block(self, parent):
        blocks = list_blocks()
        if not blocks:
            messagebox.showinfo(APP, "There are no saved signature blocks.", parent=parent)
            return
        d = self._dlg("Delete a block")
        lb = tk.Listbox(d, height=min(8, len(blocks)), width=40, activestyle="none")
        lb.pack()
        for b in blocks:
            lb.insert("end", b.get("name") or b["_id"])

        def go():
            s = lb.curselection()
            if not s:
                return
            bid = blocks[s[0]]["_id"]
            for ext in (".json", ".png"):
                try:
                    os.remove(os.path.join(blocks_dir(), bid + ext))
                except OSError:
                    pass
            d.destroy()
        self._buttons(d, "Delete", go, 1, colspan=1)

    # ------------------------------------------------------------------ highlight colours and export
    def _v3_hl_label(self, rgb):
        """name of the highlight colour closest to rgb"""
        best, bd = None, 9.0
        for i, c in enumerate(EX.HL_COLORS):
            dist = sum((a - b) ** 2 for a, b in zip(c[:3], rgb[:3]))
            if dist < bd:
                best, bd = i, dist
        if best is None or bd > 0.12:
            return _tr("Other colour")
        return self.cfg["hl_names"][best] or ("#" + "%02x%02x%02x" % tuple(int(x * 255) for x in EX.HL_COLORS[best][:3]))

    def dlg_highlight_colors(self):
        d = self._dlg("Highlight colours")
        self._label(d, "Give each highlight colour a meaning. The export groups the highlights by these names.", fg=TH.TERTIARY).grid(
            row=0, column=0, columnspan=3, sticky="w")
        rows = []
        for i in range(4):
            cb, holder = self._color_button(d, EX.HL_COLORS[i][:3])
            cb.grid(row=1 + i, column=0, padx=(0, 8), pady=3)
            nv = tk.StringVar(value=self.cfg["hl_names"][i])
            ttk.Entry(d, textvariable=nv, width=28).grid(row=1 + i, column=1, sticky="w")
            rows.append((holder, nv))

        def go():
            EX.HL_COLORS[:] = [tuple(h["rgb"]) for h, _n in rows]
            self.cfg["hl_colors"] = [list(c) for c in EX.HL_COLORS]
            self.cfg["hl_names"] = [n.get().strip()[:40] for _h, n in rows]
            save_cfg(self.cfg)
            for b, col in zip(self.hl_buttons, EX.HL_COLORS):
                b.config(bg=self._hex(col), activebackground=self._hex(col), command=lambda c=col: setattr(self, "hl_color", c))
            self.hl_color = EX.HL_COLORS[0]
            d.destroy()
        self._buttons(d, "Save", go, 5, colspan=3)

    def collect_highlights(self):
        out = []
        for pno in range(len(self.doc)):
            pg = self.doc[pno]
            for a in pg.annots() or []:
                if a.type[0] != fitz.PDF_ANNOT_HIGHLIGHT:
                    continue
                try:
                    verts = a.vertices or []
                    quads = [fitz.Quad(verts[i:i + 4]) for i in range(0, len(verts), 4)] if verts else []
                    text = " ".join(pg.get_textbox(q.rect).strip().replace("\n", " ") for q in quads) if quads \
                        else pg.get_textbox(a.rect).strip().replace("\n", " ")
                except Exception:
                    text = ""
                col = (a.colors or {}).get("stroke") or (1, 1, 0)
                out.append({"page": pno + 1, "text": re.sub(r"\s+", " ", text).strip(), "note": (a.info or {}).get("content", ""),
                            "color": col, "label": self._v3_hl_label(col)})
        return out

    def dlg_export_highlights(self):
        if not self.need_doc():
            return
        items = self.collect_highlights()
        if not items:
            messagebox.showinfo(APP, "There are no highlights in this document.")
            return
        labels = []
        for it in items:
            if it["label"] not in labels:
                labels.append(it["label"])
        d = self._dlg("Export highlights")
        self._label(d, "{} highlights found. Include:".replace("{}", str(len(items)))).grid(row=0, column=0, sticky="w")
        vars_ = {}
        for i, lab in enumerate(labels):
            vars_[lab] = tk.BooleanVar(value=True)
            n = sum(1 for it in items if it["label"] == lab)
            tk.Checkbutton(d, text=f"{lab}  ({n})", variable=vars_[lab], bg=BG, activebackground=BG, fg=TH.LABEL,
                           selectcolor=TH.CARD, font=(UI_FONT, 9)).grid(row=1 + i, column=0, sticky="w")
        group = tk.BooleanVar(value=True)
        tk.Checkbutton(d, text="Group by colour meaning", variable=group, bg=BG, activebackground=BG, fg=TH.LABEL,
                       selectcolor=TH.CARD, font=(UI_FONT, 9)).grid(row=len(labels) + 1, column=0, sticky="w", pady=(8, 0))

        def go():
            chosen = [it for it in items if vars_[it["label"]].get()]
            if not chosen:
                return
            base = os.path.splitext(os.path.basename(self.path or "document"))[0]
            p = filedialog.asksaveasfilename(parent=d, defaultextension=".docx", initialfile=base + "_highlights.docx",
                                             filetypes=[("Word document", "*.docx"), ("Text file", "*.txt"), ("Markdown", "*.md")])
            if not p:
                return
            d.destroy()
            try:
                self.write_highlights(p, chosen, labels if group.get() else None, base)
            except Exception as ex:
                messagebox.showerror(APP, f"Could not save:\n{ex}")
                return
            self.status.config(text=_tr("Highlights exported."))
        self._buttons(d, "Export", go, len(labels) + 2, colspan=1)

    def write_highlights(self, path, items, order, title):
        groups = [(None, items)] if not order else [(lab, [i for i in items if i["label"] == lab]) for lab in order]
        groups = [(lab, its) for lab, its in groups if its]
        ext = os.path.splitext(path)[1].lower()
        clean = lambda s: re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", s or "")
        title = clean(title)
        for it in items:
            it["text"], it["note"] = clean(it["text"]), clean(it["note"])
        if ext == ".docx":
            import docx
            doc = docx.Document()
            doc.add_heading(title, level=1)
            for lab, its in groups:
                if lab:
                    doc.add_heading(lab, level=2)
                for it in its:
                    para = doc.add_paragraph(style="List Bullet")
                    para.add_run(f"p. {it['page']}: ").bold = True
                    para.add_run(it["text"] or "(no text)")
                    if it["note"]:
                        para.add_run("\n" + it["note"]).italic = True
            doc.save(path)
            return
        lines = [title, ""] if ext != ".md" else ["# " + title, ""]
        for lab, its in groups:
            if lab:
                lines += [(lab if ext != ".md" else "## " + lab), ""]
            for it in its:
                lines.append(f"- p. {it['page']}: {it['text'] or '(no text)'}")
                if it["note"]:
                    lines.append(f"    ({it['note']})")
            lines.append("")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("\n".join(lines))

    # ------------------------------------------------------------------ look up / translate
    def dlg_lookup(self, text=None):
        if text is None:
            try:
                text = self._sel_text()
            except Exception:
                text = ""
        d = self._dlg("Look up / translate")
        self._label(d, "Text").grid(row=0, column=0, sticky="nw", pady=3)
        t = tk.Text(d, width=44, height=4, wrap="word", font=(UI_FONT, 10), relief="solid", bd=1)
        t.grid(row=0, column=1, columnspan=2, pady=3)
        t.insert("1.0", (text or "").strip()[:1500])
        langs = (("en", "English"), ("pl", "Polski"), ("de", "Deutsch"), ("es", "Español"), ("it", "Italiano"), ("fr", "Français"))
        import claudius_i18n as i18n
        target = tk.StringVar(value=dict(langs).get(i18n.current, "English"))
        self._label(d, "Translate into").grid(row=1, column=0, sticky="w", pady=3)
        ttk.Combobox(d, textvariable=target, values=[n for _c, n in langs], state="readonly", width=18).grid(row=1, column=1, sticky="w")
        self._label(d, "The text you look up is sent to the website you choose. Nothing is sent otherwise.", fg=TH.TERTIARY).grid(
            row=2, column=0, columnspan=3, sticky="w", pady=(4, 0))

        def go(kind):
            q = t.get("1.0", "end").strip()
            if not q:
                return
            code = {n: c for c, n in langs}.get(target.get(), "en")
            enc = urllib.parse.quote(q[:1500])
            if kind == "google":
                url = f"https://translate.google.com/?sl=auto&tl={code}&text={enc}&op=translate"
            elif kind == "deepl":
                url = f"https://www.deepl.com/translator#auto/{code}/{enc}"
            elif kind == "dict":
                word = urllib.parse.quote(q.split()[0][:60]) if q.split() else enc
                url = f"https://{code}.wiktionary.org/wiki/{word}"
            else:
                url = f"https://duckduckgo.com/?q={enc}"
            webbrowser.open(url)
        f = tk.Frame(d, bg=BG)
        f.grid(row=3, column=0, columnspan=3, sticky="e", pady=(12, 0))
        for label, kind in (("Translate (Google)", "google"), ("Translate (DeepL)", "deepl"), ("Dictionary", "dict"), ("Web search", "web")):
            ttk.Button(f, text=label, command=lambda k=kind: go(k)).pack(side="left", padx=3)
        ttk.Button(f, text="Close", command=d.destroy).pack(side="left", padx=(8, 0))

    # ------------------------------------------------------------------ redaction check
    def dlg_redaction_check(self):
        if not self.need_doc():
            return
        self.commit_floats()
        d = self._dlg("Check redaction")
        self._label(d, "Words or phrases that must NOT be in the document any more (one per line, optional):").grid(row=0, column=0, sticky="w")
        terms = tk.Text(d, width=60, height=4, font=(UI_FONT, 10), relief="solid", bd=1)
        terms.grid(row=1, column=0, pady=4)
        tree = ttk.Treeview(d, columns=("lvl", "where", "what"), show="headings", height=11)
        for c, t, w in (("lvl", "", 34), ("where", "Where", 170), ("what", "Finding", 400)):
            tree.heading(c, text=_tr(t) if t else "")
            tree.column(c, width=S(w), anchor="w" if c != "lvl" else "center")
        tree.grid(row=2, column=0, pady=4)
        summary = self._label(d, "")
        summary.grid(row=3, column=0, sticky="w")

        def run():
            tree.delete(*tree.get_children())
            words = [x.strip() for x in terms.get("1.0", "end").splitlines() if x.strip()]
            upd, close = self._progress("Checking...", len(self.doc) + 1)
            try:
                rows = self.check_redaction(words, upd)
            finally:
                close()
            bad = sum(1 for r in rows if r[0] == "x")
            warn = sum(1 for r in rows if r[0] == "!")
            for lvl, where, what in rows:
                tree.insert("", "end", values=({"x": "✖", "!": "!", "ok": "✔"}[lvl], where, what))
            if not rows:
                tree.insert("", "end", values=("✔", "", _tr("Nothing suspicious found.")))
            summary.config(text=_tr("{} problem(s), {} warning(s).").replace("{}", str(bad), 1).replace("{}", str(warn), 1))
        f = tk.Frame(d, bg=BG)
        f.grid(row=4, column=0, sticky="e", pady=(10, 0))
        ttk.Button(f, text="Check now", command=run).pack(side="left", padx=4)
        ttk.Button(f, text="Clean hidden data...", command=lambda: (d.destroy(), self.dlg_clean_metadata())).pack(side="left", padx=4)
        ttk.Button(f, text="Close", command=d.destroy).pack(side="left")

    def check_redaction(self, words, upd=None):
        rows = []
        doc = self.doc
        low = [w.lower() for w in words]

        def has(txt):
            t = (txt or "").lower()
            return [w for w, lw in zip(words, low) if lw in t]
        for pno in range(len(doc)):
            pg = doc[pno]
            where = f"{_tr('Page')} {pno + 1}"
            n_red = sum(1 for a in (pg.annots() or []) if a.type[0] == fitz.PDF_ANNOT_REDACT)
            if n_red:
                rows.append(("x", where, _tr("{} redaction mark(s) are not applied yet - the text below them is still there.").replace("{}", str(n_red))))
            try:
                hidden = 0
                white = 0
                for tr_ in pg.get_texttrace():
                    chars = len(tr_.get("chars") or [])
                    if tr_.get("type") == 3:
                        hidden += chars
                    elif tr_.get("color") == (1.0, 1.0, 1.0) or tr_.get("color") == 16777215:
                        white += chars
                if hidden:
                    rows.append(("!", where, _tr("{} invisible character(s) (for example an OCR text layer).").replace("{}", str(hidden))))
                if white:
                    rows.append(("!", where, _tr("{} white character(s) - may be hidden text on a white background.").replace("{}", str(white))))
            except Exception:
                pass
            if words:
                text = pg.get_text()
                for w in has(text):
                    rows.append(("x", where, _tr("The text \"{}\" is still on the page.").replace("{}", w)))
                try:
                    raw = pg.read_contents().decode("latin-1", "ignore").lower()
                    for w, lw in zip(words, low):
                        if lw in raw and w not in has(text):
                            rows.append(("!", where, _tr("\"{}\" appears in the page data.").replace("{}", w)))
                except Exception:
                    pass
                for a in pg.annots() or []:
                    for w in has((a.info or {}).get("content", "") + " " + (a.info or {}).get("title", "")):
                        rows.append(("x", where, _tr("The text \"{}\" is in a comment.").replace("{}", w)))
                for wd in pg.widgets() or []:
                    for w in has(str(wd.field_value or "") + " " + str(wd.field_name or "")):
                        rows.append(("x", where, _tr("The text \"{}\" is in a form field.").replace("{}", w)))
            if upd:
                upd(pno + 1)
        md = doc.metadata or {}
        info = [f"{k}: {v}" for k, v in md.items() if v and k in ("title", "author", "subject", "keywords", "creator", "producer")]
        if info:
            rows.append(("!", _tr("Document info"), "; ".join(info)[:200]))
        for w in has(" ".join(str(v) for v in md.values())):
            rows.append(("x", _tr("Document info"), _tr("The text \"{}\" is in the document info.").replace("{}", w)))
        try:
            xml = doc.get_xml_metadata()
            if xml:
                rows.append(("!", "XMP", _tr("The file has XMP metadata.")))
                for w in has(xml):
                    rows.append(("x", "XMP", _tr("The text \"{}\" is in the XMP metadata.").replace("{}", w)))
        except Exception:
            pass
        try:
            for lvl, title, page in doc.get_toc():
                for w in has(title):
                    rows.append(("x", _tr("Bookmarks"), _tr("The text \"{}\" is in a bookmark.").replace("{}", w)))
        except Exception:
            pass
        try:
            names = doc.embfile_names()
            if names:
                rows.append(("!", _tr("Attachments"), ", ".join(names)[:200]))
                for nme in names:
                    try:
                        blob = doc.embfile_get(nme).decode("utf-8", "ignore")
                        for w in has(blob):
                            rows.append(("x", _tr("Attachments"), _tr("The text \"{}\" is in the attachment {}.").replace("{}", w, 1).replace("{}", nme, 1)))
                    except Exception:
                        pass
        except Exception:
            pass
        try:
            off = [o.get("text") for o in (doc.get_ocgs() or {}).values() if o.get("on") is False]
            if off:
                rows.append(("!", _tr("Layers"), _tr("Hidden layers: {}").replace("{}", ", ".join(str(x) for x in off)[:150])))
        except Exception:
            pass
        try:
            if self.path and os.path.exists(self.path) and os.path.getsize(self.path) < 300 * 1024 * 1024:
                with open(self.path, "rb") as fh:
                    if fh.read().count(b"%%EOF") > 1:
                        rows.append(("!", _tr("File"), _tr("The saved file contains older versions of the document. Use Save As to write a clean copy.")))
        except Exception:
            pass
        return rows

    # ------------------------------------------------------------------ metadata cleaner
    def dlg_clean_metadata(self):
        if not self.need_doc():
            return
        d = self._dlg("Clean document data")
        md = self.doc.metadata or {}
        shown = [(k, v) for k, v in md.items() if v and k in ("title", "author", "subject", "keywords", "creator", "producer", "creationDate", "modDate")]
        self._label(d, "Now in the document:", bold=True).grid(row=0, column=0, sticky="w")
        self._label(d, "\n".join(f"{k}: {v}"[:90] for k, v in shown) or _tr("(no document info)"), fg=TH.SECOND).grid(row=1, column=0, sticky="w", pady=(0, 8))
        opts = [("metadata", "Document info (title, author, programs used ...)", True),
                ("xml_metadata", "XMP metadata", True),
                ("attached_files", "Attached files", False),
                ("javascript", "Scripts (JavaScript)", True),
                ("thumbnails", "Small page pictures (thumbnails)", True),
                ("hidden_text", "Hidden text", True),
                ("comments", "All comments and markup", False),
                ("links", "Links", False)]
        v = {k: tk.BooleanVar(value=on) for k, _t, on in opts}
        for i, (k, t, _on) in enumerate(opts):
            tk.Checkbutton(d, text=t, variable=v[k], bg=BG, activebackground=BG, fg=TH.LABEL, selectcolor=TH.CARD,
                           font=(UI_FONT, 9)).grid(row=2 + i, column=0, sticky="w")

        def go():
            sel = {k: bool(x.get()) for k, x in v.items()}
            d.destroy()
            self.do_clean_metadata(sel)
        self._buttons(d, "Clean", go, 12, colspan=1)

    def do_clean_metadata(self, sel):
        def work():
            if sel.get("comments"):
                for pg in self.doc:
                    for a in list(pg.annots() or []):
                        pg.delete_annot(a)
            self.doc.scrub(metadata=sel["metadata"], xml_metadata=sel["xml_metadata"], attached_files=sel["attached_files"],
                           embedded_files=sel["attached_files"], javascript=sel["javascript"], thumbnails=sel["thumbnails"],
                           hidden_text=sel["hidden_text"], remove_links=sel["links"], redactions=False, clean_pages=False,
                           reset_fields=False, reset_responses=False, redact_images=0)
        return self._run_edit(work, rebuild=True, msg=_tr("Document data cleaned."))
