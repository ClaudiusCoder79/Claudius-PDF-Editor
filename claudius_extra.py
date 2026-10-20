"""
Claudius PDF Editor - extra tools (page tools, markup, signing, forms, conversion, OCR ...).

This module is imported by claudius.py - keep both files in the same folder.
"""
import datetime
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import textwrap
import time
import webbrowser
import tkinter as tk
import tkinter.font
from tkinter import ttk, filedialog, messagebox, simpledialog, colorchooser
from claudius_i18n import tr as _tr

try:
    import pymupdf as fitz
except ImportError:
    import fitz
from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageOps, ImageTk

APP = "Claudius PDF Editor"
UI_FONT = "Segoe UI"
import claudius_theme as TH
from claudius_theme import round_rect
BG = TH.BG
ACCENT = TH.ACCENT
ACTIVE_BG = TH.TINT_DOWN
HOVER_BG = TH.HOVER
SIDEBAR_BG = TH.BG
LINE_COL = TH.LINE_SOFT


def _detect_scale():
    """Windows: tell the system this program handles high-DPI screens itself (otherwise Windows stretches
    the whole window as a bitmap, which makes text and icons look blurry).  Returns the UI scale factor."""
    env = os.environ.get("CLAUDIUS_SCALE")
    if env:
        try:
            return max(0.75, float(env))
        except ValueError:
            pass
    if sys.platform == "win32":
        try:
            import ctypes
            try:
                ctypes.windll.shcore.SetProcessDpiAwareness(1)
            except Exception:
                ctypes.windll.user32.SetProcessDPIAware()
            try:
                dpi = ctypes.windll.user32.GetDpiForSystem()
            except Exception:
                dpi = 96
            return max(1.0, dpi / 96.0)
        except Exception:
            pass
    return 1.0


SCALE = _detect_scale()

if sys.platform == "win32":                    # own identity in the taskbar (icon, grouping, pinning)
    try:
        import ctypes as _ct
        _ct.windll.shell32.SetCurrentProcessExplicitAppUserModelID("Claudius.PdfEditor.1")
    except Exception:
        pass


def S(n):
    """scale a pixel size for the current screen"""
    return int(round(n * SCALE))


CELL_H = S(196)                    # thumbnail cell height - must match claudius.py
RES = getattr(Image, "Resampling", Image)
LANCZOS, BICUBIC = RES.LANCZOS, RES.BICUBIC
CONFIG_DIR = os.path.join(os.path.expanduser("~"), ".claudius")
NO_WINDOW = 0x08000000 if sys.platform == "win32" else 0        # hide console windows of helper programs

ALL_PERMS = 0
for _n in ("PRINT", "COPY", "MODIFY", "ANNOTATE", "FORM", "ACCESSIBILITY", "ASSEMBLE", "PRINT_HQ"):
    ALL_PERMS |= getattr(fitz, "PDF_PERM_" + _n, 0)

IMG_EXT = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp")
OFFICE_EXT = (".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".rtf", ".odt", ".ods", ".odp")

STAMPS = [("APPROVED", (0.10, 0.55, 0.20)), ("REJECTED", (0.80, 0.10, 0.10)), ("DRAFT", (0.35, 0.40, 0.55)),
          ("CONFIDENTIAL", (0.80, 0.10, 0.10)), ("FINAL", (0.10, 0.55, 0.20)), ("COMPLETED", (0.10, 0.55, 0.20)),
          ("FOR REVIEW", (0.15, 0.40, 0.75)), ("SIGN HERE", (0.90, 0.50, 0.05)), ("URGENT", (0.80, 0.10, 0.10)),
          ("COPY", (0.35, 0.40, 0.55)), ("PAID", (0.10, 0.55, 0.20)), ("VOID", (0.80, 0.10, 0.10))]


# ======================================================================================
#  small pure helpers (no GUI) - these are also what the automated tests call directly
# ======================================================================================
def parse_range_groups(text, n):
    """'1-3, 5, 7-' -> [[0,1,2],[4],[6..n-1]].  Also understands all / odd / even.  Raises ValueError."""
    text = (text or "").strip().lower()
    if text in ("", "all"):
        return [list(range(n))]
    groups = []
    for part in re.split(r"[;,]", text):
        part = part.strip()
        if not part:
            continue
        if part == "all":
            groups.append(list(range(n)))
            continue
        if part == "odd":
            groups.append(list(range(0, n, 2)))
            continue
        if part == "even":
            groups.append(list(range(1, n, 2)))
            continue
        m = re.fullmatch(r"(\d*)\s*-\s*(\d*)", part)
        if m and (m.group(1) or m.group(2)):
            a = int(m.group(1)) if m.group(1) else 1
            b = int(m.group(2)) if m.group(2) else n
        elif part.isdigit():
            a = b = int(part)
        else:
            raise ValueError(f"I can't understand '{part}'. Use something like 1-3, 5, 8-")
        if a > b:
            a, b = b, a
        if a < 1 or b > n:
            raise ValueError(f"Pages must be between 1 and {n} (you typed {part}).")
        groups.append(list(range(a - 1, b)))
    if not groups:
        raise ValueError("No pages selected.")
    return groups


def parse_pages(text, n):
    seen, out = set(), []
    for g in parse_range_groups(text, n):
        for p in g:
            if p not in seen:
                seen.add(p)
                out.append(p)
    return out


def pages_to_text(pages):
    pages = sorted(set(pages))
    out, i = [], 0
    while i < len(pages):
        j = i
        while j + 1 < len(pages) and pages[j + 1] == pages[j] + 1:
            j += 1
        out.append(str(pages[i] + 1) if i == j else f"{pages[i] + 1}-{pages[j] + 1}")
        i = j + 1
    return ", ".join(out)


def runs(pages):
    """[0,1,2,5,6] -> [(0,2),(5,6)] contiguous runs"""
    pages = sorted(pages)
    out = []
    for p in pages:
        if out and p == out[-1][1] + 1:
            out[-1] = (out[-1][0], p)
        else:
            out.append((p, p))
    return out


def human_size(nbytes):
    for unit in ("bytes", "KB", "MB", "GB"):
        if nbytes < 1024 or unit == "GB":
            return f"{nbytes:.0f} {unit}" if unit == "bytes" else f"{nbytes:.1f} {unit}"
        nbytes /= 1024.0


def expand_tokens(text, number, total, fname=""):
    now = datetime.datetime.now()
    return (text.replace("{page}", str(number)).replace("{pages}", str(total))
            .replace("{date}", now.strftime("%Y-%m-%d")).replace("{time}", now.strftime("%H:%M"))
            .replace("{file}", fname))


def safe_name(s):
    return re.sub(r'[\\/:*?"<>|]+', "_", s).strip() or "file"


# ---------------------------------------------------------------------- fonts
_FONT_DIRS = [r"C:\Windows\Fonts", "/Library/Fonts", "/System/Library/Fonts/Supplemental",
              "/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/truetype/crosextra",
              "/usr/share/fonts/truetype/liberation", os.path.expanduser("~/.fonts")]
if os.environ.get("LOCALAPPDATA"):            # Windows 10/11: fonts installed "for me only"
    _FONT_DIRS.append(os.path.join(os.environ["LOCALAPPDATA"], "Microsoft", "Windows", "Fonts"))
_FONT_NAMES = {
    "regular": ["arial.ttf", "Arial.ttf", "DejaVuSans.ttf", "Carlito-Regular.ttf", "LiberationSans-Regular.ttf"],
    "bold": ["arialbd.ttf", "Arial Bold.ttf", "DejaVuSans-Bold.ttf", "Carlito-Bold.ttf", "LiberationSans-Bold.ttf"],
    "italic": ["ariali.ttf", "Arial Italic.ttf", "DejaVuSans-Oblique.ttf", "Carlito-Italic.ttf"],
}
SCRIPT_FONTS = [("Script", ["segoesc.ttf", "Zapfino.ttf"]), ("Handwriting", ["Inkfree.ttf", "LHANDW.TTF"]),
                ("Brush", ["BRUSHSCI.TTF"]), ("Calligraphy", ["MTCORSVA.TTF", "ITCEDSCR.TTF"])]


def _find_in_dirs(names):
    for d in _FONT_DIRS:
        for n in names:
            p = os.path.join(d, n)
            if os.path.exists(p):
                return p
    return None


def find_font(kind="regular"):
    return _find_in_dirs(_FONT_NAMES[kind])


def signature_fonts():
    out = []
    for label, names in SCRIPT_FONTS:
        p = _find_in_dirs(names)
        if p:
            out.append((label, p))
    if not out:                                   # no script fonts installed: fall back to italic
        p = find_font("italic") or find_font("regular")
        if p:
            out.append(("Italic", p))
    return out


def load_pil_font(path, size):
    try:
        if path:
            return ImageFont.truetype(path, size)
    except Exception:
        pass
    try:
        return ImageFont.load_default(size)
    except TypeError:
        return ImageFont.load_default()


def needs_unicode(text):
    return any(ord(c) > 255 for c in text)


def font_kwargs(text="", bold=False):
    """Keyword arguments for Page.insert_text / insert_textbox.  Uses a TrueType font only when the
    text needs characters outside the standard PDF fonts."""
    if needs_unicode(text):
        p = find_font("bold" if bold else "regular")
        if p:
            return {"fontname": "UF" + ("B" if bold else "R"), "fontfile": p}
    return {"fontname": "hebo" if bold else "helv"}


_FONT_CACHE = {}


def text_width(text, size, fk):
    if "fontfile" in fk:
        f = _FONT_CACHE.get(fk["fontfile"])
        if f is None:
            f = _FONT_CACHE[fk["fontfile"]] = fitz.Font(fontfile=fk["fontfile"])
        return f.text_length(text, fontsize=size)
    return fitz.get_text_length(text, fontname=fk["fontname"], fontsize=size)


def put_text(pg, x, y, text, size, color, fk, align="left"):
    """Write text so it is upright for the viewer.  (x, y) is the baseline start in VIEW coordinates
    (what you see on screen), which also works on rotated pages."""
    w = text_width(text, size, fk)
    if align == "center":
        x -= w / 2
    elif align == "right":
        x -= w
    pt = fitz.Point(x, y) * pg.derotation_matrix
    pg.insert_text(pt, text, fontsize=size, color=color, rotate=pg.rotation, **fk)


def base14_name(fontname, flags):
    n = (fontname or "").lower()
    bold = bool(flags & 16) or "bold" in n or "black" in n or "heavy" in n
    ital = bool(flags & 2) or "italic" in n or "oblique" in n
    if "courier" in n or "mono" in n or "consol" in n or bool(flags & 8):
        fam = "co"
    elif "times" in n or "georgia" in n or "garamond" in n or "palatino" in n or "cambria" in n or \
            ("serif" in n and "sans" not in n) or (bool(flags & 4) and "arial" not in n and "helvet" not in n):
        fam = "ti"
    else:
        fam = "he"
    table = {("he", 0, 0): "helv", ("he", 1, 0): "hebo", ("he", 0, 1): "heit", ("he", 1, 1): "hebi",
             ("ti", 0, 0): "tiro", ("ti", 1, 0): "tibo", ("ti", 0, 1): "tiit", ("ti", 1, 1): "tibi",
             ("co", 0, 0): "cour", ("co", 1, 0): "cobo", ("co", 0, 1): "coit", ("co", 1, 1): "cobi"}
    return table[(fam, int(bold), int(ital))]


# --------------------------------------------------------------- generated images
def make_text_image(text, font_path, size, color, pad=12):
    fnt = load_pil_font(font_path, size)
    probe = ImageDraw.Draw(Image.new("RGBA", (10, 10)))
    l, t, r, b = probe.textbbox((0, 0), text, font=fnt)
    img = Image.new("RGBA", (int(r - l + 2 * pad), int(b - t + 2 * pad)), (0, 0, 0, 0))
    ImageDraw.Draw(img).text((pad - l, pad - t), text, font=fnt, fill=tuple(int(c * 255) for c in color) + (255,))
    return img


def make_stamp_image(text, color, subtitle=None):
    col = tuple(int(c * 255) for c in color) + (255,)
    f1 = load_pil_font(find_font("bold"), 120)
    f2 = load_pil_font(find_font("regular"), 46)
    probe = ImageDraw.Draw(Image.new("RGBA", (10, 10)))
    l, t, r, b = probe.textbbox((0, 0), text, font=f1)
    tw, th = r - l, b - t
    sw = sh = 0
    if subtitle:
        l2, t2, r2, b2 = probe.textbbox((0, 0), subtitle, font=f2)
        sw, sh = r2 - l2, b2 - t2
    W = max(tw, sw) + 150
    H = th + (sh + 34 if subtitle else 0) + 110
    img = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((6, 6, W - 7, H - 7), radius=38, outline=col, width=11)
    d.rounded_rectangle((26, 26, W - 27, H - 27), radius=26, outline=col, width=4)
    y = 55
    d.text(((W - tw) / 2 - l, y - t), text, font=f1, fill=col)
    if subtitle:
        d.text(((W - sw) / 2 - l2, y + th + 24 - t2), subtitle, font=f2, fill=col)
    return img


def make_watermark_image(text, color, opacity, angle):
    img = make_text_image(text, find_font("bold"), 220, color, pad=20)
    if opacity < 100:
        a = img.getchannel("A").point(lambda v: v * opacity // 100)
        img.putalpha(a)
    if abs(angle) % 360 > 0.01:
        img = img.rotate(angle, resample=BICUBIC, expand=True, fillcolor=(0, 0, 0, 0))
    bbox = img.getchannel("A").getbbox()
    return img.crop(bbox) if bbox else img


def apply_opacity(img, opacity):
    img = img.convert("RGBA")
    if opacity < 100:
        img.putalpha(img.getchannel("A").point(lambda v: v * opacity // 100))
    return img


def remove_white(img, thr=235):
    """Make (nearly) white pixels transparent - for scanned signatures."""
    img = img.convert("RGBA")
    soft = 25
    lut = [0 if v >= thr else (255 if v < thr - soft else int((thr - v) * 255 / soft)) for v in range(256)]
    mask = img.convert("L").point(lut)
    img.putalpha(ImageChops.multiply(img.getchannel("A"), mask))
    bbox = img.getchannel("A").getbbox()
    return img.crop(bbox) if bbox else img



# ---------------------------------------------------------------------- compression
COMPRESS_LEVELS = {            # name -> (dpi_threshold, dpi_target, jpeg quality)
    "Low  (best quality)": (250, 200, 85),
    "Medium  (recommended)": (200, 150, 70),
    "High  (smallest file)": (150, 96, 55),
}


def compress_bytes(data, level="Medium  (recommended)", strip_meta=False):
    thr, target, q = COMPRESS_LEVELS[level]
    doc = fitz.open(stream=data, filetype="pdf")
    if doc.needs_pass:
        raise ValueError("The file is password protected.")
    if hasattr(doc, "rewrite_images"):
        doc.rewrite_images(dpi_threshold=thr, dpi_target=target, quality=q)
    if strip_meta:
        doc.set_metadata({})
        try:
            doc.del_xml_metadata()
        except Exception:
            pass
    out = doc.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    doc.close()
    return out


# ------------------------------------------------------------------------------ OCR
def find_tesseract():
    exe = shutil.which("tesseract")
    if exe:
        return exe
    for p in (r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
              os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
              "/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract", "/usr/bin/tesseract"):
        if os.path.exists(p):
            return p
    return None


def tessdata_dir():
    cands = []
    try:
        g = fitz.get_tessdata()
        if g:
            cands.append(g)
    except Exception:
        pass
    env = os.environ.get("TESSDATA_PREFIX")
    if env:
        cands += [env, os.path.join(env, "tessdata")]
    exe = find_tesseract()
    if exe:
        cands.append(os.path.join(os.path.dirname(exe), "tessdata"))
    cands += ["/usr/share/tesseract-ocr/5/tessdata", "/usr/share/tesseract-ocr/4.00/tessdata", "/usr/share/tessdata",
              "/usr/local/share/tessdata", "/opt/homebrew/share/tessdata"]
    for c in cands:
        try:
            if c and os.path.isdir(c) and any(f.endswith(".traineddata") for f in os.listdir(c)):
                return c
        except OSError:
            pass
    return None


def ocr_languages(tessdata):
    try:
        langs = sorted(f[:-12] for f in os.listdir(tessdata) if f.endswith(".traineddata") and f != "osd.traineddata")
    except Exception:
        langs = []
    return langs or ["eng"]


def ocr_document(doc, pages, lang, tessdata, dpi=250, skip_text=True, progress=None):
    """Return a NEW document where the chosen pages are searchable (invisible OCR text over the scan)."""
    pages = set(pages)
    out = fitz.open()
    n = len(doc)
    done = 0
    for k in range(n):
        pg = doc[k]
        if k in pages and not (skip_text and pg.get_text().strip()):
            pix = pg.get_pixmap(dpi=safe_dpi(pg, dpi))
            data = pix.pdfocr_tobytes(language=lang, tessdata=tessdata)
            o = fitz.open("pdf", data)
            out.insert_pdf(o)
            o.close()
            done += 1
        else:
            out.insert_pdf(doc, from_page=k, to_page=k)
        if progress:
            progress(k + 1, n)
    try:
        out.set_metadata(doc.metadata)
        toc = doc.get_toc()
        if toc:
            out.set_toc(toc)
    except Exception:
        pass
    out._ocr_pages = done
    return out


# ------------------------------------------------------------------ conversions (from PDF)
def pdf_to_images(doc, pages, fmt, dpi, outdir, base, progress=None):
    files = []
    ext = {"PNG": "png", "JPG": "jpg", "TIFF": "tif", "BMP": "bmp"}[fmt]
    width = len(str(len(doc)))
    for k, pno in enumerate(pages):
        pix = doc[pno].get_pixmap(dpi=safe_dpi(doc[pno], dpi), alpha=False)
        path = os.path.join(outdir, f"{base}_page{str(pno + 1).zfill(width)}.{ext}")
        if fmt == "PNG":
            pix.save(path)
        else:
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            img.save(path, {"JPG": "JPEG", "TIFF": "TIFF", "BMP": "BMP"}[fmt], **({"quality": 92} if fmt == "JPG" else {}))
        files.append(path)
        if progress:
            progress(k + 1, len(pages))
    return files


def pdf_to_text(doc, pages, path):
    parts = [doc[p].get_text("text", sort=True) for p in pages]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n\n".join(parts))


def pdf_to_html(doc, pages, path):
    body = "\n".join(doc[p].get_text("html") for p in pages)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("<!DOCTYPE html>\n<html><head><meta charset='utf-8'><title>%s</title></head>\n<body>\n%s\n</body></html>"
                 % (os.path.basename(path), body))


def pdf_to_xml(doc, pages, path):
    body = "\n".join(doc[p].get_text("xml") for p in pages)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("<?xml version='1.0' encoding='UTF-8'?>\n<document>\n%s\n</document>" % body)


def _clean_font_name(name):
    name = (name or "").split("+")[-1]
    name = re.split(r"[-,]", name)[0]
    return re.sub(r"(MT|PS)$", "", name) or "Arial"


def pdf_to_docx(doc, pages, path, progress=None):
    import docx
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    first = doc[pages[0]].rect
    sec = d.sections[0]
    sec.page_width, sec.page_height = Pt(first.width), Pt(first.height)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Pt(36))
    avail = max(100, first.width - 72)
    for idx, pno in enumerate(pages):
        data = doc[pno].get_text("dict", sort=True)
        for b in data["blocks"]:
            if b["type"] == 0:
                par = d.add_paragraph()
                par.paragraph_format.space_after = Pt(4)
                lines = b["lines"]
                for li, ln in enumerate(lines):
                    for sp in ln["spans"]:
                        if not sp["text"]:
                            continue
                        run = par.add_run(sp["text"])
                        run.font.size = Pt(max(4, round(sp["size"], 1)))
                        run.bold = bool(sp["flags"] & 16) or "bold" in sp["font"].lower()
                        run.italic = bool(sp["flags"] & 2) or "italic" in sp["font"].lower()
                        run.font.name = _clean_font_name(sp["font"])
                        c = sp["color"]
                        run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
                    if li < len(lines) - 1:
                        par.add_run(" ")
            elif b["type"] == 1:
                try:
                    wpt = min(avail, b["bbox"][2] - b["bbox"][0])
                    d.add_picture(io.BytesIO(b["image"]), width=Pt(max(10, wpt)))
                except Exception:
                    pass
        if idx < len(pages) - 1:
            d.add_page_break()
        if progress:
            progress(idx + 1, len(pages))
    d.save(path)


def _cell_value(v):
    if v is None:
        return None
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+[.,]\d+", s):
        return float(s.replace(",", "."))
    return s


def pdf_to_xlsx(doc, pages, path, progress=None):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ntables = 0
    for k, pno in enumerate(pages):
        page = doc[pno]
        try:
            tabs = page.find_tables().tables
        except Exception:
            tabs = []
        for t, tab in enumerate(tabs, 1):
            ws = wb.create_sheet(f"P{pno + 1}_T{t}"[:31])
            for r, row in enumerate(tab.extract(), 1):
                for c, val in enumerate(row, 1):
                    v = _cell_value(val)
                    if v is not None and v != "":
                        ws.cell(row=r, column=c, value=v)
            ntables += 1
        if not tabs:                                                # no table: keep the text, one line per row
            lines = [ln for ln in page.get_text("text", sort=True).splitlines() if ln.strip()]
            if lines:
                ws = wb.create_sheet(f"P{pno + 1}_text"[:31])
                for r, ln in enumerate(lines, 1):
                    ws.cell(row=r, column=1, value=ln)
        if progress:
            progress(k + 1, len(pages))
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    wb.save(path)
    return ntables


def pdf_to_pptx(doc, pages, path, mode="picture", dpi=150, progress=None):
    from pptx import Presentation
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    prs = Presentation()
    first = doc[pages[0]].rect
    prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    for k, pno in enumerate(pages):
        page = doc[pno]
        slide = prs.slides.add_slide(blank)
        s = prs.slide_width / (page.rect.width * 12700.0)             # fit the page onto the slide
        if mode == "picture":
            pix = page.get_pixmap(dpi=safe_dpi(page, dpi))
            slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0, prs.slide_width,
                                     int(page.rect.height * 12700 * s))
            try:
                slide.notes_slide.notes_text_frame.text = page.get_text()
            except Exception:
                pass
        else:
            data = page.get_text("dict")
            for b in data["blocks"]:
                x0, y0, x1, y1 = b["bbox"]
                if b["type"] == 1:
                    try:
                        slide.shapes.add_picture(io.BytesIO(b["image"]), Emu(int(x0 * 12700 * s)), Emu(int(y0 * 12700 * s)),
                                                 Emu(int((x1 - x0) * 12700 * s)), Emu(int((y1 - y0) * 12700 * s)))
                    except Exception:
                        pass
                    continue
                for ln in b["lines"]:
                    text = "".join(sp["text"] for sp in ln["spans"])
                    if not text.strip():
                        continue
                    lx0, ly0, lx1, ly1 = ln["bbox"]
                    sp0 = ln["spans"][0]
                    tb = slide.shapes.add_textbox(Emu(int(lx0 * 12700 * s)), Emu(int(ly0 * 12700 * s)),
                                                  Emu(int((lx1 - lx0 + 8) * 12700 * s)), Emu(int((ly1 - ly0) * 12700 * s)))
                    tf = tb.text_frame
                    tf.word_wrap = False
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    run = tf.paragraphs[0].add_run()
                    run.text = text
                    run.font.size = Pt(max(4, sp0["size"] * s))
                    run.font.bold = bool(sp0["flags"] & 16) or "bold" in sp0["font"].lower()
                    run.font.italic = bool(sp0["flags"] & 2)
                    run.font.name = _clean_font_name(sp0["font"])
                    c = sp0["color"]
                    run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
        if progress:
            progress(k + 1, len(pages))
    prs.save(path)


# --------------------------------------------------------------------- conversions (to PDF)
def text_file_to_pdf(path, out, size=10.5, margin=54):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as fh:
                text = fh.read()
            break
        except UnicodeDecodeError:
            continue
    fk = font_kwargs(text)
    doc = fitz.open()
    pw, ph = 595.28, 841.89
    maxw = pw - 2 * margin
    lh = size * 1.3
    lines = []
    for raw in text.replace("\t", "    ").splitlines():
        if not raw.strip():
            lines.append("")
            continue
        cur = ""
        for word in re.findall(r"\S+\s*", raw):
            if text_width((cur + word).rstrip(), size, fk) <= maxw:
                cur += word
            else:
                if cur:
                    lines.append(cur.rstrip())
                while text_width(word.rstrip(), size, fk) > maxw:            # a very long word
                    cut = len(word)
                    while cut > 1 and text_width(word[:cut], size, fk) > maxw:
                        cut -= 1
                    lines.append(word[:cut])
                    word = word[cut:]
                cur = word
        lines.append(cur.rstrip())
    per_page = max(1, int((ph - 2 * margin) // lh))
    for start in range(0, max(1, len(lines)), per_page):
        pg = doc.new_page(width=pw, height=ph)
        y = margin + size
        for ln in lines[start:start + per_page]:
            if ln:
                pg.insert_text((margin, y), ln, fontsize=size, **fk)
            y += lh
    doc.save(out)
    doc.close()


MAX_PIXELS = 30_000_000          # largest bitmap Claudius PDF Editor will make from one page (protects against huge / hostile pages)
SAFE_URI_SCHEMES = ("http", "https", "mailto")


def safe_dpi(page, dpi, max_pixels=MAX_PIXELS):
    """lower the dpi when the page would otherwise need an enormous bitmap"""
    r = page.rect
    w, h = max(r.width, 1) / 72.0, max(r.height, 1) / 72.0
    return max(10, min(dpi, int((max_pixels / (w * h)) ** 0.5)))


def safe_zoom(page, zoom, max_pixels=MAX_PIXELS):
    r = page.rect
    return min(zoom, (max_pixels / (max(r.width, 1) * max(r.height, 1))) ** 0.5)


def safe_uri(uri):
    """only plain web / e-mail links may be opened; anything else (javascript:, file:, ms-msdt: ...) is refused"""
    uri = "".join(ch for ch in (uri or "") if ch >= " " and ch != "\x7f").strip()
    scheme = uri.split(":", 1)[0].lower() if ":" in uri else ""
    return uri if scheme in SAFE_URI_SCHEMES else None


def sweep_stale_temp(max_age_hours=24):
    """remove leftovers of earlier runs (print files, conversion work folders)"""
    import glob
    now = time.time()
    for pat in ("claudius_print_*.pdf", "claudius_cv_*", "claudius_lo_*"):
        for p in glob.glob(os.path.join(tempfile.gettempdir(), pat)):
            try:
                if now - os.path.getmtime(p) > max_age_hours * 3600:
                    shutil.rmtree(p, ignore_errors=True) if os.path.isdir(p) else os.remove(p)
            except OSError:
                pass


def scrub_phrases(doc, phrases):
    """redaction must not leave the removed words in the document properties, XMP data or bookmark titles"""
    phrases = [p for p in (ph.strip() for ph in phrases) if len(p) >= 3]
    if not phrases:
        return
    def clean(text):
        out = text
        for ph in phrases:
            out = re.sub(re.escape(ph), "", out, flags=re.I)
        return out
    try:
        meta = dict(doc.metadata or {})
        new = {k: clean(v) if isinstance(v, str) else v for k, v in meta.items()}
        if new != meta:
            doc.set_metadata(new)
    except Exception:
        pass
    try:
        xml = doc.get_xml_metadata() or ""
        if xml and any(ph.lower() in xml.lower() for ph in phrases):
            doc.del_xml_metadata()
    except Exception:
        pass
    try:
        toc = doc.get_toc(simple=True)
        new = [[lv, clean(t), pg] for lv, t, pg in toc]
        if new != [list(x) for x in toc]:
            doc.set_toc(new)
    except Exception:
        pass


def find_soffice():
    for n in ("soffice", "libreoffice"):
        p = shutil.which(n)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.exists(p):
            return p
    return None


_PS_OFFICE = {
    ".doc": "$a=New-Object -ComObject Word.Application;$a.Visible=$false;$d=$a.Documents.Open($env:CLAUDIUS_SRC);$d.SaveAs([ref]$env:CLAUDIUS_DST,[ref]17);$d.Close();$a.Quit()",
    ".xls": "$a=New-Object -ComObject Excel.Application;$a.Visible=$false;$d=$a.Workbooks.Open($env:CLAUDIUS_SRC);$d.ExportAsFixedFormat(0,$env:CLAUDIUS_DST);$d.Close($false);$a.Quit()",
    ".ppt": "$a=New-Object -ComObject PowerPoint.Application;$d=$a.Presentations.Open($env:CLAUDIUS_SRC,$true,$false,$false);$d.SaveAs($env:CLAUDIUS_DST,32);$d.Close();$a.Quit()",
}


def office_to_pdf(path, outdir):
    """Word / Excel / PowerPoint / RTF / OpenDocument -> PDF, through LibreOffice (preferred) or Microsoft Office."""
    base = os.path.splitext(os.path.basename(path))[0]
    dst = os.path.join(outdir, base + ".pdf")
    exe = find_soffice()
    if exe:
        prof = tempfile.mkdtemp(prefix="claudius_lo_")
        try:
            uri = "file:///" + prof.replace("\\", "/").lstrip("/")
            cmd = [exe, "-env:UserInstallation=" + uri, "--headless", "--convert-to", "pdf", "--outdir", outdir, os.path.abspath(path)]
            subprocess.run(cmd, capture_output=True, timeout=240, creationflags=NO_WINDOW)
        finally:
            shutil.rmtree(prof, ignore_errors=True)
        if os.path.exists(dst):
            return dst
    if sys.platform == "win32":
        ext = os.path.splitext(path)[1].lower()
        key = {".docx": ".doc", ".rtf": ".doc", ".odt": ".doc", ".xlsx": ".xls", ".ods": ".xls", ".pptx": ".ppt", ".odp": ".ppt"}.get(ext, ext)
        if key in _PS_OFFICE:
            # the file names travel in environment variables - they are never pasted into the script (no quoting tricks possible)
            env = dict(os.environ, CLAUDIUS_SRC=os.path.abspath(path), CLAUDIUS_DST=os.path.abspath(dst))
            subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", _PS_OFFICE[key]], capture_output=True,
                           timeout=240, creationflags=NO_WINDOW, env=env)
            if os.path.exists(dst):
                return dst
    raise RuntimeError("To turn Word, Excel or PowerPoint files into PDF you need LibreOffice (free) or "
                       "Microsoft Office installed on this computer.")




# ======================================================================================
#  part 2: markup, edit text, redact, links, forms, signatures, stamps, extract text
# ======================================================================================
EXTRA_HINTS = {
    "underline": "Underline: drag over the text you want underlined.",
    "strike": "Strikethrough: drag over the text.",
    "squiggly": "Squiggly underline: drag over the text.",
    "sticky": "Sticky note: click where the note should go.",
    "textbox": "Text box: drag a box on the page, then type. Drag the frame to move it.",
    "link": "Link: drag over the area that should become clickable.",
    "redact": "Redact: drag over content to remove it permanently (black box).",
    "edittext": "Edit text: click a line to change it, drag a line to move it. Shift = whole paragraph.",
    "extract": "Extract text: drag over an area to copy its text (uses OCR on scanned pages).",
    "formtext": "Text field: drag a box to add a fillable text field.",
    "formcheck": "Check box: click where the check box should go.",
    "date": "Date: click where today's date should be written.",
    "editimage": "Edit pictures: click a picture to replace, delete or save it - drag it to move it.",
    "formcombo": "Drop-down list: drag a box to add a drop-down field.",
    "formradio": "Radio button: click where the round button should go.",
}

RECT_TOOLS = ("underline", "strike", "squiggly", "textbox", "link", "redact", "extract", "formtext", "formcombo")
CLICK_TOOLS = ("sticky", "edittext", "date", "formcheck", "formradio", "editimage")
EXTRA_TOOLS = RECT_TOOLS + CLICK_TOOLS
HL_COLORS = [(1.0, 0.92, 0.0), (0.45, 0.95, 0.35), (1.0, 0.55, 0.8), (0.4, 0.85, 1.0)]


def _int_color(c):
    return (((c >> 16) & 255) / 255, ((c >> 8) & 255) / 255, (c & 255) / 255)


def _rot_from_dir(d):
    x, y = d
    if abs(x) > 0.9 and abs(y) < 0.2:
        return 0 if x > 0 else 180
    if abs(y) > 0.9 and abs(x) < 0.2:
        return 90 if y < 0 else 270
    return None


_DOWN = {0: (0, 1), 90: (1, 0), 180: (0, -1), 270: (-1, 0)}


class ExtraMarkup:
    """Second half of ExtraTools (merged below)."""

    # ---------------------------------------------------------------- routing from the mouse handlers
    def _canvas_pt(self, cx, cy):
        return fitz.Point((cx - self.off_x) / self.zoom, (cy - self.off_y) / self.zoom) * self.page().derotation_matrix

    def _extra_release(self, t, x0, y0, x1, y1, shift=False):
        """Called by the main window for the tools above (canvas coordinates)."""
        moved = abs(x1 - x0) > 4 and abs(y1 - y0) > 4
        if t == "editimage":
            return self._v3_release_image(x0, y0, x1, y1)
        if t in RECT_TOOLS:
            if not moved:
                return
            rect = self.to_pdf_rect(x0, y0, x1, y1)
            {"underline": lambda: self.markup_rect("underline", rect),
             "strike": lambda: self.markup_rect("strike", rect),
             "squiggly": lambda: self.markup_rect("squiggly", rect),
             "textbox": lambda: self.add_textbox(rect),
             "link": lambda: self.add_link(rect),
             "redact": lambda: self.redact_rect(rect),
             "extract": lambda: self.extract_text(rect),
             "formtext": lambda: self.add_form_field("text", rect),
             "formcombo": lambda: self.v3_form_combo(rect)}[t]()
        else:
            pt = self._canvas_pt(x1, y1)
            if t == "sticky":
                self.add_sticky(pt)
            elif t == "edittext":
                p0 = self._canvas_pt(x0, y0)
                if math.hypot(x1 - x0, y1 - y0) > 6 and self._mvl and self._mvl[1]:
                    self.move_text_line(p0, pt, block=shift)
                else:
                    self.edit_text_at(pt, block=shift)
            elif t == "date":
                self.add_date(x1, y1)
            elif t == "formcheck":
                self.add_form_field("check", fitz.Rect(pt.x - 7, pt.y - 7, pt.x + 7, pt.y + 7))
            elif t == "formradio":
                self.v3_form_radio(pt)

    # ---------------------------------------------------------------- text markup
    def words_in(self, rect):
        pg = self.page()
        out = []
        for w in pg.get_text("words"):
            r = fitz.Rect(w[:4])
            c = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
            if rect.contains(c) or (r & rect).get_area() > 0.5 * r.get_area():
                out.append(w)
        return out

    def markup_rect(self, kind, rect):
        pg = self.page()
        words = self.words_in(rect)
        lines = {}
        for w in words:
            k = (w[5], w[6])
            r = fitz.Rect(w[:4])
            lines[k] = lines[k] | r if k in lines else r
        quads = [r.quad for r in lines.values()] or [rect.quad]
        add = {"highlight": pg.add_highlight_annot, "underline": pg.add_underline_annot,
               "strike": pg.add_strikeout_annot, "squiggly": pg.add_squiggly_annot}[kind]
        col = getattr(self, "hl_color", HL_COLORS[0]) if kind == "highlight" else self.color

        def work():
            a = add(quads=quads)
            if a is not None:
                a.set_colors(stroke=col)
                a.update()
        if self._run_edit(work) and not words:
            self.status.config(text="No text found there - marked the area instead.")

    # ---------------------------------------------------------------- sticky notes, text boxes, links, date
    def add_sticky(self, pt):
        r = self.ask_text("Sticky note", "Note text:", height=6)
        if not r or not r[0].strip():
            return
        txt = r[0].strip()

        def work():
            a = self.page().add_text_annot(pt, txt, icon="Note")
            a.set_colors(stroke=self.color, fill=(1, 0.93, 0.4))
            a.update()
        self._run_edit(work, msg="Note added - hover over it in other PDF viewers to read it.")

    def add_textbox(self, rect):
        r = self.ask_text("Text box", "Text:", check="Draw a border", check_default=False, height=7)
        if not r or not r[0].strip():
            return
        txt, border = r[0], r[1]
        size0 = max(6, self.font_size.get())
        fk = font_kwargs(txt)
        pg = self.page()

        def work():
            if border:
                pg.draw_rect(rect, color=self.color, width=max(0.5, self.line_width.get() / 2))
            inner = rect + (2, 2, -2, -2) if rect.width > 8 and rect.height > 8 else rect
            s = size0
            while s >= 5:
                shape_ok = pg.insert_textbox(inner, txt, fontsize=s, color=self.color, rotate=pg.rotation,
                                             align=0, **fk)
                if shape_ok >= 0:
                    return
                s -= 1
            raise RuntimeError("The text does not fit in that box - draw a bigger box.")
        self._run_edit(work)

    def add_link(self, rect):
        d = self._dlg("Add link")
        kind = tk.StringVar(value="web")
        url = tk.StringVar(value="https://")
        pno = tk.StringVar(value=str(self.page_no + 1))
        tk.Radiobutton(d, text="Web address", variable=kind, value="web", bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=0, column=0, sticky="w")
        ttk.Entry(d, textvariable=url, width=36).grid(row=0, column=1, padx=6, pady=3)
        tk.Radiobutton(d, text="Page in this document", variable=kind, value="page", bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=1, column=0, sticky="w")
        ttk.Spinbox(d, from_=1, to=len(self.doc), width=6, textvariable=pno).grid(row=1, column=1, sticky="w", padx=6)
        res = {}

        def ok():
            res["v"] = (kind.get(), url.get().strip(), pno.get())
            d.destroy()
        self._buttons(d, "Add link", ok, 2, colspan=2)
        d.wait_window()
        if "v" not in res:
            return
        k, u, p = res["v"]
        if k == "web":
            if len(u) < 4 or u in ("https://", "http://"):
                return
            link = {"kind": fitz.LINK_URI, "from": rect, "uri": u}
        else:
            try:
                n = int(p) - 1
                assert 0 <= n < len(self.doc)
            except Exception:
                messagebox.showwarning(APP, "That page number does not exist.")
                return
            link = {"kind": fitz.LINK_GOTO, "from": rect, "page": n, "to": fitz.Point(0, 0), "zoom": 0}
        self._run_edit(lambda: self.page().insert_link(link), thumbs=[], msg="Link added.")

    def add_date(self, cx, cy):
        txt = datetime.date.today().strftime("%d/%m/%Y")
        r = self.ask_text("Date", "Text to write:", initial=txt, height=1)
        if not r or not r[0].strip():
            return
        t = r[0].strip().splitlines()[0]
        vx, vy = (cx - self.off_x) / self.zoom, (cy - self.off_y) / self.zoom
        size = self.font_size.get()
        self._run_edit(lambda: put_text(self.page(), vx, vy, t, size, self.color, font_kwargs(t)))

    # ---------------------------------------------------------------- redaction
    def redact_rect(self, rect):
        def work():
            pg = self.page()
            gone = re.findall(r"\S{4,}", pg.get_text("text", clip=rect))
            pg.add_redact_annot(rect, fill=(0, 0, 0))
            pg.apply_redactions(images=getattr(fitz, "PDF_REDACT_IMAGE_PIXELS", 2))
            scrub_phrases(self.doc, gone)
        self._run_edit(work, msg="Area redacted - the content underneath is permanently removed.")

    def dlg_redact_text(self):
        if not self.need_doc():
            return
        r = self.ask_text("Redact text", "Text to remove everywhere in the document (one phrase per line):", height=5)
        if not r or not r[0].strip():
            return
        words = [w.strip() for w in r[0].splitlines() if w.strip()]
        total = []
        for i, pg in enumerate(self.doc):
            for w in words:
                for h in pg.search_for(w):
                    total.append((i, h))
        if not total:
            messagebox.showinfo(APP, "Nothing found.")
            return
        if not messagebox.askyesno(APP, f"{len(total)} match(es) found on {len({p for p, _ in total})} page(s).\n\n"
                                        "Permanently remove them? (This can be undone until you save.)"):
            return
        self.do_redact_hits(total, words)

    def do_redact_hits(self, hits, words=()):
        pages = sorted({p for p, _ in hits})

        def work():
            for p in pages:
                pg = self.doc[p]
                for q, h in hits:
                    if q == p:
                        pg.add_redact_annot(h, fill=(0, 0, 0))
                pg.apply_redactions(images=getattr(fitz, "PDF_REDACT_IMAGE_PIXELS", 2))
            scrub_phrases(self.doc, words)
        self._run_edit(work, rebuild=True, msg=f"{len(hits)} item(s) redacted.")

    # ---------------------------------------------------------------- edit existing text
    def find_line(self, pt):
        pg = self.page()
        for b in pg.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            for l in b["lines"]:
                if (fitz.Rect(l["bbox"]) + (-2, -2, 2, 2)).contains(pt) and l["spans"]:
                    return b, l
        return None, None

    def edit_text_at(self, pt, block=False):
        b, l = self.find_line(pt)
        if not l:
            self.status.config(text="There is no text at that spot (scanned pages need OCR first).")
            return
        rot = _rot_from_dir(l["dir"])
        if rot is None:
            messagebox.showinfo(APP, "This text is slanted - editing is only supported for horizontal or vertical text.")
            return
        if block:
            lines = [l2 for l2 in b["lines"] if l2["spans"]]
            old = "\n".join("".join(s["text"] for s in l2["spans"]) for l2 in lines)
            area = fitz.Rect(b["bbox"])
        else:
            lines = [l]
            old = "".join(s["text"] for s in l["spans"])
            area = fitz.Rect(l["bbox"])
        sp = l["spans"][0]
        size = sp["size"]
        col = _int_color(sp["color"])
        bold = bool(sp["flags"] & 16) or "bold" in sp["font"].lower()
        r = self.ask_text("Edit text", "Change the text:", initial=old, height=6 if block else 3)
        if r is None or r[0].strip() == old.strip():
            return
        new = r[0].rstrip("\n")
        pg = self.page()
        try:
            import claudius_v3
            fk = claudius_v3.match_font(self.doc, pg, sp, new)
        except Exception:
            fk = font_kwargs(new, bold) if needs_unicode(new) else {"fontname": base14_name(sp["font"], sp["flags"])}

        def work():
            red = fitz.Rect(area)
            red.y0 += area.height * 0.08 if not block else 1
            red.y1 -= area.height * 0.08 if not block else 1
            red.x0 += 0.5
            red.x1 -= 0.5
            pg.add_redact_annot(red, fill=False)
            pg.apply_redactions(images=getattr(fitz, "PDF_REDACT_IMAGE_NONE", 0),
                                graphics=getattr(fitz, "PDF_REDACT_LINE_ART_NONE", 0))
            kw = dict(fk)
            if "fontbuffer" in kw:
                pg.insert_font(fontname=kw["fontname"], fontbuffer=kw.pop("fontbuffer"))
            if block:
                s = size
                while s >= 4:
                    rc = pg.insert_textbox(fitz.Rect(area.x0, area.y0, area.x1 + 2, area.y1 + 4 * size), new,
                                           fontsize=s, color=col, rotate=rot, align=0, **kw) if rot == 0 else \
                        pg.insert_textbox(area, new, fontsize=s, color=col, rotate=rot, align=0, **kw)
                    if rc >= 0:
                        return
                    s -= 0.5
                raise RuntimeError("The new text is too long for this paragraph.")
            ox, oy = sp["origin"]
            dx, dy = _DOWN[rot]
            for k, line in enumerate(new.split("\n")):
                pg.insert_text(fitz.Point(ox + dx * size * 1.2 * k, oy + dy * size * 1.2 * k), line, fontsize=size,
                               color=col, rotate=rot, **kw)
        self._run_edit(work, msg="Text changed.")

    # ---------------------------------------------------------------- extract text
    def extract_text(self, rect):
        pg = self.page()
        txt = pg.get_textbox(rect).strip()
        note = ""
        if not txt:
            td = tessdata_dir()
            if td and find_tesseract() or td:
                try:
                    langs = ocr_languages(td)
                    lang = "eng" if "eng" in langs else langs[0]
                    tp = pg.get_textpage_ocr(flags=0, language=lang, dpi=250, full=True, tessdata=td)
                    txt = pg.get_text("text", clip=rect, textpage=tp).strip()
                    note = "(recognized with OCR)"
                except Exception as ex:
                    note = f"(OCR failed: {ex})"
            else:
                note = "(no text on the page - install Tesseract to read scanned pages)"
        d = self._dlg("Extracted text")
        self._label(d, "Extracted text {}".format(note)).grid(row=0, column=0, sticky="w")
        t = tk.Text(d, width=64, height=12, wrap="word", font=(UI_FONT, 10), relief="solid", bd=1)
        t.grid(row=1, column=0, pady=6)
        t.insert("1.0", txt)

        def copy():
            self.clipboard_clear()
            self.clipboard_append(t.get("1.0", "end-1c"))
            self.status.config(text="Text copied to the clipboard.")
        f = tk.Frame(d, bg=BG)
        f.grid(row=2, column=0, sticky="e", pady=(8, 0))
        ttk.Button(f, text="Copy", command=copy).pack(side="left", padx=4)
        ttk.Button(f, text="Close", command=d.destroy).pack(side="left")
        t.focus_set()

    # ---------------------------------------------------------------- view mode: links and forms
    def _view_release(self, e):
        """Called for a click (without drag) with the Hand tool.  Returns True when it did something."""
        vc = self._view_click
        self._view_click = None
        if not vc or self.doc is None or abs(e.x - vc[0]) > 4 or abs(e.y - vc[1]) > 4:
            return False
        pt = self._canvas_pt(self.canvas.canvasx(e.x), self.canvas.canvasy(e.y))
        pg = self.page()
        for w in pg.widgets() or []:
            if w.rect.contains(pt):
                self.fill_widget(w)
                return True
        for lk in pg.get_links():
            if fitz.Rect(lk["from"]).contains(pt):
                if lk.get("uri"):
                    uri = safe_uri(lk["uri"])
                    if uri is None:
                        messagebox.showwarning(APP, _tr("Blocked a link that is not a web or e-mail address:\n{}").replace("{}", lk["uri"][:200]))
                    elif messagebox.askyesno(APP, _tr("Open this link?\n\n{}").replace("{}", uri[:300])):
                        webbrowser.open(uri)
                elif lk.get("page", -1) >= 0:
                    self.goto(lk["page"])
                return True
        return False

    def fill_widget(self, w):
        xref, t = w.xref, w.field_type
        name = w.field_name or "field"

        def find():
            self._pg_keep = self.page()                 # keep the page object alive while the widget is used
            for x in self._pg_keep.widgets() or []:
                if x.xref == xref:
                    return x

        if t == fitz.PDF_WIDGET_TYPE_TEXT:
            multi = bool(w.field_flags & 4096)
            if multi:
                r = self.ask_text("Fill field", name, initial=w.field_value or "", height=6)
                val = r[0] if r else None
            else:
                val = simpledialog.askstring(APP, name, initialvalue=w.field_value or "", parent=self)
            if val is None:
                return

            def work():
                x = find()
                x.field_value = val
                x.update()
        elif t in (fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.PDF_WIDGET_TYPE_RADIOBUTTON):
            on = w.on_state()
            cur = w.field_value
            is_on = cur not in ("Off", False, "", None, 0)

            def work():
                x = find()
                if t == fitz.PDF_WIDGET_TYPE_RADIOBUTTON:
                    x.field_value = on
                else:
                    x.field_value = False if is_on else on
                x.update()
        elif t in (fitz.PDF_WIDGET_TYPE_COMBOBOX, fitz.PDF_WIDGET_TYPE_LISTBOX):
            vals = [v if isinstance(v, str) else v[-1] for v in (w.choice_values or [])]
            d = self._dlg(name)
            lb = tk.Listbox(d, height=min(10, max(3, len(vals))), width=36, exportselection=False)
            for v in vals:
                lb.insert("end", v)
            lb.grid(row=0, column=0)
            res = {}

            def ok():
                if lb.curselection():
                    res["v"] = vals[lb.curselection()[0]]
                d.destroy()
            self._buttons(d, "OK", ok, 1, colspan=1)
            d.wait_window()
            if "v" not in res:
                return

            def work():
                x = find()
                x.field_value = res["v"]
                x.update()
        else:
            return
        self._run_edit(work, msg=f"Filled: {name}")

    def add_form_field(self, kind, rect):
        pg = self.page()
        n = sum(1 for _ in (pg.widgets() or [])) + 1
        w = fitz.Widget()
        if kind == "text":
            w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            w.field_name = f"text_{self.page_no + 1}_{n}"
            w.text_fontsize = 11
            w.fill_color = (0.93, 0.96, 1)
        else:
            w.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX
            w.field_name = f"check_{self.page_no + 1}_{n}"
            w.field_value = False
        w.rect = rect
        w.border_color = (0.2, 0.2, 0.2)
        w.border_width = 1
        w.text_color = (0, 0, 0)
        self._run_edit(lambda: pg.add_widget(w), msg="Form field added - click it with the Hand tool to fill it.")

    def toggle_fields(self):
        self.show_fields.set(not self.show_fields.get())
        self.render()

    # ---------------------------------------------------------------- right click on an annotation
    def annot_menu(self, e):
        if self.doc is None:
            return False
        pt = self._canvas_pt(self.canvas.canvasx(e.x), self.canvas.canvasy(e.y))
        pg = self.page()
        hit = None
        for a in pg.annots() or []:
            if (a.rect + (-2, -2, 2, 2)).contains(pt):
                hit = a
        link = None
        if hit is None:
            for lk in pg.get_links():
                if fitz.Rect(lk["from"]).contains(pt):
                    link = lk
        pic = self._v3_image_at(pt) if hasattr(self, "_v3_image_at") else None
        if hit is None and link is None and pic is None:
            return False
        m = tk.Menu(self, tearoff=0)
        if hit is None and link is None:
            self.v3_picture_items(m, fitz.Rect(pic["bbox"]), pic["xref"])
            m.tk_popup(e.x_root, e.y_root)
            return True
        if hit is not None:
            xref = hit.xref
            if hit.type[0] == fitz.PDF_ANNOT_TEXT or hit.info.get("content"):
                m.add_command(label="Edit note…", command=lambda: self.edit_annot_text(xref))
            m.add_command(label="Delete this annotation", command=lambda: self.delete_annot(xref))
        else:
            m.add_command(label="Delete this link", command=lambda: self.delete_link(link))
        if pic is not None:
            m.add_separator()
            self.v3_picture_items(m, fitz.Rect(pic["bbox"]), pic["xref"])
        m.tk_popup(e.x_root, e.y_root)
        return True

    def _annot_by_xref(self, xref):
        self._pg_keep = self.page()
        for a in self._pg_keep.annots() or []:
            if a.xref == xref:
                return a

    def delete_annot(self, xref):
        self._run_edit(lambda: self.page().delete_annot(self._annot_by_xref(xref)), msg="Annotation deleted.")

    def edit_annot_text(self, xref):
        a = self._annot_by_xref(xref)
        r = self.ask_text("Edit note", "Note text:", initial=a.info.get("content", ""), height=6)
        if r is None:
            return

        def work():
            b = self._annot_by_xref(xref)
            b.set_info(content=r[0])
            b.update()
        self._run_edit(work, thumbs=[])

    def delete_link(self, link):
        self._run_edit(lambda: self.page().delete_link(link), thumbs=[], msg="Link deleted.")

    # ---------------------------------------------------------------- signatures / initials
    def sig_dir(self):
        p = os.path.join(CONFIG_DIR, "signatures")
        os.makedirs(p, exist_ok=True)
        return p

    def saved_signatures(self, kind):
        try:
            return sorted(os.path.join(self.sig_dir(), f) for f in os.listdir(self.sig_dir())
                          if f.startswith(kind + "_") and f.endswith(".png"))
        except OSError:
            return []

    def show_sign_menu(self, kind):
        if not self.need_doc():
            return
        word = "signature" if kind == "sign" else "initials"
        m = tk.Menu(self, tearoff=0)
        saved = self.saved_signatures(kind)
        self._menu_imgs = []
        for p in saved:
            try:
                im = Image.open(p).convert("RGBA")
                im.thumbnail((140, 40), LANCZOS)
                bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
                bg.alpha_composite(im)
                ph = ImageTk.PhotoImage(bg)
                self._menu_imgs.append(ph)
                m.add_command(image=ph, command=lambda p=p: self.insert_saved(p, word))
            except Exception:
                continue
        if saved:
            m.add_separator()
        m.add_command(label=f"Create new {word}…", command=lambda: self.dlg_new_signature(kind))
        if kind == "sign":
            self._v3_block_menu(m)
        if saved:
            dm = tk.Menu(m, tearoff=0)
            for i, p in enumerate(saved):
                dm.add_command(label=f"Delete saved {word} {i + 1}", command=lambda p=p: self._del_sig(p))
            m.add_cascade(label=f"Delete a saved {word}", menu=dm)
        m.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

    def _del_sig(self, p):
        try:
            os.remove(p)
        except OSError:
            pass

    def insert_saved(self, path, name):
        try:
            self.place_image(Image.open(path).convert("RGBA"), name)
            self.status.config(text=f"Drag the {name} into place, then press Enter (or Apply) to fix it on the page.")
        except Exception as ex:
            messagebox.showerror(APP, f"Could not load that {name}:\n{ex}")

    def dlg_new_signature(self, kind):
        word = "signature" if kind == "sign" else "initials"
        d = self._dlg(f"Create {word}")
        nb = ttk.Notebook(d)
        nb.grid(row=0, column=0, columnspan=2)
        state = {"img": None, "color": (0.05, 0.1, 0.45)}
        ink_row = tk.Frame(d, bg=BG)
        ink_row.grid(row=1, column=0, sticky="w", pady=6)
        self._label(ink_row, "Ink color").pack(side="left")
        cb, ch = self._color_button(ink_row, state["color"])
        cb.pack(side="left", padx=6)

        # --- type
        tf = tk.Frame(nb, bg=BG, padx=10, pady=10)
        nb.add(tf, text="  Type  ")
        tv = tk.StringVar(value="")
        te = ttk.Entry(tf, textvariable=tv, width=40, font=(UI_FONT, 12))
        te.grid(row=0, column=0, columnspan=2, pady=4)
        fonts = signature_fonts()
        fv = tk.StringVar(value=fonts[0][0] if fonts else "")
        ttk.Combobox(tf, textvariable=fv, values=[f[0] for f in fonts], state="readonly", width=18).grid(row=1, column=0, sticky="w")
        prev = tk.Label(tf, bg="white", fg="#6b7480", width=52, height=4, relief="solid", bd=1)
        prev.grid(row=2, column=0, columnspan=2, pady=8)

        def type_img():
            txt = tv.get().strip()
            if not txt:
                return None
            path = dict(fonts).get(fv.get()) or find_font("italic")
            return make_text_image(txt, path, 140, ch["rgb"])

        # --- draw
        df = tk.Frame(nb, bg=BG, padx=10, pady=10)
        nb.add(df, text="  Draw  ")
        cv = tk.Canvas(df, width=480, height=170, bg="white", relief="solid", bd=1, cursor="pencil")
        cv.grid(row=0, column=0, columnspan=2)
        strokes = []

        def dpress(e):
            strokes.append([(e.x, e.y)])

        def dmove(e):
            if strokes:
                x0, y0 = strokes[-1][-1]
                cv.create_line(x0, y0, e.x, e.y, width=3, fill=self._hex(ch["rgb"]), capstyle="round", smooth=True)
                strokes[-1].append((e.x, e.y))
        cv.bind("<ButtonPress-1>", dpress)
        cv.bind("<B1-Motion>", dmove)

        def clear():
            cv.delete("all")
            strokes.clear()
        ttk.Button(df, text="Clear", command=clear).grid(row=1, column=0, sticky="w", pady=4)
        self._label(df, "Draw with the mouse (or a pen / touch screen).", fg=TH.TERTIARY).grid(row=1, column=1, sticky="e")

        def draw_img():
            pts = [s for s in strokes if len(s) > 0]
            if not pts:
                return None
            k = 3
            img = Image.new("RGBA", (480 * k, 170 * k), (0, 0, 0, 0))
            dr = ImageDraw.Draw(img)
            col = tuple(int(c * 255) for c in ch["rgb"]) + (255,)
            for s in pts:
                p = [(x * k, y * k) for x, y in s]
                if len(p) == 1:
                    dr.ellipse((p[0][0] - 5, p[0][1] - 5, p[0][0] + 5, p[0][1] + 5), fill=col)
                else:
                    dr.line(p, fill=col, width=9, joint="curve")
                    for q in (p[0], p[-1]):
                        dr.ellipse((q[0] - 4, q[1] - 4, q[0] + 4, q[1] + 4), fill=col)
            bb = img.getchannel("A").getbbox()
            if not bb:
                return None
            pad = 12
            return img.crop((max(0, bb[0] - pad), max(0, bb[1] - pad), bb[2] + pad, bb[3] + pad))

        # --- image
        imf = tk.Frame(nb, bg=BG, padx=10, pady=10)
        nb.add(imf, text="  Image  ")
        ipath = tk.StringVar(value="")
        rw = tk.BooleanVar(value=True)
        ttk.Button(imf, text="Choose picture…", command=lambda: ipath.set(
            filedialog.askopenfilename(parent=d, filetypes=[("Images", "*.png *.jpg *.jpeg *.bmp *.gif")]) or ipath.get())
        ).grid(row=0, column=0, sticky="w")
        self._label(imf, "").grid(row=0, column=1)
        il = self._label(imf, "Pick a photo or scan of your signature on white paper.", fg=TH.TERTIARY)
        il.grid(row=1, column=0, columnspan=2, sticky="w", pady=4)
        tk.Checkbutton(imf, text="Remove the white background", variable=rw, bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=2, column=0, columnspan=2, sticky="w")
        ipath.trace_add("write", lambda *a: il.config(text=os.path.basename(ipath.get())))

        def image_img():
            if not ipath.get():
                return None
            im = Image.open(ipath.get())
            im = ImageOps.exif_transpose(im).convert("RGBA")
            return remove_white(im) if rw.get() else im

        def refresh_prev(*_):
            try:
                im = type_img()
            except Exception:
                im = None
            if im is None:
                prev.config(image="", text="Type your name above")
                return
            t = im.copy()
            t.thumbnail((420, 70), LANCZOS)
            bg = Image.new("RGBA", t.size, (255, 255, 255, 255))
            bg.alpha_composite(t)
            ph = ImageTk.PhotoImage(bg)
            prev.image = ph
            prev.config(image=ph, text="", width=ph.width(), height=ph.height())
        tv.trace_add("write", refresh_prev)
        fv.trace_add("write", refresh_prev)
        refresh_prev()
        keep = tk.BooleanVar(value=True)
        tk.Checkbutton(d, text=f"Save for next time", variable=keep, bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=1, column=1, sticky="e")

        def ok():
            try:
                tab = nb.index(nb.select())
                img = [type_img, draw_img, image_img][tab]()
            except Exception as ex:
                messagebox.showerror(APP, f"Could not create it:\n{ex}", parent=d)
                return
            if img is None:
                messagebox.showinfo(APP, "Nothing to use yet.", parent=d)
                return
            if keep.get():
                try:
                    img.save(os.path.join(self.sig_dir(), f"{kind}_{int(time.time())}.png"))
                except OSError:
                    pass
            d.destroy()
            self.insert_saved_image(img, word)
        self._buttons(d, "Use", ok, 2)

    def insert_saved_image(self, img, word):
        self.place_image(img.convert("RGBA"), word)
        self.status.config(text=f"Drag the {word} into place, then press Enter (or Apply) to fix it on the page.")

    # ---------------------------------------------------------------- stamps
    def show_stamp_menu(self):
        if not self.need_doc():
            return
        m = tk.Menu(self, tearoff=0)
        self._stamp_date = getattr(self, "_stamp_date", tk.BooleanVar(value=False))
        for text, col in STAMPS:
            m.add_command(label=text.title(), foreground=self._hex(col), command=lambda t=text, c=col: self.add_stamp(t, c))
        m.add_separator()
        m.add_command(label="Custom stamp…", command=self.add_custom_stamp)
        m.add_checkbutton(label="Add today's date to the stamp", variable=self._stamp_date)
        m.tk_popup(self.winfo_pointerx(), self.winfo_pointery())

    def add_stamp(self, text, col):
        sub = datetime.date.today().strftime("%d %b %Y") if self._stamp_date.get() else None
        try:
            from claudius_i18n import tr as _tr
            text = _tr(text.title()).upper()                      # the stamp itself speaks the interface language
        except Exception:
            pass
        self.place_image(make_stamp_image(text, col, sub), "stamp")
        self.status.config(text="Drag the stamp into place, then press Enter (or Apply).")

    def add_custom_stamp(self):
        r = self.ask_text("Custom stamp", "Stamp text:", height=1)
        if r and r[0].strip():
            self.add_stamp(r[0].strip().splitlines()[0].upper(), self.color)


# ======================================================================================
#  part 3: convert, create PDF, OCR, find bar, start screen
# ======================================================================================
CONVERT_KINDS = {            # key: (title, extension, needs-package)
    "images": ("PDF to images", "", None),
    "word": ("PDF to Word", ".docx", "python-docx"),
    "excel": ("PDF to Excel", ".xlsx", "openpyxl"),
    "pptx": ("PDF to PowerPoint", ".pptx", "python-pptx"),
    "text": ("PDF to text", ".txt", None),
    "html": ("PDF to HTML", ".html", None),
    "xml": ("PDF to XML", ".xml", None),
}


class ExtraConvert:
    """Third part of ExtraTools (merged below)."""

    # ---------------------------------------------------------------- PDF -> other formats
    def dlg_convert(self, kind):
        title, ext, pkg = CONVERT_KINDS[kind]
        files = []
        d = self._dlg(title)
        use_cur = tk.BooleanVar(value=self.doc is not None)
        r = 0
        if self.doc is not None:
            tk.Radiobutton(d, text="The open document", variable=use_cur, value=True, bg=BG, activebackground=BG,
                           font=(UI_FONT, 9)).grid(row=r, column=0, columnspan=4, sticky="w")
            r += 1
        tk.Radiobutton(d, text="Other PDF files (batch)", variable=use_cur, value=False, bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=r, column=0, columnspan=4, sticky="w")
        r += 1
        lb = tk.Listbox(d, height=4, width=52)
        lb.grid(row=r, column=0, columnspan=3, sticky="we", pady=4)

        def add():
            ps = filedialog.askopenfilenames(parent=d, filetypes=[("PDF files", "*.pdf")])
            for p in ps:
                if p not in files:
                    files.append(p)
                    lb.insert("end", p)
            if ps:
                use_cur.set(False)

        def rem():
            for i in reversed(lb.curselection()):
                lb.delete(i)
                files.pop(i)
        bf = tk.Frame(d, bg=BG)
        bf.grid(row=r, column=3, sticky="n", padx=4, pady=4)
        ttk.Button(bf, text="Add…", command=add, width=8).pack(pady=1)
        ttk.Button(bf, text="Remove", command=rem, width=8).pack(pady=1)
        r += 1
        pv = tk.StringVar(value="all")
        self._label(d, "Pages (open document)").grid(row=r, column=0, sticky="w", pady=3)
        ttk.Entry(d, textvariable=pv, width=16).grid(row=r, column=1, sticky="w")
        r += 1
        fmt = tk.StringVar(value="PNG")
        dpi = tk.IntVar(value=150)
        mode = tk.StringVar(value="Pictures of the pages")
        if kind == "images":
            self._label(d, "Format").grid(row=r, column=0, sticky="w", pady=3)
            ttk.Combobox(d, textvariable=fmt, values=["PNG", "JPG", "TIFF", "BMP"], state="readonly", width=8).grid(row=r, column=1, sticky="w")
            self._label(d, "Resolution (dpi)").grid(row=r, column=2, sticky="e")
            ttk.Spinbox(d, from_=72, to=600, increment=50, width=6, textvariable=dpi).grid(row=r, column=3, sticky="w")
            r += 1
        if kind == "pptx":
            self._label(d, "Slides contain").grid(row=r, column=0, sticky="w", pady=3)
            ttk.Combobox(d, textvariable=mode, values=["Pictures of the pages", "Editable text"], state="readonly",
                         width=22).grid(row=r, column=1, columnspan=2, sticky="w")
            r += 1
        out = tk.StringVar(value="")
        self._label(d, "Save to folder").grid(row=r, column=0, sticky="w", pady=3)
        ttk.Entry(d, textvariable=out, width=36).grid(row=r, column=1, columnspan=2, sticky="we")
        ttk.Button(d, text="Browse…", width=8, command=lambda: out.set(
            filedialog.askdirectory(parent=d) or out.get())).grid(row=r, column=3)
        r += 1
        self._label(d, "(empty = next to the original file)", fg=TH.TERTIARY).grid(row=r, column=1, columnspan=3, sticky="w")
        r += 1

        def go():
            srcs = []
            if use_cur.get() and self.doc is not None:
                self.commit_floats()
                pages = self._pages_or_warn(pv.get(), d)
                if pages is None:
                    return
                srcs.append((self.doc, pages, self.path))
            else:
                if not files:
                    messagebox.showinfo(APP, "Add at least one PDF file.", parent=d)
                    return
                for p in files:
                    doc = self._open_pdf_file(p, d)
                    if doc is None:
                        return
                    srcs.append((doc, list(range(len(doc))), p))
            folder = out.get().strip()
            d.destroy()
            self.do_convert(kind, srcs, folder, fmt.get(), dpi.get(), "text" if mode.get().startswith("Editable") else "picture")
        self._buttons(d, "Convert", go, r)

    def do_convert(self, kind, srcs, folder="", fmt="PNG", dpi=150, pptx_mode="picture"):
        """srcs: list of (doc, [page indexes], original_path_or_None).  Returns the list of created paths."""
        title, ext, pkg = CONVERT_KINDS[kind]
        made = []
        total = sum(len(p) for _, p, _ in srcs)
        upd, close = self._progress(title, total)
        done_before = [0]
        errors = []
        try:
            for doc, pages, src in srcs:
                base = safe_name(os.path.splitext(os.path.basename(src))[0]) if src else "document"
                outdir = folder or (os.path.dirname(src) if src else os.path.expanduser("~"))
                if kind == "images" and len(pages) > 1:
                    pass
                os.makedirs(outdir, exist_ok=True)

                def prog(i, n, off=done_before[0]):
                    upd(off + i, total)
                try:
                    if kind == "images":
                        made += pdf_to_images(doc, pages, fmt, dpi, outdir, base, prog)
                    else:
                        path = self._unique(os.path.join(outdir, base + ext))
                        if kind == "word":
                            pdf_to_docx(doc, pages, path, prog)
                        elif kind == "excel":
                            pdf_to_xlsx(doc, pages, path, prog)
                        elif kind == "pptx":
                            pdf_to_pptx(doc, pages, path, pptx_mode, 150, prog)
                        elif kind == "text":
                            pdf_to_text(doc, pages, path)
                        elif kind == "html":
                            pdf_to_html(doc, pages, path)
                        elif kind == "xml":
                            pdf_to_xml(doc, pages, path)
                        made.append(path)
                except ImportError:
                    close()
                    messagebox.showerror(APP, f"This conversion needs the '{pkg}' package.\n\nInstall it with:\n"
                                              f"    pip install {pkg}")
                    return made
                except Exception as ex:
                    errors.append(f"{os.path.basename(src or 'document')}: {ex}")
                done_before[0] += len(pages)
                if doc is not self.doc:
                    doc.close()
        finally:
            close()
        if errors:
            messagebox.showwarning(APP, "Some files could not be converted:\n\n{}".format("\n".join(errors)))
        if made:
            where = os.path.dirname(made[0])
            self.status.config(text=f"Created {len(made)} file(s) in {where}")
            if messagebox.askyesno(APP, f"Done - {len(made)} file(s) created.\n\nShow them in the folder?"):
                self.reveal(made[0] if len(made) == 1 else where)
        return made

    @staticmethod
    def _unique(path):
        if not os.path.exists(path):
            return path
        b, e = os.path.splitext(path)
        i = 2
        while os.path.exists(f"{b} ({i}){e}"):
            i += 1
        return f"{b} ({i}){e}"

    # ---------------------------------------------------------------- other formats -> PDF
    def file_to_pdfdoc(self, path, workdir):
        """Open any supported file as a fitz PDF document (images, text, Office)."""
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            return self._open_pdf_file(path)
        if ext in IMG_EXT:
            im = ImageOps.exif_transpose(Image.open(path))
            im = im.convert("RGB") if im.mode not in ("RGB", "L") else im
            w, h = im.width * 72 / 150, im.height * 72 / 150
            k = min(1.0, 842 / max(w, h)) if max(w, h) > 842 else 1.0
            doc = fitz.open()
            pg = doc.new_page(width=max(50, w * k), height=max(50, h * k))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=92)
            pg.insert_image(pg.rect, stream=buf.getvalue())
            return doc
        if ext in (".txt", ".log", ".csv", ".md"):
            out = os.path.join(workdir, "t.pdf")
            text_file_to_pdf(path, out)
            return fitz.open(out)
        if ext in OFFICE_EXT:
            return fitz.open(office_to_pdf(path, workdir))
        raise ValueError(f"Unsupported file type: {ext}")

    def build_pdf_from_files(self, paths):
        """Combine files into one new document (in memory)."""
        work = tempfile.mkdtemp(prefix="claudius_cv_")
        out = fitz.open()
        try:
            for p in paths:
                sub = self.file_to_pdfdoc(p, work)
                if sub is None:
                    continue
                out.insert_pdf(sub)
                sub.close()
        finally:
            shutil.rmtree(work, ignore_errors=True)
        return out

    def dlg_create_pdf(self):
        paths = filedialog.askopenfilenames(
            title="Choose files to turn into a PDF",
            filetypes=[("Supported files", "*.png *.jpg *.jpeg *.bmp *.gif *.tif *.tiff *.webp *.txt *.doc *.docx *.xls *.xlsx *.ppt *.pptx *.rtf *.odt *.pdf"),
                       ("All files", "*.*")])
        if not paths:
            return
        self._create_from(list(paths))

    def _create_from(self, paths):
        upd, close = self._progress("Creating PDF", len(paths))
        try:
            doc = self.build_pdf_from_files(paths)
        except Exception as ex:
            close()
            messagebox.showerror(APP, f"Could not create the PDF:\n{ex}")
            return
        close()
        if len(doc) == 0:
            messagebox.showerror(APP, "Nothing could be converted.")
            return
        self._adopt_doc(doc, None, modified=True)
        self.status.config(text=f"Created a {len(doc)}-page PDF from {len(paths)} file(s). Use Save As to keep it.")

    def open_non_pdf(self, path):
        self._create_from([path])

    # ---------------------------------------------------------------- OCR
    def dlg_ocr(self):
        if not self.need_doc():
            return
        td = tessdata_dir()
        langs = ocr_languages(td) if td else []
        if not td or not langs:
            messagebox.showinfo(APP, "OCR needs the free Tesseract program.\n\n"
                                     "1. Install Tesseract (Windows installer from github.com/UB-Mannheim/tesseract/wiki).\n"
                                     "2. Restart Claudius PDF Editor.\n\nThen choose OCR again.")
            return
        d = self._dlg("Make scanned pages searchable (OCR)")
        v = self._range_row(d, 0, "all")
        self._label(d, "Language").grid(row=2, column=0, sticky="w", pady=3)
        lv = tk.StringVar(value="eng" if "eng" in langs else langs[0])
        ttk.Combobox(d, textvariable=lv, values=langs, state="readonly", width=10).grid(row=2, column=1, sticky="w")
        self._label(d, "Resolution (dpi)").grid(row=3, column=0, sticky="w", pady=3)
        dv = tk.IntVar(value=250)
        ttk.Spinbox(d, from_=150, to=400, increment=50, width=6, textvariable=dv).grid(row=3, column=1, sticky="w")
        skip = tk.BooleanVar(value=True)
        tk.Checkbutton(d, text="Skip pages that already contain text", variable=skip, bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=4, column=0, columnspan=4, sticky="w")

        def go():
            pages = self._pages_or_warn(v.get(), d)
            if pages is None:
                return
            lang, dpi, sk = lv.get(), dv.get(), skip.get()
            d.destroy()
            self.do_ocr(pages, lang, td, dpi, sk)
        self._buttons(d, "Start OCR", go, 5)

    def do_ocr(self, pages, lang, tessdata, dpi=250, skip=True):
        self.commit_floats()
        upd, close = self._progress("Recognizing text…", len(self.doc))
        try:
            new = ocr_document(self.doc, pages, lang, tessdata, dpi, skip, upd)
        except Exception as ex:
            close()
            messagebox.showerror(APP, f"OCR failed:\n{ex}")
            return
        close()
        self._push_snapshot()
        self.doc = new
        self._rebuild()
        self.status.config(text=f"OCR finished - {getattr(new, '_ocr_pages', 0)} page(s) are now searchable.")

    # ---------------------------------------------------------------- find bar
    def show_find(self):
        if not self.need_doc():
            return
        if self.find_frame is None:
            f = tk.Frame(self.canvas.master, bg=BG, highlightbackground=LINE_COL, highlightthickness=1)
            tk.Label(f, text="Find", bg=BG, font=(UI_FONT, 9)).pack(side="left", padx=(10, 4), pady=4)
            e = ttk.Entry(f, textvariable=self.find_var, width=30)
            e.pack(side="left")
            e.bind("<Return>", lambda ev: self.find_step(-1 if ev.state & 1 else 1))
            e.bind("<Escape>", lambda ev: self.hide_find())
            self.find_entry = e
            ttk.Button(f, text="▲", width=3, command=lambda: self.find_step(-1)).pack(side="left", padx=(6, 0))
            ttk.Button(f, text="▼", width=3, command=lambda: self.find_step(1)).pack(side="left")
            self.find_case = tk.BooleanVar(value=False)
            tk.Checkbutton(f, text="Match case", variable=self.find_case, bg=BG, activebackground=BG,
                           font=(UI_FONT, 8), command=lambda: self.find_run(keep=True)).pack(side="left", padx=8)
            tk.Label(f, textvariable=self.find_info, bg=BG, fg=TH.SECOND, font=(UI_FONT, 9)).pack(side="left", padx=6)
            x = tk.Label(f, text="✕", bg=BG, fg=TH.SECOND, cursor="hand2", padx=8)
            x.pack(side="right")
            x.bind("<Button-1>", lambda ev: self.hide_find())
            self.find_frame = f
        self.find_frame.pack(side="top", fill="x", before=self.canvas)
        self.find_entry.focus_set()
        self.find_entry.selection_range(0, "end")

    def hide_find(self):
        if self.find_frame is not None:
            self.find_frame.pack_forget()
        self.find_hits = []
        self.find_idx = -1
        self.render()
        self.canvas.focus_set()

    def find_run(self, keep=False):
        """Search the whole document for the text in the box."""
        q = self.find_var.get()
        self.find_hits = []
        if not q or not self.doc:
            self.find_info.set("")
            self.render()
            return
        flags = fitz.TEXT_DEHYPHENATE if False else 0
        for i, pg in enumerate(self.doc):
            for r in pg.search_for(q, flags=flags):
                if self.find_case.get():
                    if q not in pg.get_textbox(r):
                        continue
                self.find_hits.append((i, r))
        if not self.find_hits:
            self.find_idx = -1
            self.find_info.set(_tr("Not found"))
            self.render()
            return
        # first hit on or after the current page
        idx = next((k for k, (p, _) in enumerate(self.find_hits) if p >= self.page_no), 0)
        self.find_idx = idx
        self._find_show()

    def find_step(self, step):
        q = self.find_var.get()
        if not q:
            return
        if not self.find_hits or getattr(self, "_find_q", None) != q:
            self._find_q = q
            self.find_run()
            return
        self.find_idx = (self.find_idx + step) % len(self.find_hits)
        self._find_show()

    def _find_show(self):
        p, r = self.find_hits[self.find_idx]
        self.find_info.set(_tr("{} of {}").replace("{}", str(self.find_idx + 1), 1).replace("{}", str(len(self.find_hits)), 1))
        if p != self.page_no:
            self.goto(p)
        else:
            self.render()
        rr = r * self.page().rotation_matrix
        top, bottom = self.canvas.yview()
        sr = self.canvas.cget("scrollregion").split()
        if len(sr) == 4:
            H = float(sr[3])
            y = self.off_y + rr.y0 * self.zoom
            if y < top * H or y > bottom * H:
                self.canvas.yview_moveto(max(0, (y - 80) / H))

    def find_text(self):                      # kept for the old name used by the ribbon / Ctrl+F
        self.show_find()

    # ---------------------------------------------------------------- drawn after every page render
    def _after_render(self):
        if self.doc is None:
            return
        pg = self.page()
        if self.show_fields.get():
            for w in pg.widgets() or []:
                rr = w.rect * pg.rotation_matrix
                self.canvas.create_rectangle(self.off_x + rr.x0 * self.zoom, self.off_y + rr.y0 * self.zoom,
                                             self.off_x + rr.x1 * self.zoom, self.off_y + rr.y1 * self.zoom,
                                             outline="#2b6cb8", fill="#9cc3f2", stipple="gray25", tags="pg")
        if self.find_hits:
            m = pg.rotation_matrix
            for k, (p, r) in enumerate(self.find_hits):
                if p != self.page_no:
                    continue
                rr = r * m
                cur = (k == self.find_idx)
                self.canvas.create_rectangle(self.off_x + rr.x0 * self.zoom, self.off_y + rr.y0 * self.zoom,
                                             self.off_x + rr.x1 * self.zoom, self.off_y + rr.y1 * self.zoom,
                                             outline="#ff7a00" if cur else "#e6b800", width=2 if cur else 1,
                                             fill="#ff9d3a" if cur else "#ffe14d", stipple="gray50", tags="pg")

    # ---------------------------------------------------------------- start screen
    def _draw_start_screen(self):
        c = self.canvas
        w, h = max(c.winfo_width(), 600), max(c.winfo_height(), 400)
        cx = w // 2
        top = max(S(30), h // 2 - S(240))
        c.create_image(cx, top + S(60), image=self.logo128)
        c.create_text(cx, top + S(150), text="Claudius PDF Editor", font=(UI_FONT, 26, "bold"), fill=TH.LABEL)
        c.create_text(cx, top + S(184), text="Open a PDF to start  (Ctrl + O)",
                      font=(UI_FONT, 10), fill=TH.SECOND)
        actions = [("📂  Open PDF", self.open_dialog), ("➕  Create PDF from files", self.dlg_create_pdf),
                   ("🖼  Photos to PDF", self.photos_to_pages)]
        try:
            from claudius_i18n import tr as _tr
        except Exception:
            _tr = lambda s: s
        _f = tk.font.Font(family=UI_FONT, size=10)
        bw, gap = max(S(210), max(_f.measure(_tr(t)) for t, _c in actions) + S(44)), S(14)
        x0 = cx - (len(actions) * bw + (len(actions) - 1) * gap) // 2
        for i, (t, cmd) in enumerate(actions):
            x = x0 + i * (bw + gap)
            tag = f"sb{i}"
            round_rect(c, x, top + S(214), x + bw, top + S(260), S(14), fill=TH.CARD, outline=TH.LINE_SOFT, tags=tag)
            c.create_text(x + bw // 2, top + S(237), text=t, font=(UI_FONT, 10, "bold"), fill=ACCENT, tags=tag)
            c.tag_bind(tag, "<Button-1>", lambda e, f=cmd: f())
            c.tag_bind(tag, "<Enter>", lambda e, t=tag: (c.itemconfig(t + "&&polygon", fill=TH.TINT), c.config(cursor="hand2")))
            c.tag_bind(tag, "<Leave>", lambda e, t=tag: (c.itemconfig(t + "&&polygon", fill=TH.CARD), c.config(cursor="")))
        rec = self.load_recents()
        if rec:
            c.create_text(cx - S(300), top + S(300), text="Recent files", font=(UI_FONT, 10, "bold"), fill=TH.SECOND, anchor="w")
            for i, p in enumerate(rec[:6]):
                y = top + S(328) + i * S(26)
                tag = f"rc{i}"
                c.create_text(cx - S(300), y, text=os.path.basename(p), font=(UI_FONT, 10), fill=ACCENT, anchor="w", tags=tag)
                c.create_text(cx + S(40), y, text=os.path.dirname(p)[-40:], font=(UI_FONT, 8), fill=TH.TERTIARY, anchor="w", tags=tag)
                c.tag_bind(tag, "<Button-1>", lambda e, p=p: self._open_recent(p))
                c.tag_bind(tag, "<Enter>", lambda e: c.config(cursor="hand2"))
                c.tag_bind(tag, "<Leave>", lambda e: c.config(cursor=""))
        c.config(scrollregion=(0, 0, w, h))


# ======================================================================================
#  the mixin that is added to the Claudius PDF Editor main window
# ======================================================================================
# ======================================================================================
#  Windows 11 integration (every part is a harmless no-op on other systems)
# ======================================================================================
def colorref(hex_color):
    """'#2b6cb8' -> Windows COLORREF (0x00BBGGRR)"""
    h = hex_color.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return r | (g << 8) | (b << 16)


def parse_geometry(geo):
    m = re.fullmatch(r"(\d+)x(\d+)([+-]-?\d+)([+-]-?\d+)", geo or "")
    return tuple(int(v.replace("+-", "-")) for v in m.groups()) if m else None


try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    _DnDBase = TkinterDnD.DnDWrapper
except Exception:                     # optional package - without it drag & drop is simply off
    TkinterDnD = DND_FILES = None

    class _DnDBase:
        pass


class Win11Tools(_DnDBase):
    """Taskbar identity, matching title bar, smooth touchpad scrolling, drag & drop, window memory,
    'Open with' registration."""

    def _win11_init(self):
        self._wheel_acc = 0.0
        self._dnd_ok = False
        self.win_polish()
        self.enable_drop()

    # ---- title bar in the program colour (Windows 11 22000+; ignored by older Windows)
    def win_polish(self):
        if sys.platform != "win32":
            return
        try:
            import ctypes
            self.update_idletasks()
            hwnd = ctypes.windll.user32.GetParent(self.winfo_id()) or self.winfo_id()
            dwm = ctypes.windll.dwmapi

            def put(attr, value):
                v = ctypes.c_int(value)
                dwm.DwmSetWindowAttribute(hwnd, attr, ctypes.byref(v), ctypes.sizeof(v))
            put(20, 1 if TH.IS_DARK else 0)    # DWMWA_USE_IMMERSIVE_DARK_MODE
            put(35, colorref(TH.NAV_BG))       # DWMWA_CAPTION_COLOR
            put(36, colorref(TH.LABEL))        # DWMWA_TEXT_COLOR
            put(34, colorref(TH.LINE_SOFT))    # DWMWA_BORDER_COLOR
        except Exception:
            pass

    # ---- remember window size / position / maximised state
    def _geo_file(self):
        return os.path.join(CONFIG_DIR, "window.json")

    def restore_geometry(self):
        if not getattr(self, "cfg", {}).get("remember_window", True):
            return
        try:
            with open(self._geo_file(), "r", encoding="utf-8") as fh:
                d = json.load(fh)
            if abs(d.get("scale", 1) - SCALE) > 0.01:
                return
            g = parse_geometry(d.get("geometry"))
            if not g:
                return
            w, h, x, y = g
            sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
            w, h = min(w, sw), min(h, sh)
            x = max(0, min(x, sw - 120))                  # keep the window reachable
            y = max(0, min(y, sh - 120))
            self.geometry(f"{w}x{h}+{x}+{y}")
            if d.get("zoomed"):
                self.after(50, self._maximize)
        except Exception:
            pass

    def _maximize(self):
        try:
            self.state("zoomed")
        except tk.TclError:
            try:
                self.attributes("-zoomed", True)
            except tk.TclError:
                pass

    def save_geometry(self):
        if not getattr(self, "cfg", {}).get("remember_window", True):
            return
        try:
            zoomed = self.state() == "zoomed"
            if zoomed:
                self.state("normal")
                self.update_idletasks()
            data = {"geometry": self.geometry(), "zoomed": zoomed, "scale": SCALE}
            os.makedirs(CONFIG_DIR, exist_ok=True)
            with open(self._geo_file(), "w", encoding="utf-8") as fh:
                json.dump(data, fh)
        except Exception:
            pass

    # ---- mouse wheel / touchpad: smooth, one step per 120 units of movement
    def wheel_steps(self, delta):
        """Precision touchpads send many small deltas; add them up so one 'notch' = one scroll step."""
        self._wheel_acc += delta
        steps = int(self._wheel_acc / 120)
        self._wheel_acc -= steps * 120
        return steps

    # ---- drag & drop files onto the window (needs the optional tkinterdnd2 package)
    def enable_drop(self):
        try:
            if TkinterDnD is None:
                return
            self.TkdndVersion = TkinterDnD._require(self)
            self.drop_target_register(DND_FILES)
            self.dnd_bind("<<Drop>>", lambda e: self.handle_drop(self.tk.splitlist(e.data)))
            self._dnd_ok = True
        except Exception:
            self._dnd_ok = False

    def handle_drop(self, paths):
        paths = [p for p in paths if os.path.isfile(p)]
        if not paths:
            return
        low = [os.path.splitext(p)[1].lower() for p in paths]
        if self.doc is not None and all(e in IMG_EXT for e in low):      # photos dropped on an open document
            for p in paths:
                try:
                    self.place_image(ImageOps.exif_transpose(Image.open(p)).convert("RGBA"), os.path.basename(p))
                except Exception as ex:
                    messagebox.showerror(APP, f"Could not add {os.path.basename(p)}:\n{ex}")
            self.status.config(text="Photos added - drag them into place, then press Enter to apply.")
            return
        if len(paths) == 1:
            self.open_path(paths[0])
        else:
            self._create_from(paths)

    # ---- 'Open with' list / default app (Windows)
    def _register_pdf_app(self):
        """per-user registration: file type class, 'Open with' entry and the Default apps capabilities"""
        import winreg
        if getattr(sys, "frozen", False):
            cmd = f'"{sys.executable}" "%1"'
            icon = sys.executable
        else:
            pyw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
            cmd = f'"{pyw if os.path.exists(pyw) else sys.executable}" "{os.path.abspath(sys.argv[0])}" "%1"'
            icon = os.path.abspath(os.path.join(os.path.dirname(sys.argv[0]), "claudius.ico"))
        root = r"Software\Classes\CLAUDIUS.PDF"
        cap = r"Software\CLAUDIUS\Capabilities"
        for sub, name, val in ((root, "", "PDF document (Claudius PDF Editor)"),
                               (root + r"\DefaultIcon", "", icon),
                               (root + r"\shell\open", "", "Open with Claudius PDF Editor"),
                               (root + r"\shell\open\command", "", cmd),
                               (cap, "ApplicationName", "Claudius PDF Editor"),
                               (cap, "ApplicationDescription", "PDF editor"),
                               (cap + r"\FileAssociations", ".pdf", "CLAUDIUS.PDF"),
                               (r"Software\RegisteredApplications", "CLAUDIUS", cap)):
            with winreg.CreateKey(winreg.HKEY_CURRENT_USER, sub) as k:
                winreg.SetValueEx(k, name, 0, winreg.REG_SZ, val)
        with winreg.CreateKey(winreg.HKEY_CURRENT_USER, r"Software\Classes\.pdf\OpenWithProgids") as k:
            winreg.SetValueEx(k, "CLAUDIUS.PDF", 0, winreg.REG_NONE, b"")
        try:
            import ctypes
            ctypes.windll.shell32.SHChangeNotify(0x08000000, 0, None, None)
        except Exception:
            pass

    def register_open_with(self):
        if sys.platform != "win32":
            return
        if not messagebox.askyesno(APP, "Add Claudius PDF Editor to the 'Open with' list for PDF files?\n\n"
                                        "(Only your own Windows account is changed. You can still pick the "
                                        "default PDF app yourself afterwards in Windows Settings.)"):
            return
        try:
            self._register_pdf_app()
            messagebox.showinfo(APP, "Done. Right-click a PDF in File Explorer > Open with > Claudius PDF Editor.")
        except Exception as ex:
            messagebox.showerror(APP, f"Could not register:\n{ex}")

    def choose_default_app(self):
        """make Claudius PDF Editor selectable in Windows 'Default apps' and jump to that page (Windows itself asks the user to confirm)"""
        if sys.platform != "win32":
            return
        try:
            self._register_pdf_app()
        except Exception as ex:
            messagebox.showerror(APP, f"Could not register:\n{ex}")
            return
        messagebox.showinfo(APP, "Windows lets only you decide which app opens PDF files, so the last step is yours:\n\n"
                                 "1. In the Windows window that opens next, find .pdf (or type .pdf in the search box).\n"
                                 "2. Click the current app and choose Claudius PDF Editor.\n"
                                 "3. Press Set default.")
        for uri in ("ms-settings:defaultapps?registeredAppUser=CLAUDIUS", "ms-settings:defaultapps"):
            try:
                os.startfile(uri)
                return
            except Exception:
                continue
        messagebox.showinfo(APP, "Open Windows Settings > Apps > Default apps and choose Claudius PDF Editor for .pdf.")


class ExtraTools(ExtraMarkup, ExtraConvert, Win11Tools):
    """Page tools, document tools, markup, signing, forms, conversion, OCR ... (mixed into the main window)."""

    # ------------------------------------------------------------------ setup / helpers
    def _extra_init(self):
        try:
            sweep_stale_temp()
        except Exception:
            pass
        self.sel_pages = set()
        self._anchor = 0
        self._tdrag = None
        self.redo_stack = []
        self.protect = None                  # dict(user_pw, owner_pw, perms) applied when saving
        self.find_frame = None
        self.find_hits = []
        self.find_idx = -1
        self.find_var = tk.StringVar()
        self.find_info = tk.StringVar(value="")
        self.show_fields = tk.BooleanVar(value=False)
        self._view_click = None
        self.bm_items = {}
        self.sidebar_mode = "pages"
        self.recents = self.load_recents()

    def _dlg(self, title):
        d = tk.Toplevel(self)
        d.title(title)
        d.transient(self)
        d.configure(bg=BG, padx=18, pady=14)
        d.resizable(False, False)
        d.geometry("+%d+%d" % (self.winfo_rootx() + 170, self.winfo_rooty() + 110))
        d.bind("<Escape>", lambda e: d.destroy())
        try:
            d.grab_set()
        except tk.TclError:
            d.after(80, lambda: self._late_grab(d))
        return d

    @staticmethod
    def _late_grab(d):
        try:
            if d.winfo_exists():
                d.grab_set()
        except tk.TclError:
            pass

    @staticmethod
    def _label(parent, text, bold=False, fg=None, **kw):
        if fg is not None and fg == TH.TERTIARY:
            kw.setdefault("wraplength", S(440))        # hint texts wrap instead of being cut off
        fg = fg or TH.LABEL
        return tk.Label(parent, text=text, bg=BG, fg=fg, font=(UI_FONT, 9, "bold" if bold else "normal"),
                        anchor="w", justify="left", **kw)

    def _buttons(self, d, ok_text, ok_cmd, row, colspan=4):
        f = tk.Frame(d, bg=BG)
        f.grid(row=row, column=0, columnspan=colspan, sticky="e", pady=(14, 0))
        ttk.Button(f, text=ok_text, command=ok_cmd).pack(side="left", padx=4)
        ttk.Button(f, text="Cancel", command=d.destroy).pack(side="left")
        return f

    def _color_button(self, parent, initial):
        holder = {"rgb": tuple(initial)}
        b = tk.Button(parent, width=4, relief="solid", bd=1, cursor="hand2", bg=self._hex(initial),
                      activebackground=self._hex(initial))

        def pick():
            c = colorchooser.askcolor(color=self._hex(holder["rgb"]), parent=parent.winfo_toplevel(), title="Choose color")
            if c and c[0]:
                holder["rgb"] = tuple(v / 255 for v in c[0])
                b.config(bg=self._hex(holder["rgb"]), activebackground=self._hex(holder["rgb"]))
        b.config(command=pick)
        return b, holder

    def _pages_or_warn(self, text, parent=None):
        try:
            pages = parse_pages(text, len(self.doc))
            if not pages:
                raise ValueError("No pages selected.")
            return pages
        except ValueError as ex:
            messagebox.showwarning(APP, str(ex), parent=parent or self)
            return None

    def _range_row(self, d, row, default="all", label="Pages", colspan=3):
        """a 'Pages: [all]' entry with a hint - returns the StringVar"""
        v = tk.StringVar(value=default)
        self._label(d, label).grid(row=row, column=0, sticky="w", pady=3)
        e = ttk.Entry(d, textvariable=v, width=24)
        e.grid(row=row, column=1, columnspan=colspan, sticky="w", pady=3)
        self._label(d, f"of {len(self.doc)}   (e.g. all, 1-3, 5, 8-, odd, even)", fg=TH.TERTIARY).grid(
            row=row + 1, column=1, columnspan=colspan, sticky="w")
        return v

    def _progress(self, title, total):
        """Small progress window.  Returns (update(done, total), close())."""
        d = tk.Toplevel(self)
        d.title(title)
        d.transient(self)
        d.configure(bg=BG, padx=20, pady=16)
        d.resizable(False, False)
        d.geometry("+%d+%d" % (self.winfo_rootx() + 300, self.winfo_rooty() + 200))
        lab = self._label(d, title)
        lab.pack(anchor="w")
        bar = ttk.Progressbar(d, length=320, maximum=max(1, total))
        bar.pack(pady=(8, 0))
        d.update()

        def upd(done, tot=None, text=None):
            bar["maximum"] = max(1, tot or total)
            bar["value"] = done
            if text:
                lab.config(text=text)
            d.update()

        def close():
            try:
                d.destroy()
            except tk.TclError:
                pass
        return upd, close

    def ask_text(self, title, prompt, initial="", check=None, check_default=True, height=7):
        """Multi-line text dialog.  Returns (text, checked) or None."""
        d = self._dlg(title)
        out = {}
        self._label(d, prompt).grid(row=0, column=0, sticky="w")
        t = tk.Text(d, width=56, height=height, wrap="word", font=(UI_FONT, 10), relief="solid", bd=1)
        t.grid(row=1, column=0, pady=6)
        t.insert("1.0", initial)
        t.focus_set()
        cv = tk.BooleanVar(value=check_default)
        if check:
            tk.Checkbutton(d, text=check, variable=cv, bg=BG, activebackground=BG, font=(UI_FONT, 9)).grid(row=2, column=0, sticky="w")

        def ok():
            out["v"] = (t.get("1.0", "end-1c"), cv.get())
            d.destroy()
        self._buttons(d, "OK", ok, 3, colspan=1)
        d.wait_window()
        return out.get("v")

    def reveal(self, path):
        """Show a file / folder in the system file manager."""
        try:
            if sys.platform == "win32":
                if os.path.isdir(path):
                    os.startfile(path)
                else:
                    subprocess.Popen(["explorer", "/select,", os.path.normpath(path)])
            elif sys.platform == "darwin":
                subprocess.Popen(["open", "-R" if os.path.isfile(path) else "", path] if os.path.isfile(path) else ["open", path])
            else:
                subprocess.Popen(["xdg-open", path if os.path.isdir(path) else os.path.dirname(path)])
        except Exception:
            pass

    def _open_pdf_file(self, path, parent=None):
        """Open a PDF from disk (asking for the password if needed).  Returns a fitz.Document or None."""
        try:
            doc = fitz.open(path)
        except Exception as ex:
            messagebox.showerror(APP, f"Could not open {os.path.basename(path)}:\n{ex}", parent=parent or self)
            return None
        if doc.needs_pass and self.unlock_document(doc, os.path.basename(path), parent) is None:
            return None
        return doc

    def _adopt_doc(self, doc, path=None, modified=False):
        """Make `doc` the document being edited."""
        self.doc, self.path = doc, path
        self.floats, self.sel = [], None
        self.page_no = 0
        self.sel_pages = {0}
        self.undo_stack.clear()
        self.redo_stack.clear()
        self.protect = None
        self.modified = modified
        self.update_idletasks()
        self.fit_page(render=False)
        self.build_thumbs()
        self.render()
        self.update_ui()

    def _rebuild(self):
        if self.doc is None:
            return
        self.page_no = max(0, min(self.page_no, len(self.doc) - 1))
        self._cont_raw = None
        if self.fit_mode:
            self._apply_fit(render=False)
        self.build_thumbs()
        self.render()
        self.update_ui()

    def _run_edit(self, fn, rebuild=False, msg=None, thumbs=None):
        """Snapshot for undo, run fn(); on error restore the document and show the message."""
        self._snap_label = msg
        try:
            self.snapshot()
        finally:
            self._snap_label = None
        try:
            fn()
        except Exception as ex:
            data, pno = self.undo_stack.pop()
            self.doc = fitz.open(stream=data, filetype="pdf")
            self.page_no = min(pno, len(self.doc) - 1)
            self._rebuild()
            messagebox.showerror(APP, f"That did not work:\n{ex}")
            return False
        if rebuild:
            self._rebuild()
        else:
            self.render()
            for i in (thumbs if thumbs is not None else [self.page_no]):
                self.refresh_thumb(i)
            self.update_ui()
        if msg:
            self.status.config(text=msg)
        return True

    def redo(self):
        if not self.redo_stack:
            self.status.config(text="Nothing to redo.")
            return
        data, pno = self.redo_stack.pop()
        self.undo_stack.append((self.doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE), self.page_no))
        self.doc = fitz.open(stream=data, filetype="pdf")
        self.page_no = min(pno, len(self.doc) - 1)
        self.modified = True
        self._rebuild()

    # ------------------------------------------------------------------ config / recent files
    def load_recents(self):
        try:
            with open(os.path.join(CONFIG_DIR, "recent.json"), "r", encoding="utf-8") as fh:
                return [p for p in json.load(fh) if os.path.exists(p)][:10]
        except Exception:
            return []

    def add_recent(self, path):
        if not path:
            return
        path = os.path.abspath(path)
        self.recents = [path] + [p for p in self.recents if os.path.abspath(p) != path]
        self.recents = self.recents[:10]
        try:
            os.makedirs(CONFIG_DIR, exist_ok=True)
            with open(os.path.join(CONFIG_DIR, "recent.json"), "w", encoding="utf-8") as fh:
                json.dump(self.recents, fh)
        except Exception:
            pass

    def _fill_recent_menu(self):
        m = self.recent_menu
        m.delete(0, "end")
        self.recents = self.load_recents()
        if not self.recents:
            m.add_command(label="(no recent files)", state="disabled")
        for p in self.recents:
            m.add_command(label=os.path.basename(p), command=lambda p=p: self._open_recent(p))
        if self.recents:
            m.add_separator()
            m.add_command(label="Clear list of recent files", command=self.clear_recents)

    def clear_recents(self):
        """privacy: forget which files were opened"""
        self.recents = []
        try:
            os.remove(os.path.join(CONFIG_DIR, "recent.json"))
        except OSError:
            pass
        if self.doc is None:
            self.render()

    def _open_recent(self, p):
        self.open_path(p)

    # ------------------------------------------------------------------ save options / encryption
    def _save_kwargs(self):
        kw = dict(garbage=3, deflate=True)
        if self.protect:
            kw.update(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=self.protect["user_pw"],
                      owner_pw=self.protect["owner_pw"], permissions=self.protect["perms"])
        else:
            kw["encryption"] = fitz.PDF_ENCRYPT_NONE
        return kw

    # ================================================================== page selection (thumbnails)
    def selected_pages(self):
        if not self.doc:
            return []
        n = len(self.doc)
        sp = sorted(p for p in self.sel_pages if 0 <= p < n)
        return sp or [self.page_no]

    def _thumb_press(self, e):
        if not self.doc:
            return
        n = len(self.doc)
        i = int(self.thumbs.canvasy(e.y) // CELL_H)
        if not 0 <= i < n:
            return
        ctrl, shift = bool(e.state & 0x4), bool(e.state & 0x1)
        keep = False
        if shift:
            lo, hi = sorted((self._anchor, i))
            self.sel_pages = set(range(lo, hi + 1))
        elif ctrl:
            self.sel_pages = set(self.sel_pages) ^ {i}
            if not self.sel_pages:
                self.sel_pages = {i}
            self._anchor = i
        else:
            if i in self.sel_pages and len(self.sel_pages) > 1:
                keep = True                       # might be the start of dragging a group
            else:
                self.sel_pages = {i}
            self._anchor = i
        self._tdrag = {"i": i, "start": (e.x, e.y), "moved": False, "keep": keep, "mod": ctrl or shift}
        sel = set(self.sel_pages)
        if not ctrl:
            self.goto(i)
        self.sel_pages = sel
        self._select_thumb()

    def _thumb_drag(self, e):
        d = self._tdrag
        if not d or not self.doc:
            return
        if not d["moved"]:
            if abs(e.y - d["start"][1]) < 6 and abs(e.x - d["start"][0]) < 6:
                return
            d["moved"] = True
        y = self.thumbs.canvasy(e.y)
        gap = max(0, min(len(self.doc), int((y + CELL_H / 2) // CELL_H)))
        d["gap"] = gap
        self.thumbs.delete("tdrop")
        yy = gap * CELL_H + 6
        self.thumbs.create_line(6, yy, max(40, self.thumbs.winfo_width() - 6), yy, fill=ACCENT, width=3, tags="tdrop")
        h = self.thumbs.winfo_height()
        if e.y < 18:
            self.thumbs.yview_scroll(-1, "units")
        elif e.y > h - 18:
            self.thumbs.yview_scroll(1, "units")

    def _thumb_release(self, e):
        d = self._tdrag
        self._tdrag = None
        self.thumbs.delete("tdrop")
        if not d or not self.doc:
            return
        if d["moved"] and "gap" in d:
            if d["i"] not in self.sel_pages:
                self.sel_pages = {d["i"]}
            self.reorder_pages(self.selected_pages(), d["gap"])
        elif d["keep"] and not d["mod"]:
            self.sel_pages = {d["i"]}
            self.goto(d["i"])
            self._select_thumb()

    def reorder_pages(self, pages, gap):
        """Move `pages` so they sit just before original page index `gap` (len(doc) = at the end)."""
        n = len(self.doc)
        sel = sorted(pages)
        order = [p for p in range(gap) if p not in sel] + sel + [p for p in range(gap, n) if p not in sel]
        if order == list(range(n)):
            return False
        self.snapshot()
        self.doc.select(order)
        pos = [order.index(p) for p in sel]
        self.page_no = pos[0]
        self.sel_pages = set(pos)
        self._rebuild()
        self.status.config(text=f"Moved {len(sel)} page(s).")
        return True

    def _thumb_menu(self, e):
        if not self.doc:
            return
        i = int(self.thumbs.canvasy(e.y) // CELL_H)
        if 0 <= i < len(self.doc) and i not in self.sel_pages:
            self.sel_pages = {i}
            self.goto(i)
            self.sel_pages = {i}
            self._select_thumb()
        m = tk.Menu(self, tearoff=0)
        m.add_command(label="Rotate right", command=lambda: self.rotate_selected(90))
        m.add_command(label="Rotate left", command=lambda: self.rotate_selected(-90))
        m.add_separator()
        m.add_command(label="Extract pages...", command=self.dlg_extract)
        m.add_command(label="Delete pages", command=self.delete_selected)
        m.add_command(label="Insert blank page after", command=self.insert_blank)
        m.add_separator()
        m.add_command(label="Select all pages", command=self.select_all_pages)
        m.tk_popup(e.x_root, e.y_root)

    def select_all_pages(self):
        if self.doc:
            self.sel_pages = set(range(len(self.doc)))
            self._select_thumb()

    # ================================================================== page tools
    def rotate_selected(self, deg):
        if not self.need_doc():
            return
        pages = self.selected_pages()

        def work():
            for i in pages:
                pg = self.doc[i]
                pg.set_rotation((pg.rotation + deg) % 360)
        if self.fit_mode:
            self._apply_fit(render=False)
        self._run_edit(work, thumbs=pages)
        if self.fit_mode:
            self._apply_fit(render=False)
            self.render()

    def delete_selected(self):
        if not self.need_doc():
            return
        pages = self.selected_pages()
        if len(pages) >= len(self.doc):
            messagebox.showinfo(APP, "A document needs at least one page - you can't delete them all.")
            return
        if not messagebox.askyesno(APP, f"Delete {len(pages)} page(s) ({pages_to_text(pages)})?"):
            return
        self.snapshot()
        self.doc.delete_pages(pages)
        self.page_no = min(pages[0], len(self.doc) - 1)
        self.sel_pages = {self.page_no}
        self._rebuild()
        self.status.config(text=f"Deleted {len(pages)} page(s).")

    # ---- extract
    def dlg_extract(self):
        if not self.need_doc():
            return
        n = len(self.doc)
        d = self._dlg("Extract pages")
        v = self._range_row(d, 0, pages_to_text(self.selected_pages()), "Pages to extract")
        delete_after = tk.BooleanVar(value=False)
        tk.Checkbutton(d, text="Also remove these pages from the open document", variable=delete_after,
                       bg=BG, activebackground=BG, font=(UI_FONT, 9)).grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 0))

        def go():
            pages = self._pages_or_warn(v.get(), d)
            if pages is None:
                return
            if delete_after.get() and len(pages) >= n:
                messagebox.showwarning(APP, "You can't remove every page.", parent=d)
                return
            base = os.path.splitext(os.path.basename(self.path))[0] if self.path else "document"
            p = filedialog.asksaveasfilename(parent=d, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")],
                                             initialfile=f"{base}_extract.pdf")
            if not p:
                return
            d.destroy()
            self.do_extract(pages, p, delete_after.get())
        self._buttons(d, "Extract", go, 3)

    def do_extract(self, pages, out, delete_after=False):
        self.commit_floats()
        new = fitz.open()
        for a, b in runs(pages):
            new.insert_pdf(self.doc, from_page=a, to_page=b)
        new.save(out, garbage=3, deflate=True)
        new.close()
        msg = f"Extracted {len(pages)} page(s) to {out}"
        if delete_after:
            self.snapshot()
            self.doc.delete_pages(pages)
            self.page_no = min(pages[0], len(self.doc) - 1)
            self.sel_pages = {self.page_no}
            self._rebuild()
        self.status.config(text=msg)
        return out

    # ---- split
    def dlg_split(self):
        if not self.need_doc():
            return
        n = len(self.doc)
        d = self._dlg("Split PDF")
        mode = tk.StringVar(value="n")
        nv = tk.IntVar(value=1)
        rv = tk.StringVar(value="1-%d" % max(1, n // 2) + (", %d-" % (n // 2 + 1) if n > 1 else ""))
        self._label(d, f"This document has {n} page(s).  Split it into several files:", bold=True).grid(
            row=0, column=0, columnspan=4, sticky="w", pady=(0, 6))
        tk.Radiobutton(d, text="Every", variable=mode, value="n", bg=BG, activebackground=BG).grid(row=1, column=0, sticky="w")
        f = tk.Frame(d, bg=BG)
        f.grid(row=1, column=1, columnspan=3, sticky="w")
        ttk.Spinbox(f, from_=1, to=max(1, n), width=5, textvariable=nv).pack(side="left")
        self._label(f, "  page(s)  per file  (1 = one file per page)").pack(side="left")
        tk.Radiobutton(d, text="Ranges", variable=mode, value="r", bg=BG, activebackground=BG).grid(row=2, column=0, sticky="w", pady=4)
        ttk.Entry(d, textvariable=rv, width=34).grid(row=2, column=1, columnspan=3, sticky="w")
        self._label(d, "each range becomes one file, e.g.  1-3, 4-6, 7-", fg=TH.TERTIARY).grid(row=3, column=1, columnspan=3, sticky="w")

        def go():
            try:
                if mode.get() == "n":
                    k = max(1, int(nv.get()))
                    groups = [list(range(a, min(n, a + k))) for a in range(0, n, k)]
                else:
                    groups = parse_range_groups(rv.get(), n)
            except (ValueError, tk.TclError) as ex:
                messagebox.showwarning(APP, str(ex), parent=d)
                return
            folder = filedialog.askdirectory(parent=d, title="Choose the folder for the new files")
            if not folder:
                return
            d.destroy()
            base = os.path.splitext(os.path.basename(self.path))[0] if self.path else "document"
            self.do_split(groups, folder, base)
        self._buttons(d, "Split", go, 4)

    def do_split(self, groups, folder, base):
        self.commit_floats()
        files = []
        width = len(str(len(groups)))
        for k, g in enumerate(groups, 1):
            new = fitz.open()
            for a, b in runs(g):
                new.insert_pdf(self.doc, from_page=a, to_page=b)
            path = os.path.join(folder, f"{safe_name(base)}_part{str(k).zfill(width)}.pdf")
            new.save(path, garbage=3, deflate=True)
            new.close()
            files.append(path)
        self.status.config(text=f"Split into {len(files)} file(s) in {folder}")
        if messagebox.askyesno(APP, f"Created {len(files)} file(s).\n\nOpen the folder?"):
            self.reveal(folder)
        return files

    # ---- merge several PDFs
    def dlg_merge(self):
        d = self._dlg("Merge PDF files")
        files = []
        self._label(d, "Files will be joined in this order:", bold=True).grid(row=0, column=0, columnspan=3, sticky="w")
        lb = tk.Listbox(d, width=62, height=11, selectmode="browse", activestyle="none")
        lb.grid(row=1, column=0, rowspan=6, pady=4, sticky="nsew")
        use_cur = tk.BooleanVar(value=self.doc is not None)
        out_mode = tk.StringVar(value="open")

        def refill(sel=None):
            lb.delete(0, "end")
            for p in files:
                lb.insert("end", p)
            if sel is not None and files:
                sel = max(0, min(sel, len(files) - 1))
                lb.selection_set(sel)
                lb.see(sel)

        def cur():
            s = lb.curselection()
            return s[0] if s else None

        def add():
            ps = filedialog.askopenfilenames(parent=d, title="Choose PDF files", filetypes=[("PDF files", "*.pdf")])
            files.extend(ps)
            refill()

        def up():
            i = cur()
            if i:
                files[i - 1], files[i] = files[i], files[i - 1]
                refill(i - 1)

        def down():
            i = cur()
            if i is not None and i < len(files) - 1:
                files[i + 1], files[i] = files[i], files[i + 1]
                refill(i + 1)

        def remove():
            i = cur()
            if i is not None:
                files.pop(i)
                refill(i)

        def sort():
            files.sort(key=lambda p: os.path.basename(p).lower())
            refill()
        for r, (t, c) in enumerate((("Add files...", add), ("Up", up), ("Down", down), ("Remove", remove), ("Sort by name", sort))):
            ttk.Button(d, text=t, command=c, width=13).grid(row=1 + r, column=1, padx=(8, 0), pady=1, sticky="n")
        r = 7
        if self.doc is not None:
            tk.Checkbutton(d, text="Put the open document first", variable=use_cur, bg=BG, activebackground=BG,
                           font=(UI_FONT, 9)).grid(row=r, column=0, sticky="w")
            r += 1
        tk.Radiobutton(d, text="Open the merged document here", variable=out_mode, value="open", bg=BG,
                       activebackground=BG).grid(row=r, column=0, sticky="w")
        tk.Radiobutton(d, text="Save the merged document as a new file...", variable=out_mode, value="save", bg=BG,
                       activebackground=BG).grid(row=r + 1, column=0, sticky="w")

        def go():
            count = len(files) + (1 if (self.doc is not None and use_cur.get()) else 0)
            if count < 2:
                messagebox.showinfo(APP, "Add at least two PDF files to merge.", parent=d)
                return
            out = None
            if out_mode.get() == "save":
                out = filedialog.asksaveasfilename(parent=d, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")],
                                                   initialfile="merged.pdf")
                if not out:
                    return
            docs = []
            if self.doc is not None and use_cur.get():
                self.commit_floats()
                docs.append(self.doc)
            for p in files:
                x = self._open_pdf_file(p, d)
                if x is None:
                    return
                docs.append(x)
            d.destroy()
            self.do_merge(docs, out)
        self._buttons(d, "Merge", go, r + 2, colspan=2)

    def do_merge(self, docs, out=None):
        new = fitz.open()
        for s in docs:
            new.insert_pdf(s)
        if out:
            new.save(out, garbage=3, deflate=True)
            self.status.config(text=f"Merged {len(docs)} files into {out}")
            return out
        was = self.doc
        if was is not None and any(s is was for s in docs):
            self.snapshot()
            self.doc = new
            self.page_no = 0
            self.sel_pages = {0}
            self._rebuild()
        else:
            self._adopt_doc(new, None, modified=True)
        self.status.config(text=f"Merged {len(docs)} files ({len(new)} pages).")
        return new

    # ---- insert pages from another PDF
    def dlg_insert_pdf(self):
        if not self.need_doc():
            return
        d = self._dlg("Insert pages from another PDF")
        state = {"doc": None}
        name = tk.StringVar(value="(no file chosen)")
        where = tk.StringVar(value="after")
        rng = tk.StringVar(value="all")
        self._label(d, "Source PDF").grid(row=0, column=0, sticky="w")
        self._label(d, "", fg=TH.SECOND).grid(row=0, column=1, sticky="w", padx=6)
        lab = tk.Label(d, textvariable=name, bg=BG, fg=TH.SECOND, font=(UI_FONT, 9), anchor="w", width=38)
        lab.grid(row=0, column=1, sticky="w", padx=6)

        def choose():
            p = filedialog.askopenfilename(parent=d, filetypes=[("PDF files", "*.pdf")])
            if p:
                x = self._open_pdf_file(p, d)
                if x:
                    state["doc"] = x
                    name.set(_tr("{}   ({} pages)").replace("{}", os.path.basename(p), 1).replace("{}", str(len(x)), 1))
        ttk.Button(d, text="Choose...", command=choose).grid(row=0, column=2, padx=4)
        self._label(d, "Pages to insert").grid(row=1, column=0, sticky="w", pady=6)
        ttk.Entry(d, textvariable=rng, width=22).grid(row=1, column=1, sticky="w", padx=6)
        self._label(d, "Insert").grid(row=2, column=0, sticky="nw", pady=(4, 0))
        f = tk.Frame(d, bg=BG)
        f.grid(row=2, column=1, columnspan=2, sticky="w", padx=6)
        for text, val in (("After the current page", "after"), ("Before the current page", "before"),
                          ("At the start", "start"), ("At the end", "end")):
            tk.Radiobutton(f, text=text, variable=where, value=val, bg=BG, activebackground=BG).pack(anchor="w")

        def go():
            if state["doc"] is None:
                messagebox.showinfo(APP, "Choose a PDF file first.", parent=d)
                return
            try:
                pages = parse_pages(rng.get(), len(state["doc"]))
            except ValueError as ex:
                messagebox.showwarning(APP, str(ex), parent=d)
                return
            d.destroy()
            self.do_insert_pdf(state["doc"], pages, where.get())
        self._buttons(d, "Insert", go, 3, colspan=3)

    def do_insert_pdf(self, src, pages, where="after"):
        pos = {"after": self.page_no + 1, "before": self.page_no, "start": 0, "end": len(self.doc)}[where]
        self.snapshot()
        first = pos
        for a, b in runs(pages):
            self.doc.insert_pdf(src, from_page=a, to_page=b, start_at=pos)
            pos += b - a + 1
        self.page_no = first
        self.sel_pages = set(range(first, pos))
        self._rebuild()
        self.status.config(text=f"Inserted {len(pages)} page(s).")

    # ---- crop
    def dlg_crop(self):
        if not self.need_doc():
            return
        d = self._dlg("Crop pages")
        unit = tk.StringVar(value="mm")
        vals = {k: tk.DoubleVar(value=0.0) for k in ("top", "bottom", "left", "right")}
        self._label(d, "Remove from each edge:", bold=True).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 6))
        for r, (k, t) in enumerate((("top", "Top"), ("bottom", "Bottom"), ("left", "Left"), ("right", "Right"))):
            self._label(d, t).grid(row=1 + r // 2, column=(r % 2) * 2, sticky="w", pady=2, padx=(0, 6))
            ttk.Spinbox(d, from_=0, to=500, increment=1, width=7, textvariable=vals[k]).grid(row=1 + r // 2, column=(r % 2) * 2 + 1, padx=(0, 14))
        self._label(d, "Unit").grid(row=3, column=0, sticky="w", pady=4)
        ttk.Combobox(d, textvariable=unit, values=["mm", "inch", "pt"], state="readonly", width=6).grid(row=3, column=1, sticky="w")
        self._label(d, "Apply to", bold=True).grid(row=4, column=0, sticky="w", pady=(10, 0))
        scope = tk.StringVar(value="all")
        rng = tk.StringVar(value="")
        for r, (t, v) in enumerate((("All pages", "all"), ("Current page only", "cur"), ("Selected pages", "sel"), ("Page range:", "rng"))):
            tk.Radiobutton(d, text=t, variable=scope, value=v, bg=BG, activebackground=BG).grid(row=5 + r, column=0, columnspan=2, sticky="w")
        ttk.Entry(d, textvariable=rng, width=18).grid(row=8, column=2, columnspan=2, sticky="w")

        def go():
            try:
                f = {"mm": 72 / 25.4, "inch": 72.0, "pt": 1.0}[unit.get()]
                t, b, l, r_ = (float(vals[k].get()) * f for k in ("top", "bottom", "left", "right"))
            except (tk.TclError, ValueError):
                messagebox.showwarning(APP, "Enter numbers for the margins.", parent=d)
                return
            sc = scope.get()
            if sc == "all":
                pages = list(range(len(self.doc)))
            elif sc == "cur":
                pages = [self.page_no]
            elif sc == "sel":
                pages = self.selected_pages()
            else:
                pages = self._pages_or_warn(rng.get(), d)
                if pages is None:
                    return
            d.destroy()
            self.do_crop(pages, t, b, l, r_)
        self._buttons(d, "Crop", go, 9)

    def do_crop(self, pages, top, bottom, left, right):
        bad = []

        def work():
            for i in pages:
                pg = self.doc[i]
                v = pg.rect
                nv = fitz.Rect(v.x0 + left, v.y0 + top, v.x1 - right, v.y1 - bottom)
                if nv.width < 20 or nv.height < 20:
                    bad.append(i + 1)
                    continue
                ur = nv * pg.derotation_matrix
                ur.normalize()
                pos = pg.cropbox_position
                pg.set_cropbox(ur + (pos.x, pos.y, pos.x, pos.y))
        ok = self._run_edit(work, rebuild=True, msg=f"Cropped {len(pages) - len(bad)} page(s).")
        if bad:
            messagebox.showwarning(APP, "Skipped pages where nothing would be left: {}".format(pages_to_text([b - 1 for b in bad])))
        return ok

    def crop_drag(self, x0, y0, x1, y1):
        """Crop tool: the dragged rectangle (canvas coordinates) is what stays."""
        (vx0, vy0), (vx1, vy1) = self._c2v(x0, y0), self._c2v(x1, y1)
        vx0, vx1 = sorted((vx0, vx1))
        vy0, vy1 = sorted((vy0, vy1))
        v = self.page().rect
        t, b, l, r = max(0, vy0 - v.y0), max(0, v.y1 - vy1), max(0, vx0 - v.x0), max(0, v.x1 - vx1)
        ans = messagebox.askyesnocancel(APP, "Crop to this area.\n\nYes = apply the same margins to ALL pages\n"
                                             "No = this page only")
        if ans is None:
            return
        self.set_tool("view")
        self.do_crop(list(range(len(self.doc))) if ans else [self.page_no], t, b, l, r)

    # ================================================================== document tools
    def _font_for(self, text):
        return font_kwargs(text)

    def dlg_header_footer(self):
        if not self.need_doc():
            return
        d = self._dlg("Header & footer")
        keys = ("hl", "hc", "hr", "fl", "fc", "fr")
        v = {k: tk.StringVar() for k in keys}
        v["fc"].set("")
        for c, t in enumerate(("Left", "Center", "Right")):
            self._label(d, t, bold=True).grid(row=0, column=1 + c, sticky="w", padx=3)
        for r, (t, pre) in enumerate((("Header", "h"), ("Footer", "f"))):
            self._label(d, t, bold=True).grid(row=1 + r, column=0, sticky="w", pady=3, padx=(0, 8))
            for c, side in enumerate("lcr"):
                ttk.Entry(d, textvariable=v[pre + side], width=22).grid(row=1 + r, column=1 + c, padx=3, pady=3)
        self._label(d, "You can use  {page}  {pages}  {date}  {time}  {file}", fg=TH.TERTIARY).grid(row=3, column=0, columnspan=4, sticky="w")
        size, margin, skip = tk.IntVar(value=10), tk.IntVar(value=30), tk.BooleanVar(value=False)
        self._label(d, "Font size").grid(row=4, column=0, sticky="w", pady=(10, 0))
        ttk.Spinbox(d, from_=5, to=48, width=5, textvariable=size).grid(row=4, column=1, sticky="w", pady=(10, 0), padx=3)
        self._label(d, "Margin (pt)").grid(row=4, column=2, sticky="e", pady=(10, 0))
        ttk.Spinbox(d, from_=5, to=150, width=5, textvariable=margin).grid(row=4, column=3, sticky="w", pady=(10, 0), padx=3)
        self._label(d, "Color").grid(row=5, column=0, sticky="w", pady=4)
        cb, ch = self._color_button(d, (0.2, 0.2, 0.2))
        cb.grid(row=5, column=1, sticky="w", padx=3)
        rng = self._range_row(d, 6, "all")
        tk.Checkbutton(d, text="Skip the first page", variable=skip, bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=8, column=0, columnspan=3, sticky="w")

        def _hf_get():
            return {"t": {k: v[k].get() for k in keys}, "size": size.get(), "margin": margin.get(), "color": list(ch["rgb"]), "skip": skip.get()}

        def _hf_set(p):
            for k in keys:
                v[k].set(p["t"].get(k, ""))
            size.set(p["size"])
            margin.set(p["margin"])
            skip.set(p["skip"])
            ch["rgb"] = tuple(p["color"])
            cb.config(bg=self._hex(ch["rgb"]), activebackground=self._hex(ch["rgb"]))
        self._preset_bar(d, 9, "header", _hf_get, _hf_set)

        def go():
            pages = self._pages_or_warn(rng.get(), d)
            if pages is None:
                return
            texts = {k: v[k].get() for k in keys}
            if not any(t.strip() for t in texts.values()):
                messagebox.showinfo(APP, "Type some header or footer text first.", parent=d)
                return
            d.destroy()
            self.do_header_footer(texts, size.get(), margin.get(), ch["rgb"], pages, skip.get())
        self._buttons(d, "Apply", go, 10)

    def do_header_footer(self, texts, size, margin, color, pages, skip_first=False):
        fname = os.path.splitext(os.path.basename(self.path))[0] if self.path else ""
        total = len(self.doc)
        pages = [p for p in pages if not (skip_first and p == 0)]

        def work():
            for i in pages:
                pg = self.doc[i]
                v = pg.rect
                for row, y in (("h", v.y0 + margin + size * 0.85), ("f", v.y1 - margin)):
                    for side, align, x in (("l", "left", v.x0 + margin), ("c", "center", v.x0 + v.width / 2),
                                           ("r", "right", v.x1 - margin)):
                        t = texts.get(row + side, "")
                        if t.strip():
                            t = expand_tokens(t, i + 1, total, fname)
                            put_text(pg, x, y, t, size, color, font_kwargs(t), align)
        return self._run_edit(work, thumbs=pages, msg=f"Header / footer added to {len(pages)} page(s).")

    def dlg_page_numbers(self):
        if not self.need_doc():
            return
        d = self._dlg("Page numbers")
        fmt = tk.StringVar(value="Page {n} of {total}")
        pos = tk.StringVar(value="Bottom center")
        start, size, margin, skip = tk.IntVar(value=1), tk.IntVar(value=10), tk.IntVar(value=30), tk.BooleanVar(value=False)
        self._label(d, "Format").grid(row=0, column=0, sticky="w", pady=3)
        ttk.Combobox(d, textvariable=fmt, width=22, values=["{n}", "Page {n}", "Page {n} of {total}", "{n} / {total}", "- {n} -"]
                     ).grid(row=0, column=1, sticky="w", padx=4)
        self._label(d, "Position").grid(row=1, column=0, sticky="w", pady=3)
        ttk.Combobox(d, textvariable=pos, state="readonly", width=22,
                     values=["Bottom center", "Bottom right", "Bottom left", "Top center", "Top right", "Top left"]
                     ).grid(row=1, column=1, sticky="w", padx=4)
        self._label(d, "Start at").grid(row=2, column=0, sticky="w", pady=3)
        ttk.Spinbox(d, from_=-99, to=9999, width=6, textvariable=start).grid(row=2, column=1, sticky="w", padx=4)
        self._label(d, "Font size").grid(row=3, column=0, sticky="w", pady=3)
        ttk.Spinbox(d, from_=5, to=48, width=6, textvariable=size).grid(row=3, column=1, sticky="w", padx=4)
        self._label(d, "Margin (pt)").grid(row=4, column=0, sticky="w", pady=3)
        ttk.Spinbox(d, from_=5, to=150, width=6, textvariable=margin).grid(row=4, column=1, sticky="w", padx=4)
        self._label(d, "Color").grid(row=5, column=0, sticky="w", pady=3)
        cb, ch = self._color_button(d, (0.2, 0.2, 0.2))
        cb.grid(row=5, column=1, sticky="w", padx=4)
        rng = self._range_row(d, 6, "all")
        tk.Checkbutton(d, text="Don't number the first page", variable=skip, bg=BG, activebackground=BG,
                       font=(UI_FONT, 9)).grid(row=8, column=0, columnspan=3, sticky="w")

        def _pn_get():
            return {"fmt": fmt.get(), "pos": pos.get(), "start": start.get(), "size": size.get(), "margin": margin.get(),
                    "color": list(ch["rgb"]), "skip": skip.get()}

        def _pn_set(p):
            fmt.set(p["fmt"])
            pos.set(p["pos"])
            start.set(p["start"])
            size.set(p["size"])
            margin.set(p["margin"])
            skip.set(p["skip"])
            ch["rgb"] = tuple(p["color"])
            cb.config(bg=self._hex(ch["rgb"]), activebackground=self._hex(ch["rgb"]))
        self._preset_bar(d, 9, "numbers", _pn_get, _pn_set, colspan=3)

        def go():
            pages = self._pages_or_warn(rng.get(), d)
            if pages is None:
                return
            d.destroy()
            self.do_page_numbers(fmt.get(), pos.get(), start.get(), size.get(), margin.get(), ch["rgb"], pages, skip.get())
        self._buttons(d, "Apply", go, 10, colspan=3)

    def do_page_numbers(self, fmt, position, start, size, margin, color, pages, skip_first=False):
        pages = [p for p in pages if not (skip_first and p == 0)]
        total = len(self.doc)
        vert, _, horiz = position.lower().partition(" ")
        align = {"left": "left", "center": "center", "right": "right"}[horiz]

        def work():
            for k, i in enumerate(pages):
                pg = self.doc[i]
                v = pg.rect
                t = fmt.replace("{n}", str(start + k)).replace("{total}", str(total))
                x = {"left": v.x0 + margin, "center": v.x0 + v.width / 2, "right": v.x1 - margin}[horiz]
                y = v.y1 - margin if vert == "bottom" else v.y0 + margin + size * 0.85
                put_text(pg, x, y, t, size, color, font_kwargs(t), align)
        return self._run_edit(work, thumbs=pages, msg=f"Page numbers added to {len(pages)} page(s).")

    # ---- watermark
    def dlg_watermark(self):
        if not self.need_doc():
            return
        d = self._dlg("Watermark")
        kind = tk.StringVar(value="text")
        text = tk.StringVar(value="CONFIDENTIAL")
        img = tk.StringVar(value="")
        opacity, angle, scale = tk.IntVar(value=30), tk.IntVar(value=45), tk.IntVar(value=60)
        position, layer = tk.StringVar(value="Center"), tk.StringVar(value="Above the page content")
        tk.Radiobutton(d, text="Text", variable=kind, value="text", bg=BG, activebackground=BG).grid(row=0, column=0, sticky="w")
        ttk.Entry(d, textvariable=text, width=30).grid(row=0, column=1, columnspan=3, sticky="w", padx=4)
        tk.Radiobutton(d, text="Image", variable=kind, value="image", bg=BG, activebackground=BG).grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(d, textvariable=img, width=30).grid(row=1, column=1, columnspan=2, sticky="w", padx=4)

        def pick():
            p = filedialog.askopenfilename(parent=d, filetypes=[("Images", "*.png *.jpg *.jpeg *.bmp *.gif *.webp"), ("All", "*.*")])
            if p:
                img.set(p)
                kind.set("image")
        ttk.Button(d, text="Browse...", command=pick).grid(row=1, column=3, padx=2)
        rows = (("Color", None), ("Opacity %", opacity), ("Rotation", angle), ("Size % of width", scale))
        self._label(d, "Color").grid(row=2, column=0, sticky="w", pady=3)
        cb, ch = self._color_button(d, (0.8, 0.1, 0.1))
        cb.grid(row=2, column=1, sticky="w", padx=4)
        for r, (t, var) in enumerate(rows[1:], 3):
            self._label(d, t).grid(row=r, column=0, sticky="w", pady=3)
            ttk.Spinbox(d, from_=-360 if t == "Rotation" else 1, to=360 if t == "Rotation" else 100, width=6, textvariable=var).grid(
                row=r, column=1, sticky="w", padx=4)
        self._label(d, "Position").grid(row=6, column=0, sticky="w", pady=3)
        ttk.Combobox(d, textvariable=position, state="readonly", width=14, values=["Center", "Tiled"]).grid(row=6, column=1, sticky="w", padx=4)
        self._label(d, "Layer").grid(row=7, column=0, sticky="w", pady=3)
        ttk.Combobox(d, textvariable=layer, state="readonly", width=24,
                     values=["Above the page content", "Behind the page content"]).grid(row=7, column=1, columnspan=3, sticky="w", padx=4)
        rng = self._range_row(d, 8, "all")

        def _wm_get():
            return {"kind": kind.get(), "text": text.get(), "img": img.get(), "opacity": opacity.get(), "angle": angle.get(),
                    "scale": scale.get(), "pos": position.get(), "layer": layer.get(), "color": list(ch["rgb"])}

        def _wm_set(p):
            kind.set(p["kind"])
            text.set(p["text"])
            img.set(p["img"])
            opacity.set(p["opacity"])
            angle.set(p["angle"])
            scale.set(p["scale"])
            position.set(p["pos"])
            layer.set(p["layer"])
            ch["rgb"] = tuple(p["color"])
            cb.config(bg=self._hex(ch["rgb"]), activebackground=self._hex(ch["rgb"]))
        self._preset_bar(d, 10, "watermark", _wm_get, _wm_set)

        def go():
            pages = self._pages_or_warn(rng.get(), d)
            if pages is None:
                return
            if kind.get() == "text" and not text.get().strip():
                messagebox.showinfo(APP, "Type the watermark text.", parent=d)
                return
            if kind.get() == "image" and not os.path.exists(img.get()):
                messagebox.showinfo(APP, "Choose an image file.", parent=d)
                return
            d.destroy()
            self.do_watermark(kind.get(), text.get(), img.get(), ch["rgb"], opacity.get(), angle.get(), scale.get(),
                              position.get(), layer.get().startswith("Behind"), pages)
        self._buttons(d, "Apply", go, 11)

    def do_watermark(self, kind, text, image_path, color, opacity, angle, scale, position, behind, pages):
        if kind == "text":
            base = make_watermark_image(text, color, opacity, angle)
        else:
            im = ImageOps.exif_transpose(Image.open(image_path)).convert("RGBA")
            im = apply_opacity(im, opacity)
            if abs(angle) % 360 > 0.01:
                im = im.rotate(angle, resample=BICUBIC, expand=True, fillcolor=(0, 0, 0, 0))
            base = im
        buf = io.BytesIO()
        base.save(buf, "PNG")
        data = buf.getvalue()
        ratio = base.height / base.width
        state = {"xref": 0}

        def work():
            for i in pages:
                pg = self.doc[i]
                v = pg.rect
                w = v.width * max(1, scale) / 100.0
                h = w * ratio
                if h > v.height * 0.95:
                    h = v.height * 0.95
                    w = h / ratio
                if position == "Tiled":
                    sx, sy = w * 1.25, h * 1.6
                    centers = []
                    row = 0
                    y = sy / 2
                    while y - h / 2 < v.height:
                        x = sx / 2 - (sx / 2 if row % 2 else 0)
                        while x - w / 2 < v.width:
                            centers.append((x, y))
                            x += sx
                        y += sy
                        row += 1
                else:
                    centers = [(v.width / 2, v.height / 2)]
                for cx, cy in centers:
                    r = fitz.Rect(cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2) * pg.derotation_matrix
                    r.normalize()
                    kw = dict(overlay=not behind, rotate=pg.rotation, keep_proportion=False)
                    if state["xref"]:
                        pg.insert_image(r, xref=state["xref"], **kw)
                    else:
                        state["xref"] = pg.insert_image(r, stream=data, **kw)
        return self._run_edit(work, thumbs=pages, msg=f"Watermark added to {len(pages)} page(s).")

    def dlg_background(self):
        if not self.need_doc():
            return
        d = self._dlg("Page background")
        self._label(d, "Background color").grid(row=0, column=0, sticky="w", pady=4)
        cb, ch = self._color_button(d, (1.0, 0.97, 0.80))
        cb.grid(row=0, column=1, sticky="w", padx=6)
        rng = self._range_row(d, 1, "all")

        def go():
            pages = self._pages_or_warn(rng.get(), d)
            if pages is None:
                return
            d.destroy()
            self.do_background(ch["rgb"], pages)
        self._buttons(d, "Apply", go, 3, colspan=3)

    def do_background(self, color, pages):
        def work():
            for i in pages:
                pg = self.doc[i]
                r = pg.rect * pg.derotation_matrix
                r.normalize()
                sh = pg.new_shape()
                sh.draw_rect(r)
                sh.finish(fill=color, color=None, width=0)
                sh.commit(overlay=False)
        return self._run_edit(work, thumbs=pages, msg=f"Background set on {len(pages)} page(s).")

    # ---- compress
    def dlg_compress(self):
        d = self._dlg("Compress PDF")
        sources = []                                       # paths; None = current document
        if self.doc is not None:
            sources.append(None)
        lb = tk.Listbox(d, width=60, height=6, activestyle="none")
        lb.grid(row=0, column=0, columnspan=3, pady=(0, 4))

        def refill():
            lb.delete(0, "end")
            for s in sources:
                lb.insert("end", "(the open document)" if s is None else s)

        def add():
            for p in filedialog.askopenfilenames(parent=d, filetypes=[("PDF files", "*.pdf")]):
                sources.append(p)
            refill()

        def remove():
            s = lb.curselection()
            if s:
                sources.pop(s[0])
                refill()
        f = tk.Frame(d, bg=BG)
        f.grid(row=1, column=0, columnspan=3, sticky="w")
        ttk.Button(f, text="Add PDF files...", command=add).pack(side="left")
        ttk.Button(f, text="Remove", command=remove).pack(side="left", padx=4)
        refill()
        level, strip, replace = tk.StringVar(value="Medium  (recommended)"), tk.BooleanVar(value=False), tk.BooleanVar(value=False)
        self._label(d, "Compression").grid(row=2, column=0, sticky="w", pady=(10, 3))
        ttk.Combobox(d, textvariable=level, state="readonly", width=24, values=list(COMPRESS_LEVELS)).grid(row=2, column=1, sticky="w", pady=(10, 3))
        tk.Checkbutton(d, text="Also remove document information (author, title ...)", variable=strip, bg=BG,
                       activebackground=BG, font=(UI_FONT, 9)).grid(row=3, column=0, columnspan=3, sticky="w")
        if self.doc is not None:
            tk.Checkbutton(d, text="Replace the open document instead of saving a new file", variable=replace, bg=BG,
                           activebackground=BG, font=(UI_FONT, 9)).grid(row=4, column=0, columnspan=3, sticky="w")

        def go():
            if not sources:
                messagebox.showinfo(APP, "Add a PDF file first.", parent=d)
                return
            only_current = sources == [None]
            out = None
            if only_current and not replace.get():
                base = os.path.splitext(os.path.basename(self.path))[0] if self.path else "document"
                out = filedialog.asksaveasfilename(parent=d, defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")],
                                                   initialfile=f"{base}_compressed.pdf")
                if not out:
                    return
            elif not only_current:
                out = filedialog.askdirectory(parent=d, title="Choose the folder for the compressed files")
                if not out:
                    return
            d.destroy()
            self.do_compress(sources, level.get(), strip.get(), out, replace.get() and only_current)
        self._buttons(d, "Compress", go, 5, colspan=3)

    def do_compress(self, sources, level, strip, out, replace=False):
        """sources: list of paths, None means the open document.  out: file (one source) or folder (several)."""
        report = []
        upd, close = self._progress("Compressing...", len(sources))
        try:
            for k, s in enumerate(sources):
                if s is None:
                    self.commit_floats()
                    data = self.doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE)
                    name = os.path.basename(self.path) if self.path else "document.pdf"
                else:
                    name = os.path.basename(s)
                    upd(k, len(sources), f"Compressing {name} ...")
                    with open(s, "rb") as fh:
                        data = fh.read()
                try:
                    new = compress_bytes(data, level, strip)
                except Exception as ex:
                    report.append(f"{name}: skipped ({ex})")
                    continue
                if replace and s is None:
                    self.snapshot()
                    self.doc = fitz.open(stream=new, filetype="pdf")
                    self._rebuild()
                else:
                    if len(sources) == 1 and out and not os.path.isdir(out):
                        target = out
                    else:
                        target = os.path.join(out, os.path.splitext(name)[0] + "_compressed.pdf")
                    with open(target, "wb") as fh:
                        fh.write(new)
                pct = (1 - len(new) / max(1, len(data))) * 100
                report.append(f"{name}: {human_size(len(data))} -> {human_size(len(new))}  ({pct:.0f}% smaller)")
                upd(k + 1, len(sources))
        finally:
            close()
        messagebox.showinfo(APP, "Done.\n\n{}".format("\n".join(report)))
        self.status.config(text="Compression finished.")
        return report

    # ---- flatten / protect / properties
    def do_flatten(self):
        if not self.need_doc():
            return
        n = sum(len(list(pg.annots() or [])) + len(list(pg.widgets() or [])) for pg in self.doc)
        if not n:
            messagebox.showinfo(APP, "There are no comments or form fields to flatten.")
            return
        if not messagebox.askyesno(APP, f"Flatten {n} comment(s) / form field(s) into the page?\n\n"
                                        "They can no longer be edited afterwards (you can still undo)."):
            return
        self._run_edit(lambda: self.doc.bake(annots=True, widgets=True), rebuild=True, msg="Document flattened.")

    def dlg_protect(self):
        if not self.need_doc():
            return
        d = self._dlg("Protect with a password")
        up, up2, op = tk.StringVar(), tk.StringVar(), tk.StringVar()
        self._label(d, "Password to open the file").grid(row=0, column=0, sticky="w", pady=3)
        ttk.Entry(d, textvariable=up, show="*", width=26).grid(row=0, column=1, padx=6)
        self._label(d, "Repeat the password").grid(row=1, column=0, sticky="w", pady=3)
        ttk.Entry(d, textvariable=up2, show="*", width=26).grid(row=1, column=1, padx=6)
        self._label(d, "Owner password (optional)").grid(row=2, column=0, sticky="w", pady=3)
        ttk.Entry(d, textvariable=op, show="*", width=26).grid(row=2, column=1, padx=6)
        self._label(d, "People who open the file may:", bold=True).grid(row=3, column=0, columnspan=2, sticky="w", pady=(10, 2))
        perms = {k: tk.BooleanVar(value=True) for k in ("print", "copy", "modify", "annotate")}
        for r, (k, t) in enumerate((("print", "Print"), ("copy", "Copy text and images"),
                                    ("modify", "Edit the document"), ("annotate", "Add comments"))):
            tk.Checkbutton(d, text=t, variable=perms[k], bg=BG, activebackground=BG, font=(UI_FONT, 9)).grid(
                row=4 + r, column=0, columnspan=2, sticky="w")
        self._label(d, "Permissions are only enforced when an owner password is set.", fg=TH.TERTIARY).grid(
            row=8, column=0, columnspan=2, sticky="w")

        def go():
            if not up.get():
                messagebox.showwarning(APP, "Type a password.", parent=d)
                return
            if up.get() != up2.get():
                messagebox.showwarning(APP, "The two passwords are different.", parent=d)
                return
            d.destroy()
            self.do_protect(up.get(), op.get(), {k: v.get() for k, v in perms.items()})
        self._buttons(d, "Protect and save", go, 9, colspan=2)

    def do_protect(self, user_pw, owner_pw="", allow=None, save=True):
        allow = allow or {"print": True, "copy": True, "modify": True, "annotate": True}
        p = getattr(fitz, "PDF_PERM_ACCESSIBILITY", 0)
        for k, c in (("print", ("PRINT", "PRINT_HQ")), ("copy", ("COPY",)), ("modify", ("MODIFY", "ASSEMBLE", "FORM")),
                     ("annotate", ("ANNOTATE",))):
            if allow.get(k, True):
                for name in c:
                    p |= getattr(fitz, "PDF_PERM_" + name, 0)
        self.protect = {"user_pw": user_pw, "owner_pw": owner_pw or user_pw + "#owner", "perms": p}
        self.modified = True
        self.update_ui()
        if save:
            self.save()
        self.status.config(text="Password protection will be used when saving.")

    def do_unprotect(self):
        """Take the password off and save the result as a NEW file (the original is not touched)."""
        if not self.need_doc():
            return
        enc = getattr(self.doc, "is_encrypted", False) or bool((self.doc.metadata or {}).get("encryption"))   # also files with only an owner password
        if self.protect is None and not enc:
            messagebox.showinfo(APP, "This document is not password protected.")
            return
        old = self.protect
        self.protect = None
        base = os.path.splitext(os.path.basename(self.path or "document.pdf"))[0]
        p = filedialog.asksaveasfilename(parent=self, title="Save the unprotected copy as", defaultextension=".pdf",
                                         filetypes=[("PDF files", "*.pdf")], initialfile=base + "_unlocked.pdf",
                                         initialdir=os.path.dirname(self.path) if self.path else None)
        if not p:
            self.protect = old                                   # cancelled: nothing changes
            return
        if self.path and os.path.normcase(os.path.abspath(p)) == os.path.normcase(os.path.abspath(self.path)):
            if not messagebox.askyesno(APP, "This replaces the protected original with an unprotected copy.\n\nContinue?"):
                self.protect = old
                return
        self._write(p)
        if self.modified:                                        # the save did not happen (error or declined)
            self.protect = old
            return
        self.path = p
        self.update_ui()
        self.status.config(text=_tr("Saved without a password:") + " " + os.path.basename(p))

    def dlg_properties(self):
        if not self.need_doc():
            return
        self.commit_floats()
        d = self._dlg("Document properties")
        meta = self.doc.metadata or {}
        fields = (("title", "Title"), ("author", "Author"), ("subject", "Subject"), ("keywords", "Keywords"))
        vs = {}
        for r, (k, t) in enumerate(fields):
            vs[k] = tk.StringVar(value=meta.get(k) or "")
            self._label(d, t).grid(row=r, column=0, sticky="w", pady=3)
            ttk.Entry(d, textvariable=vs[k], width=46).grid(row=r, column=1, padx=6)
        size = os.path.getsize(self.path) if self.path and os.path.exists(self.path) else None
        r0 = self.doc[0].rect
        info = [("Pages", str(len(self.doc))), ("Page size", f"{r0.width / 72 * 25.4:.0f} x {r0.height / 72 * 25.4:.0f} mm"),
                ("File size", human_size(size) if size else "-"), ("PDF version", meta.get("format") or "-"),
                ("Created with", meta.get("creator") or meta.get("producer") or "-"),
                ("Password protected", "yes" if self.protect else "no")]
        for r, (t, val) in enumerate(info, len(fields) + 1):
            self._label(d, t, fg=TH.TERTIARY).grid(row=r, column=0, sticky="w")
            self._label(d, val).grid(row=r, column=1, sticky="w", padx=6)

        def go():
            self.do_set_properties({k: vs[k].get() for k in vs})
            d.destroy()

        def clear():
            for v in vs.values():
                v.set("")
        f = self._buttons(d, "OK", go, 20, colspan=2)
        ttk.Button(f, text="Clear all", command=clear).pack(side="left", padx=(8, 0))

    def do_set_properties(self, values):
        def work():
            m = dict(self.doc.metadata or {})
            m.update(values)
            self.doc.set_metadata(m)
        self._run_edit(work, thumbs=[], msg="Document properties updated.")

    # ---- print
    def do_print(self):
        if not self.need_doc():
            return
        self.commit_floats()
        import claudius_print as cp
        names, default = cp.list_printers()
        n = len(self.doc)
        d = self._dlg("Print")
        pref = getattr(self, "cfg", {}).get("printer", "")
        pv = tk.StringVar(value=pref if pref in names else (default or (names[0] if names else "")))
        self._label(d, "Printer").grid(row=0, column=0, sticky="w", pady=3)
        if names:
            ttk.Combobox(d, textvariable=pv, values=names, state="readonly", width=34).grid(row=0, column=1, columnspan=3, sticky="w", pady=3)
        else:
            ttk.Entry(d, textvariable=pv, width=36).grid(row=0, column=1, columnspan=3, sticky="w", pady=3)
        self._label(d, "Copies").grid(row=1, column=0, sticky="w", pady=3)
        cv = tk.IntVar(value=getattr(self, "cfg", {}).get("copies", 1))
        ttk.Spinbox(d, from_=1, to=99, width=5, textvariable=cv).grid(row=1, column=1, sticky="w", pady=3)
        sel = self.selected_pages() if hasattr(self, "selected_pages") else []
        rv = self._range_row(d, 2, "all" if len(sel) <= 1 else pages_to_text(sel), "Pages")
        self._label(d, "Scale").grid(row=4, column=0, sticky="w", pady=3)
        fv = tk.StringVar(value="Fit to page")
        ttk.Combobox(d, textvariable=fv, values=["Fit to page", "Actual size"], state="readonly", width=16).grid(row=4, column=1, sticky="w", pady=3)

        def go():
            pages = self._pages_or_warn(rv.get(), d)
            if pages is None:
                return
            printer = pv.get().strip()
            try:
                copies = max(1, int(cv.get()))
            except (tk.TclError, ValueError):
                copies = 1
            fit = "actual" if fv.get() == "Actual size" else "fit"
            d.destroy()
            self.run_print(pages, printer, copies, fit)
        f = self._buttons(d, "Print", go, 5)

        def preview():
            pages = self._pages_or_warn(rv.get(), d)
            if pages is None:
                return
            fit = "actual" if fv.get() == "Actual size" else "fit"
            self.print_preview(pages, fit)
        ttk.Button(f, text="Preview...", command=preview).pack(side="left", padx=4)

    def run_print(self, pages, printer, copies=1, fit="fit"):
        import claudius_print as cp
        title = os.path.basename(self.path) if self.path else "Claudius PDF Editor"
        upd, close = self._progress("Printing...", len(pages) * copies)
        try:
            cp.print_document(self.doc, pages, printer, title, copies, fit, progress=upd)
        except Exception as ex:
            close()
            messagebox.showerror(APP, f"Could not print:\n{ex}")
            return False
        close()
        self.status.config(text="Sent to the printer.")
        return True

    # ================================================================== bookmarks panel
    def _build_sidebar_tabs(self, side):
        bar = tk.Frame(side, bg=SIDEBAR_BG)
        bar.pack(fill="x", padx=6, pady=(6, 2))
        self._sb_tabs = {}
        for key, text in (("pages", "Pages"), ("bookmarks", "Bookmarks"), ("comments", "Comments")):
            lb = tk.Label(bar, text=text, font=(UI_FONT, 9, "bold"), bg=SIDEBAR_BG, fg=TH.TERTIARY, cursor="hand2", padx=5, pady=3)
            lb.pack(side="left")
            lb.bind("<Button-1>", lambda e, k=key: self.set_sidebar(k))
            self._sb_tabs[key] = lb
        self._sb_tabs["pages"].config(fg=ACCENT)
        hide = tk.Label(bar, text="\u25C0", font=(UI_FONT, 10), bg=SIDEBAR_BG, fg=TH.TERTIARY, cursor="hand2", padx=6, pady=2)
        hide.pack(side="right")
        hide.bind("<Button-1>", lambda e: self.toggle_sidebar(False))
        hide.bind("<Enter>", lambda e: hide.config(fg=ACCENT))
        hide.bind("<Leave>", lambda e: hide.config(fg=TH.TERTIARY))
        self.bm_wrap = tk.Frame(side, bg=SIDEBAR_BG)
        self.bm_tree = ttk.Treeview(self.bm_wrap, show="tree", selectmode="browse")
        sb = ttk.Scrollbar(self.bm_wrap, orient="vertical", command=self.bm_tree.yview)
        self.bm_tree.config(yscrollcommand=sb.set)
        bf = tk.Frame(self.bm_wrap, bg=SIDEBAR_BG)
        bf.pack(side="bottom", fill="x", pady=4)
        for t, c in (("+ Add", self.bm_add), ("Rename", self.bm_rename), ("Delete", self.bm_delete)):
            b = ttk.Button(bf, text=t, command=c, width=7)
            b.pack(side="left", padx=2)
        sb.pack(side="right", fill="y")
        self.bm_tree.pack(side="left", fill="both", expand=True)
        self.bm_tree.bind("<<TreeviewSelect>>", self._bm_select)
        self._build_comments_panel(side)

    def toggle_sidebar(self, show=None):
        """hide / show the Pages - Bookmarks - Comments panel (Ctrl+B)"""
        if getattr(self, "_fs", None) or getattr(self, "presenting", False):
            return
        shown = self.side_frame.winfo_manager() == "pack"
        if show is None:
            show = not shown
        if show == shown:
            return
        if show:
            self.side_frame.pack(side="left", fill="y", before=self.page_area)
            self.side_sep.pack(side="left", fill="y", before=self.page_area)
        else:
            self.side_frame.pack_forget()
            self.side_sep.pack_forget()
        self.update_idletasks()
        self.status.config(text="" if show else _tr("Side panel hidden - press Ctrl+B to show it again."))
        if self.doc is not None and getattr(self, "fit_mode", None):
            self.after(60, lambda: self._apply_fit(render=True))

    def set_sidebar(self, mode):
        if self.side_frame.winfo_manager() != "pack":
            self.toggle_sidebar(True)
        self.sidebar_mode = mode
        for k, lb in self._sb_tabs.items():
            lb.config(fg=ACCENT if k == mode else TH.TERTIARY)
        for w in (self.pages_wrap, self.bm_wrap, self.cm_wrap):
            w.pack_forget()
        if mode == "pages":
            self.pages_wrap.pack(fill="both", expand=True)
        elif mode == "bookmarks":
            self.bm_wrap.pack(fill="both", expand=True)
            self._refresh_bookmarks()
        else:
            self.cm_wrap.pack(fill="both", expand=True)
            self._refresh_comments()

    def _refresh_bookmarks(self):
        if not hasattr(self, "bm_tree"):
            return
        t = self.bm_tree
        t.delete(*t.get_children())
        self.bm_items = {}
        if not self.doc:
            return
        try:
            toc = self.doc.get_toc(simple=True)
        except Exception:
            toc = []
        parents = {}
        for idx, (lvl, title, page) in enumerate(toc):
            parent = ""
            for l in range(lvl - 1, 0, -1):
                if l in parents:
                    parent = parents[l]
                    break
            iid = t.insert(parent, "end", text=f"{title}", open=True)
            parents[lvl] = iid
            for k in [k for k in parents if k > lvl]:
                del parents[k]
            self.bm_items[iid] = (idx, page)

    def _bm_select(self, _e=None):
        sel = self.bm_tree.selection()
        if sel and sel[0] in self.bm_items:
            page = self.bm_items[sel[0]][1]
            if page >= 1:
                self.goto(page - 1)

    def bm_add(self, title=None):
        if not self.need_doc():
            return
        if title is None:
            title = simpledialog.askstring(APP, "Bookmark name:", initialvalue=f"Page {self.page_no + 1}", parent=self)
        if not title:
            return
        toc = self.doc.get_toc(simple=True)
        toc.append([1, title, self.page_no + 1])
        toc.sort(key=lambda x: x[2])                         # keeps the existing order of equal pages (stable)
        self._run_edit(lambda: self.doc.set_toc(self._fix_toc(toc)), thumbs=[], msg="Bookmark added.")
        self._refresh_bookmarks()

    @staticmethod
    def _fix_toc(toc):
        out, prev = [], 0
        for lvl, title, page in toc:
            lvl = max(1, min(lvl, prev + 1))
            out.append([lvl, title, page])
            prev = lvl
        return out

    def _bm_current(self):
        sel = self.bm_tree.selection()
        return self.bm_items.get(sel[0]) if sel else None

    def bm_rename(self):
        cur = self._bm_current()
        if not cur:
            return
        toc = self.doc.get_toc(simple=True)
        name = simpledialog.askstring(APP, "New name:", initialvalue=toc[cur[0]][1], parent=self)
        if not name:
            return
        toc[cur[0]][1] = name
        self._run_edit(lambda: self.doc.set_toc(toc), thumbs=[])
        self._refresh_bookmarks()

    def bm_delete(self):
        cur = self._bm_current()
        if not cur:
            return
        toc = self.doc.get_toc(simple=True)
        i = cur[0]
        j = i + 1
        while j < len(toc) and toc[j][0] > toc[i][0]:
            j += 1
        new = toc[:i] + toc[j:]
        self._run_edit(lambda: self.doc.set_toc(self._fix_toc(new)), thumbs=[], msg="Bookmark deleted.")
        self._refresh_bookmarks()

